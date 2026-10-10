"""Génère les jeux d'essai FICTIFS de ``data/samples/`` (déterministe, idempotent).

Usage : ``python data/samples/generate_samples.py [--out DOSSIER]``

Tout est FICTIF : fournisseurs « fictif_grossiste_* », SKU « FICTIF-… », extensions
« Extension Fictive … », GTIN de test ``200…`` au checksum valide (SPEC §0.7), prix
inventés pour exercer les contrôles. Chaque ligne porte l'anomalie volontaire qu'elle
teste ; ``FICTIF_attendus.yaml`` donne le résultat attendu (lu par tests/test_importers.py).
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import io
import sys
from datetime import datetime
from pathlib import Path
from typing import Any
from xml.sax.saxutils import escape, quoteattr

import yaml

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "engine"))

from pokeshop.catalog import fictitious_gtin13, gtin_check_digit  # noqa: E402

J0 = "2026-10-04T06:00:00+02:00"
J1 = "2026-10-03T06:00:00+02:00"
STALE = "2026-10-02T06:00:00+02:00"
NOW_J0 = "2026-10-04T08:00:00+02:00"
NOW_J1 = "2026-10-03T08:00:00+02:00"


def g(n: int) -> str:
    """GTIN-13 FICTIF n° n."""
    return fictitious_gtin13(n)


def bad_check(gtin: str) -> str:
    """Même GTIN avec un chiffre de contrôle faux."""
    return gtin[:-1] + str((int(gtin[-1]) + 1) % 10)


def case_gtin(gtin13: str) -> str:
    """GTIN-14 FICTIF de carton (indicateur 1) dérivé d'un GTIN-13."""
    payload = "1" + gtin13[:-1]
    return payload + str(gtin_check_digit(payload))


# ----------------------------------------------------------------- CSV grossiste A

CSV_A_HEADER = [
    "sku", "ean", "designation", "langue", "extension", "format", "contenu", "scelle", "prix", "devise",
    "base_prix", "taux_tva", "unite_prix", "unites_par_colis", "colisage", "moq", "disponibilite", "stock",
    "allocation", "palier1_qte", "palier1_prix", "palier2_qte", "palier2_remise", "date_sortie",
    "pays_expedition", "date_maj", "fictif",
]

_A_BASE: dict[str, dict[str, str]] = {
    "FICTIF-A-001": {"ean": g(101), "designation": "Display 36 boosters Extension Fictive Alpha", "extension": "Fictive Alpha",
                     "format": "Display", "contenu": "36 boosters", "prix": "95,00", "colisage": "6", "moq": "6", "stock": "24",
                     "palier1_qte": "12", "palier1_prix": "92,00", "palier2_qte": "24", "palier2_remise": "5 %",
                     "date_sortie": "17.07.2026"},
    "FICTIF-A-002": {"ean": g(102), "designation": "Coffret Dresseur d'Élite Fictive Alpha", "extension": "Fictive Alpha",
                     "format": "ETB", "contenu": "9 boosters + accessoires", "prix": "38,50", "colisage": "10", "moq": "10", "stock": "40"},
    "FICTIF-A-003": {"ean": g(103), "designation": "Bundle 6 boosters Fictive Alpha", "extension": "Fictive Alpha",
                     "format": "Bundle", "contenu": "6 boosters", "prix": "21,00", "colisage": "12", "moq": "12", "stock": "36"},
    "FICTIF-A-004": {"ean": g(104), "designation": "Tripack Fictive Alpha", "extension": "Fictive Alpha",
                     "format": "Tripack", "contenu": "3 boosters + carte promo", "prix": "11,20", "colisage": "24", "moq": "24", "stock": "48"},
    "FICTIF-A-005": {"ean": g(105), "designation": "Coffret collection Fictive Bêta", "extension": "Fictive Bêta",
                     "format": "Coffret", "contenu": "4 boosters + promo", "prix": "18,90", "colisage": "6", "moq": "6", "stock": "18"},
    "FICTIF-A-006": {"ean": g(106), "designation": "Display 36 boosters Fictive Bêta", "extension": "Fictive Bêta",
                     "format": "Display", "contenu": "36 boosters", "prix": "98,00", "colisage": "6", "moq": "6", "stock": "12"},
    "FICTIF-A-007": {"ean": g(107), "designation": "ETB Fictive Bêta", "extension": "Fictive Bêta",
                     "format": "ETB", "contenu": "9 boosters", "prix": "39,00", "colisage": "10", "moq": "10", "stock": "20"},
    "FICTIF-A-008": {"ean": g(108), "designation": "Protège-cartes standard x65", "langue": "n/a", "extension": "",
                     "format": "Protège-cartes", "contenu": "65 protège-cartes", "prix": "4,20", "colisage": "20", "moq": "20", "stock": "200"},
    "FICTIF-A-009": {"ean": g(109), "designation": "Classeur 9 cases", "langue": "n/a", "extension": "",
                     "format": "Classeur", "contenu": "1 classeur 360 cartes", "prix": "9,80", "colisage": "10", "moq": "10", "stock": "30"},
    "FICTIF-A-016": {"ean": g(116), "designation": "Bundle 6 boosters Fictive Bêta", "extension": "Fictive Bêta",
                     "format": "Bundle", "contenu": "6 boosters", "prix": "21,50", "colisage": "12", "moq": "12", "stock": "24"},
    "FICTIF-A-025": {"ean": "", "designation": "Booster Fictive Gamma", "extension": "Fictive Gamma",
                     "format": "Booster", "contenu": "1 booster", "prix": "3,80", "colisage": "36", "moq": "36", "stock": "360"},
    "FICTIF-A-026": {"ean": g(126), "designation": "Booster Fictive Alpha", "extension": "Fictive Alpha",
                     "format": "Booster", "contenu": "1 booster", "prix": "3,90", "colisage": "36", "moq": "36", "stock": "720"},
    "FICTIF-A-027": {"ean": g(101), "designation": "Carton 6 displays Fictive Alpha", "extension": "Fictive Alpha",
                     "format": "Display", "contenu": "36 boosters", "prix": "570,00", "unite_prix": "carton",
                     "unites_par_colis": "6", "colisage": "", "moq": "", "stock": "4"},
    "FICTIF-A-028": {"ean": g(128), "designation": "Display 36 boosters Fictive Gamma", "extension": "Fictive Gamma",
                     "format": "Display", "contenu": "36 boosters", "prix": "97,00", "colisage": "6", "moq": "6",
                     "disponibilite": "sur allocation", "stock": "", "allocation": "12", "date_sortie": "06.11.2026"},
    "FICTIF-A-033": {"ean": g(133), "designation": "ETB Fictive Alpha édition 11 boosters", "extension": "Fictive Alpha",
                     "format": "ETB", "contenu": "11 boosters", "prix": "49,00", "colisage": "10", "moq": "10", "stock": "10"},
    "FICTIF-A-034": {"ean": g(134), "designation": "Display 36 boosters Fictive Gamma (lot B)", "extension": "Fictive Gamma",
                     "format": "Display", "contenu": "36 boosters", "prix": "99,00", "colisage": "6", "moq": "6", "stock": "6"},
    "FICTIF-A-038": {"ean": g(138), "designation": "Display 30 boosters Fictive Alpha (JP)", "langue": "Japonais",
                     "extension": "Fictive Alpha", "format": "Display", "contenu": "30 boosters", "prix": "72,00",
                     "colisage": "12", "moq": "12", "stock": "12"},
    "FICTIF-A-040": {"ean": g(140), "designation": "ETB Fictive Gamma", "extension": "Fictive Gamma",
                     "format": "ETB", "contenu": "9 boosters", "prix": "38,90", "colisage": "10", "moq": "10", "stock": "15"},
    "FICTIF-A-041": {"ean": g(141), "designation": "Bundle Fictive Gamma", "extension": "Fictive Gamma",
                     "format": "Bundle", "contenu": "6 boosters", "prix": "21,20", "colisage": "12", "moq": "12", "stock": "12"},
    "FICTIF-A-042": {"ean": g(142), "designation": "Coffret Fictive Gamma", "extension": "Fictive Gamma",
                     "format": "Coffret", "contenu": "4 boosters", "prix": "19,40", "colisage": "6", "moq": "6", "stock": "6"},
    "FICTIF-A-043": {"ean": g(143), "designation": "Tripack Fictive Gamma", "extension": "Fictive Gamma",
                     "format": "Tripack", "contenu": "3 boosters", "prix": "11,30", "colisage": "24", "moq": "24", "stock": "24"},
    "FICTIF-A-099": {"ean": g(199), "designation": "Tripack Fictive Bêta", "extension": "Fictive Bêta",
                     "format": "Tripack", "contenu": "3 boosters", "prix": "11,40", "colisage": "24", "moq": "24", "stock": "24"},
}

_A_DEFAULTS = {
    "langue": "Français", "scelle": "oui", "devise": "EUR", "base_prix": "HT", "taux_tva": "0", "unite_prix": "unité",
    "unites_par_colis": "", "disponibilite": "en stock", "allocation": "", "palier1_qte": "", "palier1_prix": "",
    "palier2_qte": "", "palier2_remise": "", "date_sortie": "", "pays_expedition": "FR", "fictif": "true",
}


def a_row(sku: str, base: str | None = None, /, *, ts: str = J0, **changes: str) -> dict[str, str]:
    """Ligne CSV A : valeurs par défaut + modèle ``base`` + modifications."""
    row = {h: "" for h in CSV_A_HEADER}
    row.update(_A_DEFAULTS)
    row.update(_A_BASE[base or sku])
    row["sku"] = sku
    row["date_maj"] = ts
    row.update(changes)
    return row


# (sku/ligne, anomalie volontaire, motifs de quarantaine attendus, anomalies attendues)
Spec = tuple[dict[str, str], str, list[str], list[str]]


def csv_a_j1() -> list[dict[str, str]]:
    """Import de la veille (référence) : 22 lignes propres."""
    skus = [
        "FICTIF-A-001", "FICTIF-A-002", "FICTIF-A-003", "FICTIF-A-004", "FICTIF-A-005", "FICTIF-A-006", "FICTIF-A-007",
        "FICTIF-A-008", "FICTIF-A-009", "FICTIF-A-016", "FICTIF-A-025", "FICTIF-A-026", "FICTIF-A-027", "FICTIF-A-028",
        "FICTIF-A-033", "FICTIF-A-034", "FICTIF-A-038", "FICTIF-A-040", "FICTIF-A-041", "FICTIF-A-042", "FICTIF-A-043",
        "FICTIF-A-099",
    ]
    return [a_row(sku, ts=J1) for sku in skus]


def csv_a_j0() -> list[Spec]:
    """Import du jour : 45 lignes, anomalies volontaires couvrant toutes les règles."""
    return [
        (a_row("FICTIF-A-001"), "Ligne propre, deux paliers (prix et %)", [], []),
        (a_row("FICTIF-A-002"), "Ligne propre", [], []),
        (a_row("FICTIF-A-003"), "Ligne propre", [], []),
        (a_row("FICTIF-A-004", prix="17,50"), "Hausse ×1,56 vs veille (signalée, pas bloquante)", [], ["PRICE_CHANGE_LARGE"]),
        (a_row("FICTIF-A-005"), "Ligne propre", [], []),
        (a_row("FICTIF-A-006", prix="980,00"), "Prix ×10 vs veille (98,00)", ["PRICE_ANOMALY"], []),
        (a_row("FICTIF-A-007", prix="3,90"), "Prix ÷10 vs veille (39,00)", ["PRICE_ANOMALY"], []),
        (a_row("FICTIF-A-008"), "Accessoire, langue « n/a » acceptée (langue NA)", [], []),
        (a_row("FICTIF-A-009"), "Accessoire sans extension (SANS_EXTENSION)", [], []),
        (a_row("FICTIF-A-010", "FICTIF-A-002", ean=g(110), devise="XYZ"), "Devise inconnue XYZ", ["UNKNOWN_CURRENCY"], []),
        (a_row("FICTIF-A-011", "FICTIF-A-002", ean=g(111), devise=""), "Devise absente", ["UNKNOWN_CURRENCY"], []),
        (a_row("FICTIF-A-012", "FICTIF-A-002", ean=g(112), base_prix=""), "HT/TTC absent", ["UNKNOWN_VAT_BASIS"], []),
        (a_row("FICTIF-A-013", "FICTIF-A-002", ean=g(113), base_prix="net"), "« net » ≠ HT : base inconnue", ["UNKNOWN_VAT_BASIS"], []),
        (a_row("FICTIF-A-014", "FICTIF-A-003", ean=g(114), prix="0"), "Prix zéro", ["NON_POSITIVE_PRICE"], []),
        (a_row("FICTIF-A-015", "FICTIF-A-003", ean=g(115), prix="-5,00"), "Prix négatif", ["NON_POSITIVE_PRICE"], []),
        (a_row("FICTIF-A-016"), "Doublon de SKU (valeurs différentes, 1/2)", ["DUPLICATE_SKU"], []),
        (a_row("FICTIF-A-016", prix="19,90"), "Doublon de SKU (valeurs différentes, 2/2)", ["DUPLICATE_SKU"], []),
        (a_row("FICTIF-A-017", "FICTIF-A-005", ean=bad_check(g(117))), "GTIN au checksum faux", ["INVALID_GTIN"], []),
        (a_row("FICTIF-A-018", "FICTIF-A-005", ean=g(118), langue="VO"), "Langue « VO » : inconnue", ["UNKNOWN_LANGUAGE"], []),
        (a_row("FICTIF-A-019", "FICTIF-A-005", ean=g(119), langue="FR/EN"), "Langues multiples", ["UNKNOWN_LANGUAGE"], []),
        (a_row("FICTIF-A-020", "FICTIF-A-003", ean=g(120), unite_prix="palette"), "Unité « palette » : unité/carton inconnue",
         ["UNKNOWN_UNIT_BASIS"], []),
        (a_row("FICTIF-A-021", "FICTIF-A-003", ean=g(121), unite_prix="carton", unites_par_colis="", colisage=""),
         "Prix au carton sans nombre d'unités ni colisage", ["UNKNOWN_UNIT_BASIS"], []),
        (a_row("FICTIF-A-022", "FICTIF-A-001", ean=g(122), palier1_prix="120,00"), "Palier plus cher que le prix de base",
         ["INVALID_TIER"], []),
        (a_row("FICTIF-A-023", "FICTIF-A-003", ean=g(123), prix="12.50"), "Point décimal dans un fichier à virgule (ambigu)",
         ["UNPARSEABLE_PRICE"], []),
        (a_row("FICTIF-A-024", "FICTIF-A-003", ean=g(124), ts=STALE), "Ligne datée de J-2 (> 24 h)", ["STALE_SOURCE"], []),
        (a_row("FICTIF-A-025"), "GTIN absent : offre acceptée, fiche en brouillon", [], ["MISSING_GTIN"]),
        (a_row("FICTIF-A-026", devise="CHF"), "Devise EUR -> CHF depuis la veille", ["CURRENCY_CHANGED"], []),
        (a_row("FICTIF-A-027"), "Prix au carton de 6 (570,00 = 6 × 95,00)", [], []),
        (a_row("FICTIF-A-028"), "Allocation annoncée sans accord écrit : ignorée", [], ["ALLOCATION_NOT_FIRM"]),
        (a_row("FICTIF-A-001", sku="", ean=g(129)), "SKU vide", ["MISSING_SKU"], []),
        (a_row("FICTIF-A-004", prix="17,50"), "Copie identique de la ligne FICTIF-A-004 : ignorée", [],
         ["DUPLICATE_IDENTICAL_ROW"]),
        (a_row("FICTIF-A-030", "FICTIF-A-026", ean=g(130), designation="Blister Fictive Alpha", format="Blister",
               contenu="1 booster + carte promo"), "Format « Blister » ambigu : offre en brouillon", [], ["UNKNOWN_FORMAT"]),
        (a_row("FICTIF-A-031", "FICTIF-A-002", ean=g(131), designation="ETB Extension Inconnue", extension="Extension Inconnue"),
         "Extension absente de la table d'alias", [], ["UNKNOWN_EXTENSION"]),
        (a_row("FICTIF-A-032", "FICTIF-A-002", ean=g(132), fictif="false"), "Ligne « non fictive » dans un jeu FICTIF",
         ["FICTIF_FLAG_MISMATCH"], []),
        (a_row("FICTIF-A-033"), "Même nom que FICTIF-A-002 mais 11 boosters (catalogue : AMBIGUOUS)", [], []),
        (a_row("FICTIF-A-034", base_prix="TTC"), "HT -> TTC depuis la veille", ["VAT_BASIS_CHANGED"], []),
        (a_row("FICTIF-A-035", "FICTIF-A-003", ean=g(135), unites_par_colis="6"), "Prix « unité » mais 6 unités par colis",
         ["UNIT_CONFLICT"], []),
        (a_row("FICTIF-A-036", "FICTIF-A-003", ean=g(136), moq="0"), "MOQ nul", ["INVALID_QUANTITY"], []),
        (a_row("FICTIF-A-037", "FICTIF-A-003", ean=g(137), moq="six"), "MOQ illisible", ["UNPARSEABLE_VALUE"], []),
        (a_row("FICTIF-A-038"), "Langue JP : offre valide, bloquée ensuite par le pricing (FR attendu)", [], []),
        (a_row("FICTIF-A-039", "FICTIF-A-003", ean=g(139), scelle="non"), "Produit non scellé : hors périmètre", [], ["NOT_SEALED"]),
        (a_row("FICTIF-A-040"), "Ligne propre", [], []),
        (a_row("FICTIF-A-041"), "Ligne propre", [], []),
        (a_row("FICTIF-A-042"), "Ligne propre", [], []),
        (a_row("FICTIF-A-043"), "Ligne propre", [], []),
    ]


def csv_a_incomplete() -> list[dict[str, str]]:
    """Import tronqué : 5 lignes au lieu de 22 la veille."""
    return [a_row(sku) for sku in ("FICTIF-A-001", "FICTIF-A-002", "FICTIF-A-003", "FICTIF-A-005", "FICTIF-A-040")]


def write_csv(rows: list[dict[str, str]]) -> bytes:
    buf = io.StringIO()
    writer = csv.DictWriter(buf, fieldnames=CSV_A_HEADER, delimiter=";", lineterminator="\n")
    writer.writeheader()
    writer.writerows(rows)
    return buf.getvalue().encode("utf-8")


# --------------------------------------------------------------- XLSX grossiste B

XLSX_B_HEADER = [
    "Référence", "EAN", "Article", "Langue", "Série", "Type", "Contenu", "Prix CHF TTC", "TVA", "Unité de prix",
    "Unités/carton", "Minimum", "Stock", "Statut", "Sortie", "Fictif",
]
XLSX_B_EXPORT = datetime(2026, 10, 4, 7, 30)  # heure locale (Europe/Zurich) dans la cellule B2


def b_row(ref: str, **cells: Any) -> dict[str, Any]:
    base: dict[str, Any] = {
        "Référence": ref, "EAN": None, "Article": None, "Langue": "Français", "Série": "Fictive Alpha", "Type": "Display",
        "Contenu": "36 boosters", "Prix CHF TTC": 105.9, "TVA": 0.081, "Unité de prix": "pièce", "Unités/carton": 6,
        "Minimum": 6, "Stock": 12, "Statut": "Disponible", "Sortie": None, "Fictif": "oui",
    }
    base.update(cells)
    return base


def xlsx_b() -> list[Spec | None]:
    """Tarif Excel : 23 lignes + 1 ligne vide (None)."""
    return [
        (b_row("FICTIF-B-001", EAN=g(201), Article="Display Fictive Alpha", Sortie=datetime(2026, 7, 17)),
         "Ligne propre (nombres Excel, date Excel)", [], []),
        (b_row("FICTIF-B-002", EAN=int(g(202)), Article="ETB Fictive Alpha", Type="ETB", Contenu="9 boosters", **{"Prix CHF TTC": 45.9},
               Minimum=10, **{"Unités/carton": 10}, Stock="10+"), "EAN numérique ; stock « 10+ »", [], ["APPROXIMATE_STOCK"]),
        (b_row("FICTIF-B-003", EAN=g(203), Article="Bundle Fictive Bêta", Série="Fictive Bêta", Type="Bundle", Contenu="6 boosters",
               **{"Prix CHF TTC": 24.5}, Minimum=12, **{"Unités/carton": 12}), "Ligne propre", [], []),
        (b_row("FICTIF-B-004", EAN=g(204), Article="Tripack Fictive Bêta", Série="Fictive Bêta", Type="Tripack", Contenu="3 boosters",
               **{"Prix CHF TTC": 12.9}, Stock=0, Statut="Rupture"), "Rupture déclarée (stock 0)", [], []),
        (b_row("FICTIF-B-005", EAN=g(205), Article="Coffret Fictive Gamma", Série="Fictive Gamma", Type="Coffret", Contenu="4 boosters",
               **{"Prix CHF TTC": 22.9}, Statut="Précommande", Sortie="06.11.2026"), "Précommande, date texte", [], []),
        (b_row("FICTIF-B-006", EAN=g(206), Article="Booster Fictive Gamma", Série="Fictive Gamma", Type="Booster", Contenu="1 booster",
               **{"Prix CHF TTC": 4.9}, Minimum=36, **{"Unités/carton": 36}, Stock=360), "Ligne propre", [], []),
        (b_row("FICTIF-B-007", EAN=g(207), Article="Deck box", Langue=None, Série=None, Type="Deck box", Contenu="1 boîte 100 cartes",
               **{"Prix CHF TTC": 7.9}, Minimum=10, **{"Unités/carton": 10}), "Accessoire sans langue ni série", [], []),
        (b_row("FICTIF-B-008", EAN=g(208), Article="Display Fictive Bêta JP", Langue="Japonais", Série="Fictive Bêta",
               **{"Prix CHF TTC": 109.9}), "Langue JP : valide ici, bloquée par le pricing FR", [], []),
        (b_row("FICTIF-B-009", EAN=case_gtin(g(201)), Article="Carton 6 displays Fictive Alpha", **{"Prix CHF TTC": "1'299.00"},
               **{"Unité de prix": "carton"}, Stock=2), "Carton : GTIN-14 de colis, prix « 1'299.00 »", [], ["CASE_LEVEL_GTIN"]),
        (b_row("FICTIF-B-010", EAN=g(210), Article="Display Fictive Alpha (promo)", **{"Prix CHF TTC": "CHF 0.00"}),
         "Prix zéro", ["NON_POSITIVE_PRICE"], []),
        (b_row("FICTIF-B-011", EAN=g(211), Article="Display Fictive Alpha", **{"Prix CHF TTC": "€ 45.00"}),
         "Symbole € dans un tarif CHF", ["CURRENCY_CONFLICT"], []),
        (b_row("FICTIF-B-012", EAN=g(212), Article="Display Fictive Alpha", **{"Prix CHF TTC": "$45"}),
         "Symbole $ ambigu", ["UNKNOWN_CURRENCY"], []),
        (b_row("FICTIF-B-013", EAN=bad_check("2000013" + str(gtin_check_digit("2000013"))), Article="Display Fictive Alpha"),
         "GTIN-8 au checksum faux", ["INVALID_GTIN"], []),
        (b_row("FICTIF-B-014", EAN=g(214), Article="ETB Fictive Gamma", Langue=None, Série="Fictive Gamma", Type="ETB",
               Contenu="9 boosters", **{"Prix CHF TTC": 46.9}), "Langue vide pour un produit de cartes", ["UNKNOWN_LANGUAGE"], []),
        (b_row("FICTIF-B-015", EAN=g(215), Article="Display Fictive Gamma", Série="Fictive Gamma", Stock=-3),
         "Stock négatif : ignoré (pas de stock inventé)", [], ["UNPARSEABLE_OPTIONAL"]),
        (b_row("FICTIF-B-016", EAN=g(216), Article="Booster Fictive Alpha", Type="Booster", Contenu="1 booster",
               **{"Prix CHF TTC": 4.9, "Unité de prix": "display"}), "Unité « display » pour un booster : ambiguë",
         ["UNKNOWN_UNIT_BASIS"], []),
        (b_row("FICTIF-B-017", EAN=g(217), Article="Display Fictive Alpha", TVA="abc"), "TVA illisible", ["UNPARSEABLE_VALUE"], []),
        (b_row("FICTIF-B-018", EAN=g(218), Article="Display Fictive Bêta", Série="Fictive Bêta", Sortie="bientôt"),
         "Date de sortie illisible : ignorée", [], ["UNPARSEABLE_OPTIONAL"]),
        (b_row("FICTIF-B-019", EAN=g(219), Article="Display Fictive Bêta", Série="Fictive Bêta", **{"Prix CHF TTC": None}),
         "Prix vide", ["MISSING_PRICE"], []),
        (b_row("FICTIF-B-003", EAN=g(203), Article="Bundle Fictive Bêta", Série="Fictive Bêta", Type="Bundle", Contenu="6 boosters",
               **{"Prix CHF TTC": 24.5}, Minimum=12, **{"Unités/carton": 12}), "Copie identique de FICTIF-B-003 : ignorée", [],
         ["DUPLICATE_IDENTICAL_ROW"]),
        None,  # ligne vide volontaire
        (b_row("FICTIF-B-021", EAN=g(221), Article="Display Fictive Gamma IT", Langue="Italien", Série="Fictive Gamma"),
         "Langue IT : valide ici, bloquée par le pricing FR", [], []),
        (b_row("FICTIF-B-022", EAN=g(222), Article="Display Fictive Gamma", Série="Fictive Gamma", Fictif="non"),
         "Ligne « non fictive » dans un jeu FICTIF", ["FICTIF_FLAG_MISMATCH"], []),
        (b_row("FICTIF-B-023", EAN=g(223), Article="Coffret Fictive Alpha", Type="Coffret", Contenu="4 boosters",
               **{"Prix CHF TTC": 24.9}, TVA="8.1 %"), "TVA en texte « 8.1 % »", [], []),
    ]


def write_xlsx(rows: list[Spec | None], path: Path) -> None:
    import openpyxl  # import local : seulement pour la génération

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tarif"
    ws["A1"] = "FICTIF — Tarif grossiste B (jeu d'essai, aucune donnée réelle)"
    ws["A2"] = "Export du"
    ws["B2"] = XLSX_B_EXPORT
    ws.append([])
    ws.append(XLSX_B_HEADER)
    for spec in rows:
        if spec is None:
            ws.append([None] * len(XLSX_B_HEADER))
            continue
        ws.append([spec[0][h] for h in XLSX_B_HEADER])
    wb.properties.creator = "FICTIF generate_samples.py"
    wb.properties.created = datetime(2026, 10, 4, 0, 0)
    wb.properties.modified = datetime(2026, 10, 4, 0, 0)
    wb.save(path)


# ----------------------------------------------------------------- XML grossiste C

XML_C_EXPORT = "2026-10-04T05:30:00+02:00"
XML_C_STALE_EXPORT = "2026-10-01T05:30:00+02:00"


def c_offer(
    sku: str | None,
    *,
    ean: str | None,
    designation: str,
    langue: str = "Français",
    extension: str = "Fictive Alpha",
    fmt: str = "Display",
    contenu: str = "36 boosters",
    prix: str = "96.40",
    devise: str = "EUR",
    base: str = "HT",
    unite: str = "unité",
    unites: str = "1",
    carton: str = "6",
    moq: str = "6",
    stock: str = "30",
    statut: str = "en stock",
    paliers: list[tuple[str, str]] | None = None,
    allocation: str | None = None,
    sortie: str | None = None,
) -> str:
    attr = f" sku={quoteattr(sku)}" if sku is not None else ""
    parts = [f"  <offre{attr}>"]
    if ean is not None:
        parts.append(f"    <ean>{escape(ean)}</ean>")
    parts += [
        f"    <designation>{escape(designation)}</designation>",
        f"    <langue>{escape(langue)}</langue>",
        f"    <extension>{escape(extension)}</extension>",
        f"    <format>{escape(fmt)}</format>",
        f"    <contenu>{escape(contenu)}</contenu>",
        "    <scelle>oui</scelle>",
        f"    <prix devise={quoteattr(devise)} base={quoteattr(base)}>{escape(prix)}</prix>",
        "    <tva>0</tva>",
        f"    <conditionnement unite={quoteattr(unite)} unites={quoteattr(unites)} carton={quoteattr(carton)}/>",
        f"    <moq>{escape(moq)}</moq>",
        f"    <stock statut={quoteattr(statut)}>{escape(stock)}</stock>",
    ]
    if paliers:
        parts.append("    <paliers>")
        parts += [f"      <palier qte={quoteattr(q)} prix={quoteattr(p)}/>" for q, p in paliers]
        parts.append("    </paliers>")
    if allocation is not None:
        parts.append(f"    <allocation>{escape(allocation)}</allocation>")
    if sortie is not None:
        parts.append(f"    <sortie>{escape(sortie)}</sortie>")
    parts.append('    <expedition pays="FR" incoterm="DAP"/>')
    parts.append("  </offre>")
    return "\n".join(parts)


def xml_c() -> list[Spec]:
    """Flux XML : 23 offres."""
    specs: list[tuple[str, str, list[str], list[str]]] = [
        (c_offer("FICTIF-C-001", ean=g(301), designation="Display 36 boosters Extension Fictive Alpha",
                 paliers=[("12", "94.00"), ("24", "92.50")], sortie="2026-07-17"), "Ligne propre, deux paliers", [], []),
        (c_offer("FICTIF-C-002", ean=g(302), designation="ETB Fictive Bêta", extension="Fictive Bêta", fmt="ETB",
                 contenu="9 boosters", prix="37.90", carton="10", moq="10"), "Ligne propre", [], []),
        (c_offer("FICTIF-C-003", ean=g(303), designation="Bundle Fictive Gamma", extension="Fictive Gamma", fmt="Bundle",
                 contenu="6 boosters", prix="20.80", carton="12", moq="12"), "Ligne propre", [], []),
        (c_offer("FICTIF-C-004", ean=g(304), designation="Tripack Fictive Alpha", fmt="Tripack", contenu="3 boosters",
                 prix="11.10", carton="24", moq="24"), "Ligne propre", [], []),
        (c_offer("FICTIF-C-005", ean=g(305), designation="Coffret Fictive Bêta", extension="Fictive Bêta", fmt="Coffret",
                 contenu="4 boosters", prix="18.70"), "Ligne propre", [], []),
        (c_offer("FICTIF-C-006", ean=g(306), designation="Booster Fictive Alpha", fmt="Booster", contenu="1 booster",
                 prix="3.85", carton="36", moq="36", stock="", statut="allocation", allocation="72"),
         "Allocation annoncée sans accord écrit : ignorée", [], ["ALLOCATION_NOT_FIRM"]),
        (c_offer("FICTIF-C-007", ean=g(307), designation="Protège-cartes x65", langue="multilingue", extension="",
                 fmt="Sleeves", contenu="65 protège-cartes", prix="4.05", carton="20", moq="20"),
         "Accessoire multilingue (langue NA)", [], []),
        (c_offer("FICTIF-C-008", ean=g(308), designation="Display Fictive Gamma EN", langue="Anglais", extension="Fictive Gamma",
                 prix="89.00"), "Langue EN : valide ici, bloquée par le pricing FR", [], []),
        (c_offer("FICTIF-C-009", ean=g(309), designation="Display Fictive Alpha", devise="GBP"), "Devise GBP non autorisée",
         ["UNKNOWN_CURRENCY"], []),
        (c_offer("FICTIF-C-010", ean=g(310), designation="Display Fictive Alpha", base=""), "Base HT/TTC vide",
         ["UNKNOWN_VAT_BASIS"], []),
        (c_offer("FICTIF-C-011", ean=g(311), designation="Display Fictive Alpha", prix="0.00"), "Prix zéro",
         ["NON_POSITIVE_PRICE"], []),
        (c_offer("FICTIF-C-012", ean=g(312), designation="ETB Fictive Alpha", fmt="ETB", contenu="9 boosters", prix="37.50"),
         "Doublon de SKU (1/2)", ["DUPLICATE_SKU"], []),
        (c_offer("FICTIF-C-012", ean=g(312), designation="ETB Fictive Alpha", fmt="ETB", contenu="9 boosters", prix="36.90"),
         "Doublon de SKU (2/2)", ["DUPLICATE_SKU"], []),
        (c_offer("FICTIF-C-013", ean=bad_check(g(313)), designation="Display Fictive Alpha"), "GTIN au checksum faux",
         ["INVALID_GTIN"], []),
        (c_offer("FICTIF-C-014", ean=g(314), designation="Display Fictive Alpha", langue="multilingue"),
         "« multilingue » refusé pour un produit de cartes", ["UNKNOWN_LANGUAGE"], []),
        (c_offer("FICTIF-C-015", ean=g(315), designation="Display Fictive Alpha", paliers=[("12", "95.00"), ("24", "97.00")]),
         "Paliers : prix qui remonte", ["INVALID_TIER"], []),
        (c_offer("FICTIF-C-016", ean=g(316), designation="Display Fictive Alpha", paliers=[("", "90.00")]),
         "Palier sans quantité", ["INVALID_TIER"], []),
        (c_offer("FICTIF-C-017", ean=g(317), designation="Display Fictive Alpha", prix="abc"), "Prix illisible",
         ["UNPARSEABLE_PRICE"], []),
        (c_offer("FICTIF-C-018", ean=g(318), designation="Display Fictive Alpha", unite="carton", unites="", carton=""),
         "Prix au carton sans nombre d'unités ni colisage", ["UNKNOWN_UNIT_BASIS"], []),
        (c_offer(None, ean=g(319), designation="Display Fictive Alpha"), "Attribut sku absent", ["MISSING_SKU"], []),
        (c_offer("FICTIF-C-020", ean=g(320), designation="Carton 6 displays Fictive Alpha", prix="570.00", unite="carton",
                 unites="6", carton="6", moq="1", stock="3"), "Prix au carton (6 displays)", [], []),
        (c_offer("FICTIF-C-021", ean=g(321), designation="Collection Classeur Fictive Alpha", fmt="Collection Classeur",
                 contenu="classeur + 5 boosters", prix="27.40"), "Coffret « Collection Classeur »", [], []),
        (c_offer("FICTIF-C-022", ean=g(322), designation="Coffret ETB Fictive Gamma", extension="Fictive Gamma", fmt="Coffret ETB",
                 contenu="9 boosters", prix="38.20", stock="", statut="rupture"), "« Coffret ETB » = ETB ; rupture", [], []),
    ]
    return [({"xml": x}, note, q, a) for x, note, q, a in specs]


def write_xml(specs: list[Spec], exported_at: str) -> bytes:
    body = "\n".join(spec[0]["xml"] for spec in specs)
    text = (
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        "<!-- FICTIF : flux de démonstration, aucune donnée réelle (data/samples/README.md) -->\n"
        f'<catalogue fournisseur="FICTIF_GROSSISTE_C" exporte_le="{exported_at}" fictif="true">\n'
        f"{body}\n"
        "</catalogue>\n"
    )
    return text.encode("utf-8")


# ------------------------------------------------------- tarif email assisté (D)

ASSISTED_SOURCE = b"FICTIF - email tarif grossiste D du 4.10.2026 (contenu de demonstration)"


def assisted_doc() -> dict[str, Any]:
    return {
        "supplier_id": "fictif_grossiste_d",
        "source_kind": "EMAIL",
        "source_ref": "FICTIF-tarif-0001@exemple.invalid",
        "document_sha256": hashlib.sha256(ASSISTED_SOURCE).hexdigest(),
        "received_at": "2026-10-04T07:00:00+02:00",
        "document_date": "2026-10-04",
        "currency": "EUR",
        "price_basis": "HT",
        "vat_rate": "0",
        "decimal_separator": ",",
        "declared_total": "1 316,80",
        "declared_shipping": "35,00",
        "extracted_by": "FICTIF — saisie manuelle de démonstration",
        "ship_from_country": "FR",
        "fictif": True,
        "extension_tables": ["data/samples/FICTIF_extensions_aliases.yaml"],
        "lines": [
            {"line_no": 1, "supplier_sku": "FICTIF-D-001", "designation": "Display 36 boosters Fictive Alpha", "gtin": g(401),
             "language": "FR", "extension": "Fictive Alpha", "format": "Display", "content": "36 boosters", "sealed": True,
             "quantity": "2", "unit_basis": "unité", "unit_price": "95,00", "line_total": "190,00"},
            {"line_no": 2, "supplier_sku": "FICTIF-D-002", "designation": "ETB Fictive Alpha", "gtin": g(402),
             "language": "FR", "extension": "Fictive Alpha", "format": "ETB", "content": "9 boosters", "sealed": True,
             "quantity": "10", "unit_basis": "unité", "unit_price": "38,00", "line_total": "380,00"},
            {"line_no": 3, "supplier_sku": "FICTIF-D-003", "designation": "Bundle Fictive Bêta", "gtin": g(403),
             "language": "Français", "extension": "Fictive Bêta", "format": "Bundle", "content": "6 boosters", "sealed": True,
             "quantity": "12", "unit_basis": "unité", "unit_price": "21,00", "line_total": "252,00"},
            {"line_no": 4, "supplier_sku": "FICTIF-D-004", "designation": "Carton 24 tripacks Fictive Bêta", "gtin": g(404),
             "language": "FR", "extension": "Fictive Bêta", "format": "Tripack", "content": "3 boosters", "sealed": True,
             "quantity": "1", "unit_basis": "carton", "units_per_pack": "24", "unit_price": "268,80", "line_total": "268,80"},
            {"line_no": 5, "supplier_sku": "FICTIF-D-005", "designation": "Protège-cartes x65", "gtin": g(405),
             "format": "Protège-cartes", "content": "65 protège-cartes", "sealed": True,
             "quantity": "20", "unit_basis": "unité", "unit_price": "4,10", "line_total": "80,00"},
            {"line_no": 6, "supplier_sku": "FICTIF-D-006", "designation": "Coffret Fictive Gamma", "gtin": None,
             "language": "FR", "extension": "Fictive Gamma", "format": "Coffret", "content": "4 boosters", "sealed": True,
             "quantity": "6", "unit_basis": "unité", "unit_price": "18,50", "line_total": "111,00"},
        ],
    }


# ------------------------------------------------------------------- écriture

FILES = {
    "csv_j1": "FICTIF_offres_grossiste_a_J-1.csv",
    "csv_j0": "FICTIF_offres_grossiste_a.csv",
    "csv_incomplete": "FICTIF_offres_grossiste_a_incomplet.csv",
    "xlsx": "FICTIF_tarif_grossiste_b.xlsx",
    "xml": "FICTIF_flux_grossiste_c.xml",
    "xml_stale": "FICTIF_flux_grossiste_c_perime.xml",
    "assisted": "FICTIF_tarif_email_assiste.yaml",
    "expected": "FICTIF_attendus.yaml",
    "readme": "README.md",
}


def _line_entries(specs: list[Spec | None], first_line: int, sku_of: Any) -> list[dict[str, Any]]:
    out: list[dict[str, Any]] = []
    line = first_line
    for spec in specs:
        if spec is None:
            line += 1
            continue
        row, note, reasons, anomalies = spec
        out.append({"line": line, "sku": sku_of(row), "note": note, "quarantine": reasons, "anomalies": anomalies})
        line += 1
    return out


def _xml_sku(row: dict[str, str]) -> str | None:
    import re

    m = re.search(r'sku="([^"]+)"', row["xml"])
    return m.group(1) if m else None


def expected() -> dict[str, Any]:
    a = csv_a_j0()
    b = xlsx_b()
    c = xml_c()
    return {
        "version": "2026-10-04",
        "statut": "FICTIF — résultats attendus des jeux d'essai",
        FILES["csv_j1"]: {"mapping": "fictif_grossiste_a", "now": NOW_J1, "status": "ACCEPTED", "rows_read": 22, "offers": 22},
        FILES["csv_j0"]: {
            "mapping": "fictif_grossiste_a", "now": NOW_J0, "baseline": FILES["csv_j1"], "status": "PARTIAL",
            "rows_read": len(a), "lines": _line_entries(list(a), 2, lambda r: r["sku"] or None),
            "import_anomalies": ["MISSING_SINCE_PREVIOUS", "NEW_SINCE_PREVIOUS", "QUARANTINE_RATIO_EXCEEDED"],
        },
        FILES["csv_incomplete"]: {
            "mapping": "fictif_grossiste_a", "now": NOW_J0, "baseline": FILES["csv_j1"], "status": "QUARANTINED",
            "rows_read": 5, "import_anomalies": ["INCOMPLETE_IMPORT", "MISSING_SINCE_PREVIOUS"],
        },
        FILES["xlsx"]: {
            "mapping": "fictif_grossiste_b", "now": NOW_J0, "status": "PARTIAL", "rows_read": len([s for s in b if s]),
            "lines": _line_entries(b, 5, lambda r: r["Référence"]),
            "import_anomalies": ["BLANK_ROWS", "QUARANTINE_RATIO_EXCEEDED"],
        },
        FILES["xml"]: {
            "mapping": "fictif_grossiste_c", "now": NOW_J0, "status": "PARTIAL", "rows_read": len(c),
            "lines": _line_entries(list(c), 1, _xml_sku),
            "import_anomalies": ["QUARANTINE_RATIO_EXCEEDED"],
        },
        FILES["xml_stale"]: {
            "mapping": "fictif_grossiste_c", "now": NOW_J0, "status": "QUARANTINED", "rows_read": len(c),
            "import_anomalies": ["STALE_SNAPSHOT"],
        },
        FILES["assisted"]: {
            "now": NOW_J0, "blocking": ["LINE_TOTAL_MATCH@5"],
            "warnings": ["GTIN@6"],
        },
    }


def readme(exp: dict[str, Any]) -> str:
    parts = [
        "# Jeux d'essai FICTIFS — imports fournisseurs",
        "",
        "> **Tout est FICTIF.** Fournisseurs `fictif_grossiste_*`, SKU `FICTIF-…`, extensions « Extension Fictive … »",
        "> (table `FICTIF_extensions_aliases.yaml`), GTIN de test `200…` au checksum valide (plage GS1 « restricted",
        "> circulation », SPEC §0.7), prix inventés. Aucun tarif, EAN ou fournisseur réel. Référence : 4 octobre 2026.",
        "",
        "Fichiers générés par `python data/samples/generate_samples.py` (déterministe ; `tests/test_importers.py` vérifie",
        "que les fichiers versionnés correspondent au générateur). Résultats attendus : `FICTIF_attendus.yaml`.",
        "",
        "| Fichier | Dictionnaire | Format | Rôle |",
        "|---|---|---|---|",
        f"| `{FILES['csv_j1']}` | `fictif_grossiste_a` | CSV `;`, virgule décimale, EUR HT | Import de la veille (référence ×10, devise, HT/TTC, lignes) |",
        f"| `{FILES['csv_j0']}` | `fictif_grossiste_a` | idem | Import du jour : une anomalie volontaire par ligne |",
        f"| `{FILES['csv_incomplete']}` | `fictif_grossiste_a` | idem | Import tronqué (5 lignes au lieu de 22) |",
        f"| `{FILES['xlsx']}` | `fictif_grossiste_b` | Excel, feuille « Tarif », en-tête ligne 4, export en B2, CHF TTC | Tarif type distributeur suisse |",
        f"| `{FILES['xml']}` | `fictif_grossiste_c` | XML, attributs, paliers répétés | Flux type API |",
        f"| `{FILES['xml_stale']}` | `fictif_grossiste_c` | idem, exporté le 1.10.2026 | Capture > 24 h |",
        f"| `{FILES['assisted']}` | — (import assisté) | fiche YAML d'un tarif email | Contrôle unités et totaux |",
        "| `FICTIF_extensions_aliases.yaml` | — | YAML | Extensions fictives (jamais en production) |",
        "",
        "Couverture des règles de quarantaine (SPEC §2.5, BP §13) : devise inconnue, HT/TTC inconnu, prix 0 ou négatif,",
        "prix ×10 ou ÷10 vs dernier import, doublons, import incomplet, GTIN invalide, langue inconnue, capture > 24 h,",
        "unité vs carton, paliers incohérents, devise ou base HT/TTC changées, ligne fictive dans un flux réel.",
        "",
    ]
    for key in ("csv_j0", "xlsx", "xml"):
        name = FILES[key]
        spec = exp[name]
        parts += [
            f"## `{name}` — statut attendu {spec['status']} ({spec['rows_read']} lignes lues)",
            "",
            "| Ligne | SKU | Anomalie volontaire | Quarantaine attendue | Signal attendu |",
            "|---|---|---|---|---|",
        ]
        for entry in spec["lines"]:
            parts.append(
                f"| {entry['line']} | {entry['sku'] or '(vide)'} | {entry['note']} | "
                f"{', '.join(entry['quarantine']) or '— (offre valide)'} | {', '.join(entry['anomalies']) or '—'} |"
            )
        parts += ["", f"Signaux de niveau import : {', '.join(spec['import_anomalies'])}.", ""]
    parts += [
        f"## `{FILES['csv_incomplete']}` et `{FILES['xml_stale']}`",
        "",
        "- Import tronqué : 5 lignes < 80 % des 22 lignes de la veille (hypothèse de seuil) => `INCOMPLETE_IMPORT`,",
        "  import entier en quarantaine, offres de la veille conservées (elles vieillissent), stock local intact.",
        "- Flux périmé : `exporte_le` = 1.10.2026 (> 24 h) => `STALE_SNAPSHOT`, import entier en quarantaine.",
        "",
        f"## `{FILES['assisted']}` (import assisté email/PDF)",
        "",
        "- Ligne 5 : 20 × 4,10 = 82,00 mais total recopié 80,00 => contrôle `LINE_TOTAL_MATCH` bloquant.",
        "- Ligne 6 : GTIN absent => avertissement (fiche brouillon), non bloquant.",
        "- Total du document : lignes + port 35,00 = 1 316,80 (cohérent avec les totaux recopiés).",
        "- La revue reste « A_VALIDER_HUMAINEMENT » ; `approve_review` refuse tant qu'un contrôle bloquant subsiste.",
        "",
        "## Validation humaine requise",
        "",
        "- [ ] Aucune décision commerciale : ces fichiers ne servent qu'aux tests.",
        "- [ ] Vérifier qu'aucun fichier de ce dossier n'est utilisé par un dictionnaire non FICTIF ni publié.",
        "- [ ] Remplacer ces jeux par l'exemple de fichier réel du premier fournisseur dès réception (BP §13), archivé hors dépôt.",
        "",
    ]
    return "\n".join(parts)


def generate(out: Path) -> dict[str, Path]:
    """Écrit tous les fichiers dans ``out`` ; renvoie leurs chemins."""
    out.mkdir(parents=True, exist_ok=True)
    paths = {k: out / v for k, v in FILES.items()}
    paths["csv_j1"].write_bytes(write_csv(csv_a_j1()))
    paths["csv_j0"].write_bytes(write_csv([s[0] for s in csv_a_j0()]))
    paths["csv_incomplete"].write_bytes(write_csv(csv_a_incomplete()))
    write_xlsx(xlsx_b(), paths["xlsx"])
    paths["xml"].write_bytes(write_xml(xml_c(), XML_C_EXPORT))
    paths["xml_stale"].write_bytes(write_xml(xml_c(), XML_C_STALE_EXPORT))
    header = "# FICTIF — fiche d'un tarif reçu par email, recopiée pour import ASSISTÉ (aucune donnée réelle).\n"
    paths["assisted"].write_text(
        header + yaml.safe_dump(assisted_doc(), allow_unicode=True, sort_keys=False), encoding="utf-8"
    )
    exp = expected()
    paths["expected"].write_text(
        "# FICTIF — résultats attendus (générés), lus par tests/test_importers.py\n"
        + yaml.safe_dump(exp, allow_unicode=True, sort_keys=False),
        encoding="utf-8",
    )
    paths["readme"].write_text(readme(exp), encoding="utf-8")
    return paths


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out", type=Path, default=Path(__file__).resolve().parent)
    args = parser.parse_args(argv)
    for path in generate(args.out).values():
        print(path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
