"""Tests des importeurs fournisseurs : parseurs, dictionnaires, quarantaine, jeux FICTIFS, import assisté.

Toutes les données sont FICTIVES (GTIN 200…, SKU FICTIF-…, extensions « Fictive … »).
"""

from __future__ import annotations

import copy
import csv
import hashlib
import io
import sys
from datetime import date, datetime, timedelta
from decimal import Decimal
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

import openpyxl
import pytest
import yaml

from pokeshop.catalog import (
    CatalogIndex,
    CatalogProduct,
    ExtensionTable,
    MatchStatus,
    expected_language_for,
    fictitious_gtin13,
    is_fictitious_gtin,
    match_offer_to_product,
    normalize_gtin,
)
from pokeshop.importers import (
    QUARANTINE_LABELS_FR,
    SNAPSHOT_LEVEL_REASONS,
    AnomalyCode,
    AssistedImportError,
    AssistedPriceList,
    BytesConnector,
    FetchError,
    FileConnector,
    ImportBaseline,
    ImporterError,
    ImportPolicy,
    ImportResult,
    ImportStatus,
    MappingError,
    MappingStatus,
    QuarantineReason,
    RawSnapshot,
    Severity,
    SourceKind,
    SupplierConnector,
    SupplierMapping,
    UnitBasis,
    ValueParseError,
    approve_review,
    import_csv,
    import_xlsx,
    import_xml,
    list_mappings,
    load_mapping,
    load_price_list,
    merge_offer_book,
    next_baseline,
    review_price_list,
    run_import,
)
from pokeshop.importers import mapping as mapping_mod
from pokeshop.importers.values import (
    UnknownCurrencyError,
    cell_text,
    excel_number,
    is_blank,
    mapped_value,
    normalize_currency,
    parse_availability,
    parse_bool,
    parse_date,
    parse_datetime,
    parse_decimal,
    parse_money,
    parse_quantity,
    parse_rate,
    parse_unit_basis,
    parse_vat_basis,
)
from pokeshop.importers.xml_importer import flatten_element
from pokeshop.models import AvailabilityStatus, DecisionStatus, PromiseKind, Reason
from pokeshop.pricing import evaluate_offer
from pokeshop.rules import load_rules
from pokeshop.stock import StockRegistry, availability_promise

ROOT = Path(__file__).resolve().parents[1]
SAMPLES = ROOT / "data" / "samples"
MAPPINGS = ROOT / "data" / "supplier_mappings"
ZRH = ZoneInfo("Europe/Zurich")
NOW = datetime(2026, 10, 4, 8, 0, tzinfo=ZRH)
J0_TS = "2026-10-04T06:00:00+02:00"
EXPECTED = yaml.safe_load((SAMPLES / "FICTIF_attendus.yaml").read_text(encoding="utf-8"))
BP_SUPPLIERS = ("asmodee_fr", "matoo_miao", "tcg_distribution", "otakuworld", "cardcosmos")
TEMPLATE_PHRASE = "TEMPLATE — à compléter dès réception d'un fichier réel ; aucun format confirmé"


def g(n: int) -> str:
    return fictitious_gtin13(n)


@pytest.fixture(scope="module")
def fictif_table() -> ExtensionTable:
    return ExtensionTable.load([ROOT / "data" / "extensions_aliases.yaml", SAMPLES / "FICTIF_extensions_aliases.yaml"])


# ----------------------------------------------------------------- utilitaires


def deep_merge(base: dict[str, Any], extra: dict[str, Any]) -> dict[str, Any]:
    out = copy.deepcopy(base)
    for key, value in extra.items():
        if isinstance(value, dict) and isinstance(out.get(key), dict):
            out[key] = deep_merge(out[key], value)
        else:
            out[key] = value
    return out


MINI_MAPPING: dict[str, Any] = {
    "supplier_id": "fictif_test",
    "mapping_version": "fictif_test-v1",
    "status": "FICTIF",
    "fictif": True,
    "file_format": "csv",
    "csv": {"delimiter": ";"},
    "numbers": {"decimal_separator": "."},
    "columns": {
        "supplier_sku": ["sku"], "price": ["prix"], "currency": ["devise"], "price_basis": ["base"],
        "unit_basis": ["unite"], "units_per_pack": ["unites"], "carton_qty": ["carton"], "moq": ["moq"],
        "language": ["langue"], "gtin": ["ean"], "format": ["format"], "extension": ["extension"],
        "content": ["contenu"], "available_qty": ["stock"], "availability": ["statut"], "source_ts": ["maj"],
        "vat_rate": ["tva"], "allocation_qty": ["allocation"], "ship_from_country": ["pays"],
    },
    "defaults": {"sealed": True},
    "extension_tables": ["data/samples/FICTIF_extensions_aliases.yaml"],
}

MINI_ROW = {
    "sku": "FICTIF-T-1", "prix": "10.00", "devise": "EUR", "base": "HT", "unite": "unité", "unites": "", "carton": "",
    "moq": "", "langue": "FR", "ean": g(900), "format": "Display", "extension": "Fictive Alpha", "contenu": "36 boosters",
    "stock": "5", "statut": "en stock", "maj": J0_TS, "tva": "0", "allocation": "", "pays": "FR",
}


def mini_mapping(**overrides: Any) -> SupplierMapping:
    return SupplierMapping.model_validate(deep_merge(MINI_MAPPING, overrides))


def csv_bytes(rows: list[dict[str, str]], header: list[str] | None = None, delimiter: str = ";") -> bytes:
    cols = header or list(MINI_ROW)
    buf = io.StringIO()
    writer = csv.DictWriter(buf, fieldnames=cols, delimiter=delimiter, lineterminator="\n", extrasaction="ignore")
    writer.writeheader()
    for row in rows:
        writer.writerow({c: row.get(c, "") for c in cols})
    return buf.getvalue().encode("utf-8")


def row(**changes: str) -> dict[str, str]:
    return {**MINI_ROW, **changes}


def run_rows(rows: list[dict[str, str]], mapping: SupplierMapping | None = None, **kwargs: Any) -> ImportResult:
    return run_import(mapping or mini_mapping(), csv_bytes(rows), now=kwargs.pop("now", NOW), **kwargs)


def reasons_of(result: ImportResult, sku: str) -> set[str]:
    out: set[str] = set()
    for q in result.quarantined:
        if q.supplier_sku == sku:
            out.update(q.reasons)
    return out


def anomaly_codes(result: ImportResult, sku: str | None = None) -> set[str]:
    return {a.code for a in result.anomalies if sku is None or a.supplier_sku == sku}


# ---------------------------------------------------------------- parseurs


class TestValues:
    @pytest.mark.parametrize(
        ("raw", "sep", "expected"),
        [
            ("12.50", ".", "12.50"), ("12,50", ",", "12.50"), ("1'234.50", ".", "1234.50"), ("1 234,50", ",", "1234.50"),
            ("1 234,5", ",", "1234.5"), ("-5", ".", "-5"), ("+3.10", ".", "3.10"), ("0", ",", "0"),
            (Decimal("4.20"), ",", "4.20"), (7, ".", "7"), ("  ", ".", None), (None, ".", None),
        ],
    )
    def test_parse_decimal(self, raw: Any, sep: str, expected: str | None) -> None:
        result = parse_decimal(raw, decimal_separator=sep)
        assert result == (None if expected is None else Decimal(expected))

    @pytest.mark.parametrize(
        ("raw", "sep"),
        [
            ("12.50", ","), ("12,50", "."), ("1.234,5", "."), ("12,3,4", ","), ("abc", "."), ("1'23.0", "."),
            ("1 2345", "."), ("1e5", "."), (True, "."), (1.5, "."), (date(2026, 1, 1), "."), ("NaN", "."),
            ("1'234'5", "."), ("1 234,5 6", ","),
        ],
    )
    def test_parse_decimal_rejects_ambiguous(self, raw: Any, sep: str) -> None:
        with pytest.raises(ValueParseError):
            parse_decimal(raw, decimal_separator=sep)

    def test_parse_decimal_rejects_non_finite_and_bad_separator(self) -> None:
        with pytest.raises(ValueParseError):
            parse_decimal(Decimal("Infinity"))
        with pytest.raises(ValueParseError):
            parse_decimal("1;2", decimal_separator=";")

    @pytest.mark.parametrize(
        ("raw", "sep", "amount", "currency"),
        [
            ("12,50 €", ",", "12.50", "EUR"), ("€12,50", ",", "12.50", "EUR"), ("CHF 1'299.00", ".", "1299.00", "CHF"),
            ("12.30 Fr.", ".", "12.30", "CHF"), ("45 EUR", ".", "45", "EUR"), ("12.50 XYZ", ".", "12.50", "XYZ"),
            ("9.90", ".", "9.90", None), (Decimal("3.3"), ".", "3.3", None), (12, ".", "12", None), ("", ".", None, None),
        ],
    )
    def test_parse_money(self, raw: Any, sep: str, amount: str | None, currency: str | None) -> None:
        got_amount, got_currency = parse_money(raw, decimal_separator=sep)
        assert got_amount == (None if amount is None else Decimal(amount))
        assert got_currency == currency

    @pytest.mark.parametrize("raw", ["$45", "45 $", "12 bananes", "CHF 12 EUR", "abc", "--"])
    def test_parse_money_errors(self, raw: str) -> None:
        with pytest.raises(ValueParseError):
            parse_money(raw)

    def test_dollar_is_unknown_currency(self) -> None:
        with pytest.raises(UnknownCurrencyError):
            parse_money("$45")

    @pytest.mark.parametrize(
        ("raw", "expected"),
        [("EUR", "EUR"), ("eur", "EUR"), ("€", "EUR"), ("euros", "EUR"), ("CHF", "CHF"), ("Fr.", "CHF"),
         ("francs suisses", "CHF"), ("usd", "USD"), ("JPY", "JPY"), (None, None), ("", None)],
    )
    def test_normalize_currency(self, raw: str | None, expected: str | None) -> None:
        assert normalize_currency(raw) == expected

    @pytest.mark.parametrize("raw", ["$", "dollars", "Yen japonais", "12"])
    def test_normalize_currency_unknown(self, raw: str) -> None:
        with pytest.raises(UnknownCurrencyError):
            normalize_currency(raw)

    def test_currency_value_map(self) -> None:
        assert normalize_currency("Euro (zone)", {"euro zone": "EUR"}) == "EUR"

    @pytest.mark.parametrize(
        ("raw", "expected"),
        [("20%", "0.2"), ("20 %", "0.2"), ("20", "0.2"), ("0.2", "0.2"), ("8.1 %", "0.081"), ("0", "0"),
         (Decimal("0.081"), "0.081"), (Decimal("8.1"), "0.081"), (None, None), ("  ", None)],
    )
    def test_parse_rate(self, raw: Any, expected: str | None) -> None:
        assert parse_rate(raw) == (None if expected is None else Decimal(expected))

    @pytest.mark.parametrize("raw", ["100 %", "-5", "150", "1", "abc"])
    def test_parse_rate_errors(self, raw: str) -> None:
        with pytest.raises(ValueParseError):
            parse_rate(raw)

    @pytest.mark.parametrize(
        ("raw", "expected"),
        [("12", (12, False)), ("10+", (10, True)), (">50", (50, True)), (">= 7", (7, True)), ("≥3", (3, True)),
         (Decimal("6.0"), (6, False)), (4, (4, False)), ("", (None, False)), (None, (None, False))],
    )
    def test_parse_quantity(self, raw: Any, expected: tuple[int | None, bool]) -> None:
        assert parse_quantity(raw) == expected

    @pytest.mark.parametrize("raw", ["-3", -3, "2.5", "six", "beaucoup"])
    def test_parse_quantity_errors(self, raw: Any) -> None:
        with pytest.raises(ValueParseError):
            parse_quantity(raw)

    @pytest.mark.parametrize(
        ("raw", "expected"),
        [("oui", True), ("Scellé", True), ("x", True), (True, True), (1, True), ("non", False), ("ouvert", False),
         (0, False), (False, False), ("peut-être", None), ("", None), (None, None), (2, None)],
    )
    def test_parse_bool(self, raw: Any, expected: bool | None) -> None:
        assert parse_bool(raw) is expected

    @pytest.mark.parametrize(
        ("raw", "expected"),
        [("HT", False), ("H.T.", False), ("hors taxes", False), ("excl. TVA", False), ("TTC", True), ("T.T.C.", True),
         ("TVA comprise", True), ("inkl. MwSt", True), ("net", None), ("", None), (None, None), (True, None)],
    )
    def test_parse_vat_basis(self, raw: Any, expected: bool | None) -> None:
        assert parse_vat_basis(raw) is expected

    @pytest.mark.parametrize(
        ("raw", "expected"),
        [("unité", UnitBasis.UNIT), ("pièce", UnitBasis.UNIT), ("PCE", UnitBasis.UNIT), ("carton", UnitBasis.PACK),
         ("colis", UnitBasis.PACK), ("PACK", UnitBasis.PACK), ("UNIT", UnitBasis.UNIT), ("palette", None),
         ("display", None), ("boîte", None), ("", None), (None, None)],
    )
    def test_parse_unit_basis(self, raw: Any, expected: UnitBasis | None) -> None:
        assert parse_unit_basis(raw) is expected

    @pytest.mark.parametrize(
        ("raw", "expected"),
        [("en stock", "IN_STOCK"), ("Disponible", "IN_STOCK"), ("stock limité", "LOW_STOCK"), ("Rupture", "OUT_OF_STOCK"),
         ("épuisé", "OUT_OF_STOCK"), ("Précommande", "PREORDER"), ("sur allocation", "ALLOCATION"),
         ("fin de série", "DISCONTINUED"), ("PREORDER", "PREORDER"), ("??", "UNKNOWN"), (None, "UNKNOWN")],
    )
    def test_parse_availability(self, raw: str | None, expected: str) -> None:
        assert parse_availability(raw) is AvailabilityStatus(expected)

    def test_value_maps(self) -> None:
        assert parse_vat_basis("prix pro", {"Prix pro": "HT"}) is False
        assert parse_unit_basis("lot de 6", {"lot de 6": "PACK"}) is UnitBasis.PACK
        assert parse_availability("dispo J+2", {"dispo j+2": "IN_STOCK"}) is AvailabilityStatus.IN_STOCK
        assert parse_bool("ja", {"ja": "oui"}) is True
        assert mapped_value("ABC", {"x": "y"}) is None

    @pytest.mark.parametrize(
        ("raw", "expected"),
        [("2026-10-04", date(2026, 10, 4)), ("04.10.2026", date(2026, 10, 4)), ("04/10/2026", date(2026, 10, 4)),
         (datetime(2026, 1, 2, 3, 4), date(2026, 1, 2)), (date(2026, 5, 6), date(2026, 5, 6)), (None, None)],
    )
    def test_parse_date(self, raw: Any, expected: date | None) -> None:
        assert parse_date(raw) == expected

    def test_parse_date_error(self) -> None:
        with pytest.raises(ValueParseError):
            parse_date("bientôt")
        with pytest.raises(ValueParseError):
            parse_date("10/04/2026", ("%Y-%m-%d",))

    def test_parse_datetime(self) -> None:
        aware = parse_datetime("2026-10-04T06:00:00+02:00", ZRH)
        assert aware is not None and aware.utcoffset() == timedelta(hours=2)
        utc = parse_datetime("2026-10-04T04:00:00Z", ZRH)
        assert utc == aware
        naive = parse_datetime("04.10.2026 06:00", ZRH)
        assert naive == aware
        assert parse_datetime(datetime(2026, 10, 4, 6, 0), ZRH) == aware
        assert parse_datetime(date(2026, 10, 4), ZRH) == datetime(2026, 10, 4, tzinfo=ZRH)
        assert parse_datetime("2026-10-04", ZRH) == datetime(2026, 10, 4, tzinfo=ZRH)
        assert parse_datetime(None, ZRH) is None
        with pytest.raises(ValueParseError):
            parse_datetime("hier", ZRH)

    def test_excel_number_and_cell_text(self) -> None:
        assert excel_number(12.0) == 12 and isinstance(excel_number(12.0), int)
        assert excel_number(105.9) == Decimal("105.9")
        assert excel_number(0.1 + 0.2) == Decimal("0.30000000000000004")
        with pytest.raises(ValueParseError):
            excel_number(float("nan"))
        assert cell_text(Decimal("1.50")) == "1.50"
        assert cell_text(True) == "true" and cell_text(False) == "false"
        assert cell_text(date(2026, 1, 2)) == "2026-01-02"
        assert cell_text("  x ") == "x" and cell_text("   ") is None and cell_text(5) == "5"
        with pytest.raises(ValueParseError):
            cell_text(1.5)  # type: ignore[arg-type]
        assert is_blank(None) and is_blank(" ") and not is_blank(0)


# ------------------------------------------------------------ dictionnaires


class TestMappings:
    def test_all_shipped_mappings_load(self) -> None:
        mappings = list_mappings(MAPPINGS)
        assert set(mappings) == {"generic", *BP_SUPPLIERS, "fictif_grossiste_a", "fictif_grossiste_b", "fictif_grossiste_c"}
        for sid, m in mappings.items():
            assert m.supplier_id == sid
            assert m.content_sha256 and len(m.content_sha256) == 64
            assert m.source_path and m.source_path.endswith(f"{sid}.yaml")

    @pytest.mark.parametrize("sid", ("generic", *BP_SUPPLIERS))
    def test_templates_are_flagged_and_assume_nothing(self, sid: str) -> None:
        text = (MAPPINGS / f"{sid}.yaml").read_text(encoding="utf-8")
        assert text.startswith(f"# {TEMPLATE_PHRASE}.")
        assert "## " not in text  # pas de Markdown dans le YAML
        assert "Validation humaine requise" in text
        m = load_mapping(sid, mappings_dir=MAPPINGS)
        assert m.status is MappingStatus.TEMPLATE and m.is_template and m.simulation_only
        assert m.description == TEMPLATE_PHRASE
        assert not m.fictif and not m.allocation_is_firm
        assert m.defaults.model_dump() == {k: None for k in m.defaults.model_dump()}
        assert m.allowed_currencies == ("CHF", "EUR")

    @pytest.mark.parametrize("sid", BP_SUPPLIERS)
    def test_bp_supplier_templates_cite_bp(self, sid: str) -> None:
        text = (MAPPINGS / f"{sid}.yaml").read_text(encoding="utf-8")
        assert "BP §2" in text and "https://" in text and "consulté le 4.10.2026" in text
        assert "PAS un partenariat" in text

    def test_fictif_mappings(self) -> None:
        for sid in ("fictif_grossiste_a", "fictif_grossiste_b", "fictif_grossiste_c"):
            m = load_mapping(sid)
            assert m.fictif and m.status is MappingStatus.FICTIF and m.simulation_only
            assert m.extension_table_paths()[0].exists()

    def test_load_by_path_env_and_errors(self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
        assert load_mapping(MAPPINGS / "generic.yaml").supplier_id == "generic"
        monkeypatch.setenv(mapping_mod.MAPPINGS_ENV_VAR, str(tmp_path))
        with pytest.raises(MappingError, match="introuvable"):
            load_mapping("generic")
        (tmp_path / "generic.yaml").write_text((MAPPINGS / "generic.yaml").read_text(encoding="utf-8"), encoding="utf-8")
        assert load_mapping("generic").supplier_id == "generic"
        with pytest.raises(MappingError, match="introuvable"):
            load_mapping("../etc/passwd")

    def test_file_name_must_match_supplier_id(self, tmp_path: Path) -> None:
        path = tmp_path / "autre_nom.yaml"
        path.write_text((MAPPINGS / "generic.yaml").read_text(encoding="utf-8"), encoding="utf-8")
        with pytest.raises(MappingError, match="nom de fichier"):
            load_mapping(path)

    def test_bad_yaml_and_non_mapping(self, tmp_path: Path) -> None:
        bad = tmp_path / "fictif_x.yaml"
        bad.write_text("columns: [\n", encoding="utf-8")
        with pytest.raises(MappingError, match="lecture impossible"):
            load_mapping(bad)
        bad.write_text("- 1\n- 2\n", encoding="utf-8")
        with pytest.raises(MappingError, match="objet"):
            load_mapping(bad)

    def test_missing_extension_table_detected(self, tmp_path: Path) -> None:
        data = deep_merge(MINI_MAPPING, {"extension_tables": ["data/samples/absent.yaml"]})
        path = tmp_path / "fictif_test.yaml"
        path.write_text(yaml.safe_dump(data, allow_unicode=True), encoding="utf-8")
        with pytest.raises(MappingError, match="table d'extensions introuvable"):
            load_mapping(path)

    @pytest.mark.parametrize(
        "overrides",
        [
            {"columns": {"prix_b2b": ["x"]}},
            {"columns": {"price": []}},
            {"supplier_id": "fictif-Bad"},
            {"supplier_id": "reel_test"},
            {"fictif": False},
            {"status": "VALIDE"},
            {"allocation_is_firm": True},
            {"required_fields": ["supplier_sku"]},
            {"required_fields": ["supplier_sku", "price", "moq_inconnu"]},
            {"required_fields": ["supplier_sku", "price", "incoterm"]},
            {"allowed_currencies": ["eur"]},
            {"defaults": {"currency": "USD"}},
            {"defaults": {"language": "UNKNOWN"}},
            {"defaults": {"vat_rate": "1.2"}},
            {"defaults": {"ship_from_country": "France"}},
            {"value_maps": {"price": {"a": "b"}}},
            {"csv": {"delimiter": ";;"}},
            {"csv": {"encoding": "klingon-8"}},
            {"xlsx": {"source_ts_cell": "B0"}},
            {"dates": {"timezone": "Mars/Olympus"}},
            {"numbers": {"decimal_separator": ",", "thousands_separators": [","]}},
            {"tiers": [{"min_qty": 6, "min_qty_column": "q", "price_column": "p"}]},
            {"tiers": [{"min_qty": 6}]},
            {"unexpected": True},
        ],
    )
    def test_invalid_mappings(self, overrides: dict[str, Any]) -> None:
        with pytest.raises(ValueError):
            mini_mapping(**overrides)

    def test_real_mapping_rules(self) -> None:
        real = deep_merge(MINI_MAPPING, {"supplier_id": "reel_test", "fictif": False, "status": "TEMPLATE"})
        real.pop("extension_tables")
        SupplierMapping.model_validate(real)
        with pytest.raises(ValueError, match="extension_tables"):
            SupplierMapping.model_validate({**real, "extension_tables": ["data/samples/FICTIF_extensions_aliases.yaml"]})
        valide = {**real, "status": "VALIDE", "validated_by": "Personne X", "validated_at": "2026-10-04"}
        with pytest.raises(ValueError, match="VALIDE"):
            SupplierMapping.model_validate(valide)
        with pytest.raises(ValueError, match="sha256"):
            SupplierMapping.model_validate({**valide, "validated_on_sample_sha256": "abc"})
        ok = SupplierMapping.model_validate({**valide, "validated_on_sample_sha256": "a" * 64})
        assert not ok.simulation_only
        firm = {**real, "allocation_is_firm": True, "allocation_evidence": "email du 4.10.2026 (FICTIF)"}
        assert SupplierMapping.model_validate(firm).allocation_is_firm

    def test_columns_for(self) -> None:
        m = mini_mapping()
        assert m.columns_for("price") == ("prix",)
        assert m.columns_for("incoterm") == ()
        with pytest.raises(MappingError):
            m.columns_for("cost")


# ----------------------------------------------------------------- captures


class TestSnapshot:
    def test_checksum_computed_and_verified(self) -> None:
        snap = RawSnapshot.capture(
            supplier_id="fictif_test", content=b"abc", source_kind=SourceKind.CSV, source_uri="memory://x", fetched_at=NOW
        )
        assert snap.checksum_sha256 == hashlib.sha256(b"abc").hexdigest()
        assert snap.byte_size == 3
        assert snap.meta().snapshot_id == f"fictif_test:{snap.checksum_sha256[:16]}"
        assert "content" not in repr(snap)
        with pytest.raises(ValueError, match="altérée"):
            RawSnapshot(supplier_id="x", source_kind=SourceKind.CSV, source_uri="m", content=b"abc", fetched_at=NOW,
                        checksum_sha256="0" * 64)
        same = RawSnapshot(supplier_id="x", source_kind=SourceKind.CSV, source_uri="m", content=b"abc", fetched_at=NOW,
                           checksum_sha256=hashlib.sha256(b"abc").hexdigest())
        assert same.checksum_sha256 == snap.checksum_sha256

    def test_naive_datetimes_refused(self) -> None:
        with pytest.raises(ValueError):
            RawSnapshot.capture(supplier_id="x", content=b"", source_kind=SourceKind.CSV, source_uri="m",
                                fetched_at=datetime(2026, 10, 4))
        with pytest.raises(ValueError):
            RawSnapshot.capture(supplier_id="x", content=b"", source_kind=SourceKind.CSV, source_uri="m", fetched_at=NOW,
                                source_ts=datetime(2026, 10, 4))

    def test_snapshot_meta_in_result(self) -> None:
        path = SAMPLES / "FICTIF_offres_grossiste_a_J-1.csv"
        result = run_import("fictif_grossiste_a", path, now=NOW)
        assert result.snapshot.checksum_sha256 == hashlib.sha256(path.read_bytes()).hexdigest()
        assert result.snapshot.source_kind is SourceKind.CSV
        assert result.snapshot.fictif and result.fictif and result.dry_run
        assert all(o.raw_ref.startswith(f"fictif_grossiste_a:{result.snapshot.checksum_sha256[:12]}#L") for o in result.offers)


# ------------------------------------------------------- jeux d'essai FICTIFS


def _import_sample(name: str) -> ImportResult:
    spec = EXPECTED[name]
    baseline = None
    if "baseline" in spec:
        ref = EXPECTED[spec["baseline"]]
        prev = run_import(ref["mapping"], SAMPLES / spec["baseline"], now=datetime.fromisoformat(ref["now"]))
        baseline = ImportBaseline.from_result(prev)
    return run_import(spec["mapping"], SAMPLES / name, now=datetime.fromisoformat(spec["now"]), baseline=baseline)


SAMPLE_FILES = [k for k, v in EXPECTED.items() if isinstance(v, dict) and "mapping" in v]


class TestSamples:
    @pytest.mark.parametrize("name", SAMPLE_FILES)
    def test_status_and_counts(self, name: str) -> None:
        spec = EXPECTED[name]
        result = _import_sample(name)
        assert result.status.value == spec["status"]
        assert result.rows_read == spec["rows_read"]
        if "offers" in spec:
            assert result.accepted_count == spec["offers"]
        dropped = len(result.anomalies_with(AnomalyCode.DUPLICATE_IDENTICAL_ROW))
        assert result.accepted_count + result.quarantined_count + dropped == result.rows_read
        import_codes = {a.code for a in result.anomalies if a.line_no is None}
        assert set(spec.get("import_anomalies", [])) <= import_codes

    @pytest.mark.parametrize("name", [n for n in SAMPLE_FILES if "lines" in EXPECTED[n]])
    def test_each_line(self, name: str) -> None:
        result = _import_sample(name)
        quarantined = {q.line_no: q for q in result.quarantined}
        anomalies: dict[int | None, set[str]] = {}
        for a in result.anomalies:
            anomalies.setdefault(a.line_no, set()).add(a.code)
        offers_lines = {int(o.raw_ref.rsplit("#L", 1)[1]) for o in result.offers}
        for entry in EXPECTED[name]["lines"]:
            line = entry["line"]
            got = sorted(quarantined[line].reasons) if line in quarantined else []
            assert got == sorted(entry["quarantine"]), (line, entry["sku"], entry["note"])
            assert set(entry["anomalies"]) <= anomalies.get(line, set()), (line, entry["sku"])
            dropped = "DUPLICATE_IDENTICAL_ROW" in entry["anomalies"]
            assert (line in offers_lines) == (not entry["quarantine"] and not dropped), (line, entry["sku"])

    def test_at_least_20_rows_per_format(self) -> None:
        for name in ("FICTIF_offres_grossiste_a.csv", "FICTIF_tarif_grossiste_b.xlsx", "FICTIF_flux_grossiste_c.xml"):
            assert EXPECTED[name]["rows_read"] >= 20

    def test_all_quarantine_rules_are_covered(self) -> None:
        covered: set[str] = set()
        for name in SAMPLE_FILES:
            result = _import_sample(name)
            covered |= {r for q in result.quarantined for r in q.reasons}
            covered |= {a.code for a in result.anomalies if a.line_no is None}
        required = {
            "UNKNOWN_CURRENCY", "UNKNOWN_VAT_BASIS", "NON_POSITIVE_PRICE", "PRICE_ANOMALY", "DUPLICATE_SKU",
            "INCOMPLETE_IMPORT", "INVALID_GTIN", "UNKNOWN_LANGUAGE", "STALE_SNAPSHOT", "STALE_SOURCE",
            "UNKNOWN_UNIT_BASIS", "UNIT_CONFLICT", "INVALID_TIER", "CURRENCY_CHANGED", "VAT_BASIS_CHANGED",
            "CURRENCY_CONFLICT", "UNPARSEABLE_PRICE", "MISSING_SKU", "MISSING_PRICE", "FICTIF_FLAG_MISMATCH",
            "INVALID_QUANTITY", "UNPARSEABLE_VALUE",
        }
        assert required <= covered, required - covered

    def test_readme_documents_every_line_and_code(self) -> None:
        readme = (SAMPLES / "README.md").read_text(encoding="utf-8")
        assert "FICTIF" in readme and "## Validation humaine requise" in readme
        for name in SAMPLE_FILES:
            assert name in readme
            for entry in EXPECTED[name].get("lines", []):
                for code in entry["quarantine"]:
                    assert code in readme

    def test_samples_are_fictitious(self) -> None:
        for name in ("FICTIF_offres_grossiste_a.csv", "FICTIF_offres_grossiste_a_J-1.csv", "FICTIF_flux_grossiste_c.xml"):
            result = _import_sample(name) if name in EXPECTED and "baseline" in EXPECTED[name] else run_import(
                EXPECTED[name]["mapping"], SAMPLES / name, now=datetime.fromisoformat(EXPECTED[name]["now"])
            )
            for offer in result.offers:
                assert offer.supplier_sku.startswith("FICTIF-")
                assert offer.supplier_id.startswith("fictif_")
                assert offer.gtin is None or is_fictitious_gtin(offer.gtin) or offer.gtin.startswith("1200")
                assert offer.extension in (None, "SANS_EXTENSION") or offer.extension.startswith("FICTIF_")
        xlsx = _import_sample("FICTIF_tarif_grossiste_b.xlsx")
        assert all(o.supplier_sku.startswith("FICTIF-B-") for o in xlsx.offers)
        for path in SAMPLES.iterdir():
            if path.is_dir():
                continue
            assert path.name.startswith("FICTIF_") or path.name in ("README.md", "generate_samples.py")

    def test_generator_is_deterministic_and_matches_committed_files(self, tmp_path: Path) -> None:
        sys.path.insert(0, str(SAMPLES))
        previous_flag = sys.dont_write_bytecode
        sys.dont_write_bytecode = True  # ne rien écrire dans data/samples/
        try:
            import generate_samples
        finally:
            sys.path.remove(str(SAMPLES))
            sys.dont_write_bytecode = previous_flag
        paths = generate_samples.generate(tmp_path)
        for key, path in paths.items():
            committed = SAMPLES / path.name
            if key == "xlsx":
                assert _xlsx_values(path) == _xlsx_values(committed)
            else:
                assert path.read_bytes() == committed.read_bytes(), path.name
        assert generate_samples.main(["--out", str(tmp_path / "again")]) == 0

    def test_result_serializes_to_json_without_floats(self) -> None:
        import json

        result = _import_sample("FICTIF_offres_grossiste_a.csv")
        data = json.loads(result.model_dump_json())
        assert data["status"] == "PARTIAL" and data["snapshot"]["checksum_sha256"] == result.snapshot.checksum_sha256
        first = data["offers"][0]
        assert first["price"] == "95.00" and isinstance(first["price"], str)
        assert ImportResult.model_validate_json(result.model_dump_json()) == result

    def test_identical_reimport_is_deterministic(self) -> None:
        a = _import_sample("FICTIF_flux_grossiste_c.xml")
        b = _import_sample("FICTIF_flux_grossiste_c.xml")
        assert a == b

    def test_report_markdown(self) -> None:
        result = _import_sample("FICTIF_offres_grossiste_a.csv")
        report = result.report_markdown()
        assert report.startswith("# Import — fictif_grossiste_a — 2026-10-04 08:00 — mode SIMULATION")
        assert "**FICTIF**" in report and "escalade E1" in report
        assert "## Quarantaine" in report and "## Anomalies" in report and "## Validation humaine requise" in report
        assert "Écarts avec l'import précédent" in report and "FICTIF-A-099" in report
        assert "Fraîcheur : 2.0 h → OK" in report
        stale = _import_sample("FICTIF_flux_grossiste_c_perime.xml").report_markdown()
        assert "périmé : achats et promesses bloqués" in stale


def _xlsx_values(path: Path) -> list[list[Any]]:
    wb = openpyxl.load_workbook(path, data_only=True)
    try:
        return [[c for c in r] for ws in wb.worksheets for r in ws.iter_rows(values_only=True)]
    finally:
        wb.close()


# ------------------------------------------------------------- règles de ligne


class TestRowRules:
    def test_clean_row(self) -> None:
        result = run_rows([row()])
        assert result.status is ImportStatus.ACCEPTED and result.escalation is None
        offer = result.offers[0]
        assert offer.price == Decimal("10.00") and offer.currency == "EUR" and offer.price_includes_vat is False
        assert offer.units_per_pack == 1 and offer.language == "FR" and offer.extension == "FICTIF_ALPHA"
        assert offer.format == "DISPLAY" and offer.content == "36 BOOSTERS" and offer.sealed is True
        assert offer.vat_rate == Decimal(0) and offer.available_qty == 5
        assert offer.source_ts == datetime.fromisoformat(J0_TS) and offer.ship_from_country == "FR"

    @pytest.mark.parametrize(
        ("changes", "reason"),
        [
            ({"prix": ""}, "MISSING_PRICE"),
            ({"prix": "0.00"}, "NON_POSITIVE_PRICE"),
            ({"prix": "-1"}, "NON_POSITIVE_PRICE"),
            ({"prix": "10,00"}, "UNPARSEABLE_PRICE"),
            ({"prix": "$10"}, "UNKNOWN_CURRENCY"),
            ({"devise": "GBP"}, "UNKNOWN_CURRENCY"),
            ({"devise": "dollars"}, "UNKNOWN_CURRENCY"),
            ({"devise": ""}, "UNKNOWN_CURRENCY"),
            ({"prix": "10 CHF"}, "CURRENCY_CONFLICT"),
            ({"base": "net"}, "UNKNOWN_VAT_BASIS"),
            ({"base": ""}, "UNKNOWN_VAT_BASIS"),
            ({"unite": ""}, "UNKNOWN_UNIT_BASIS"),
            ({"unite": "palette"}, "UNKNOWN_UNIT_BASIS"),
            ({"unite": "carton"}, "UNKNOWN_UNIT_BASIS"),
            ({"unite": "unité", "unites": "6"}, "UNIT_CONFLICT"),
            ({"unites": "x"}, "UNPARSEABLE_VALUE"),
            ({"carton": "0"}, "INVALID_QUANTITY"),
            ({"moq": "0"}, "INVALID_QUANTITY"),
            ({"moq": "beaucoup"}, "UNPARSEABLE_VALUE"),
            ({"ean": g(900)[:-1] + str((int(g(900)[-1]) + 1) % 10)}, "INVALID_GTIN"),
            ({"ean": "12345"}, "INVALID_GTIN"),
            ({"langue": "VO"}, "UNKNOWN_LANGUAGE"),
            ({"langue": ""}, "UNKNOWN_LANGUAGE"),
            ({"langue": "n/a"}, "UNKNOWN_LANGUAGE"),
            ({"tva": "beaucoup"}, "UNPARSEABLE_VALUE"),
            ({"maj": "2026-10-02T06:00:00+02:00"}, "STALE_SOURCE"),
            ({"maj": "2026-10-05T06:00:00+02:00"}, "STALE_SOURCE"),
            ({"maj": "hier"}, "UNPARSEABLE_VALUE"),
            ({"sku": ""}, "MISSING_SKU"),
        ],
    )
    def test_quarantine_reasons(self, changes: dict[str, str], reason: str) -> None:
        result = run_rows([row(**changes), row(sku="FICTIF-T-OK", ean=g(901))])
        assert result.status is ImportStatus.PARTIAL
        q = result.quarantined[0]
        assert reason in q.reasons, (q.reasons, q.details)
        assert q.line_no == 2 and q.raw["prix"] == (changes.get("prix", "10.00") or None)
        assert [o.supplier_sku for o in result.offers] == ["FICTIF-T-OK"]
        assert QuarantineReason(reason) in QUARANTINE_LABELS_FR

    @pytest.mark.parametrize(
        ("changes", "code", "check"),
        [
            ({"ean": ""}, "MISSING_GTIN", lambda o: o.gtin is None),
            ({"format": "Blister", "contenu": "1 booster"}, "UNKNOWN_FORMAT", lambda o: o.format is None),
            ({"extension": "Inconnue"}, "UNKNOWN_EXTENSION", lambda o: o.extension is None),
            ({"extension": "Fictive Alpha / Fictive Bêta"}, "AMBIGUOUS_EXTENSION", lambda o: o.extension is None),
            ({"contenu": ""}, "MISSING_CONTENT", lambda o: o.content is None),
            ({"tva": ""}, "MISSING_VAT_RATE", lambda o: o.vat_rate is None),
            ({"statut": "bof"}, "UNKNOWN_AVAILABILITY", lambda o: o.availability_status is AvailabilityStatus.UNKNOWN),
            ({"stock": "-2"}, "UNPARSEABLE_OPTIONAL", lambda o: o.available_qty is None),
            ({"stock": "50+"}, "APPROXIMATE_STOCK", lambda o: o.available_qty == 50),
            ({"allocation": "12"}, "ALLOCATION_NOT_FIRM", lambda o: o.allocation_qty is None),
            ({"pays": "France"}, "INVALID_COUNTRY", lambda o: o.ship_from_country is None),
            ({"carton": "6", "moq": "4"}, "MOQ_NOT_CARTON_MULTIPLE", lambda o: o.moq == 4),
            ({"ean": "1" + g(900)[:-1] + "9"}, "CASE_LEVEL_GTIN", lambda o: True),
        ],
    )
    def test_warnings_keep_the_offer(self, changes: dict[str, str], code: str, check: Any) -> None:
        if code == "CASE_LEVEL_GTIN":
            from pokeshop.catalog import gtin_check_digit

            payload = "1" + g(900)[:-1]
            changes = {"ean": payload + str(gtin_check_digit(payload))}
        result = run_rows([row(**changes)])
        assert result.status is ImportStatus.ACCEPTED, result.quarantined
        assert code in anomaly_codes(result, "FICTIF-T-1")
        assert check(result.offers[0])

    def test_designation_fallbacks(self) -> None:
        mapping = mini_mapping(columns={"designation": ["designation"]})
        content = csv_bytes(
            [row(format="", extension="", contenu="", designation="Display 36 boosters Extension Fictive Bêta")],
            header=[*MINI_ROW, "designation"],
        )
        offer = run_import(mapping, content, now=NOW).offers[0]
        assert (offer.format, offer.extension, offer.content) == ("DISPLAY", "FICTIF_BETA", "36 BOOSTERS")

    def test_accessory_language_and_extension(self) -> None:
        result = run_rows([row(langue="", extension="", format="Protège-cartes", contenu="65 protège-cartes")])
        offer = result.offers[0]
        assert offer.language == "NA" and offer.extension == "SANS_EXTENSION" and offer.format == "SLEEVES"
        assert expected_language_for(offer.format) == "NA"

    def test_pack_price_units(self) -> None:
        result = run_rows([row(unite="carton", unites="6", prix="57.00")])
        offer = result.offers[0]
        assert offer.units_per_pack == 6 and offer.carton_qty == 6
        from_carton = run_rows([row(unite="carton", unites="", carton="12", prix="114.00")]).offers[0]
        assert from_carton.units_per_pack == 12

    def test_quantities_in_cartons(self) -> None:
        mapping = mini_mapping(quantities={"stock": "PACK", "moq": "PACK", "tiers": "PACK", "allocation": "PACK"},
                               allocation_is_firm=True, allocation_evidence="FICTIF : accord écrit simulé",
                               tiers=[{"min_qty_column": "q1", "price_column": "p1"}])
        header = [*MINI_ROW, "q1", "p1"]
        content = csv_bytes([row(carton="6", stock="3", moq="2", allocation="4", q1="5", p1="9.50")], header=header)
        offer = run_import(mapping, content, now=NOW).offers[0]
        assert (offer.available_qty, offer.moq, offer.allocation_qty) == (18, 12, 24)
        assert offer.tier_discounts[0].min_qty == 30
        unknown = run_import(mapping, csv_bytes([row(stock="3", moq="2", q1="", p1="")], header=header), now=NOW)
        assert "UNKNOWN_UNIT_BASIS" in reasons_of(unknown, "FICTIF-T-1")

    def test_stock_in_cartons_without_size_is_ignored(self) -> None:
        mapping = mini_mapping(quantities={"stock": "PACK"})
        result = run_import(mapping, csv_bytes([row(stock="3")]), now=NOW)
        assert result.offers[0].available_qty is None
        assert "UNPARSEABLE_OPTIONAL" in anomaly_codes(result)

    def test_firm_allocation_requires_parseable_qty(self) -> None:
        mapping = mini_mapping(allocation_is_firm=True, allocation_evidence="FICTIF")
        assert run_import(mapping, csv_bytes([row(allocation="8")]), now=NOW).offers[0].allocation_qty == 8
        bad = run_import(mapping, csv_bytes([row(allocation="huit")]), now=NOW)
        assert "UNPARSEABLE_VALUE" in reasons_of(bad, "FICTIF-T-1")

    @pytest.mark.parametrize(
        ("cells", "ok"),
        [
            ({"q1": "6", "p1": "9.50", "q2": "12", "p2": "9.00"}, True),
            ({"q1": "6", "p1": "9.50", "q2": "", "p2": ""}, True),
            ({"q1": "6", "p1": "11.00", "q2": "", "p2": ""}, False),
            ({"q1": "6", "p1": "9.50", "q2": "12", "p2": "9.80"}, False),
            ({"q1": "6", "p1": "9.50", "q2": "6", "p2": "9.00"}, False),
            ({"q1": "", "p1": "9.50", "q2": "", "p2": ""}, False),
            ({"q1": "6", "p1": "", "q2": "", "p2": ""}, False),
            ({"q1": "6", "p1": "0", "q2": "", "p2": ""}, False),
            ({"q1": "x", "p1": "9", "q2": "", "p2": ""}, False),
            ({"q1": "6", "p1": "9.50 CHF", "q2": "", "p2": ""}, False),
            ({"q1": "0", "p1": "9.50", "q2": "", "p2": ""}, False),
        ],
    )
    def test_price_tiers(self, cells: dict[str, str], ok: bool) -> None:
        mapping = mini_mapping(tiers=[{"min_qty_column": "q1", "price_column": "p1"}, {"min_qty_column": "q2", "price_column": "p2"}])
        result = run_import(mapping, csv_bytes([row(**cells)], header=[*MINI_ROW, "q1", "p1", "q2", "p2"]), now=NOW)
        assert (result.status is ImportStatus.ACCEPTED) is ok, result.quarantined
        if not ok:
            assert reasons_of(result, "FICTIF-T-1") == {"INVALID_TIER"}

    @pytest.mark.parametrize(("pct2", "ok"), [("10 %", True), ("5", True), ("2 %", False), ("120 %", False)])
    def test_percent_tiers(self, pct2: str, ok: bool) -> None:
        mapping = mini_mapping(tiers=[{"min_qty": 6, "discount_pct_column": "r1"}, {"min_qty": 12, "discount_pct_column": "r2"}])
        result = run_import(mapping, csv_bytes([row(r1="5 %", r2=pct2)], header=[*MINI_ROW, "r1", "r2"]), now=NOW)
        assert (result.status is ImportStatus.ACCEPTED) is ok
        if ok:
            assert [t.min_qty for t in result.offers[0].tier_discounts] == [6, 12]

    def test_tier_column_absent_is_warning(self) -> None:
        mapping = mini_mapping(tiers=[{"min_qty": 6, "price_column": "absente"}])
        result = run_rows([row()], mapping)
        assert result.status is ImportStatus.ACCEPTED
        assert any(a.code == "MAPPING_MISMATCH" and a.severity is Severity.WARNING for a in result.anomalies)

    def test_value_maps_in_import(self) -> None:
        mapping = mini_mapping(value_maps={
            "price_basis": {"prix pro": "HT"}, "language": {"VF Québec": "FR"}, "format": {"Boîte 36": "DISPLAY"},
            "extension": {"Alpha FX": "Fictive Alpha"}, "unit_basis": {"u.": "UNIT"},
        })
        result = run_rows([row(base="prix pro", langue="VF Québec", format="Boîte 36", extension="Alpha FX", unite="u.")], mapping)
        offer = result.offers[0]
        assert offer.price_includes_vat is False and offer.language == "FR"
        assert offer.format == "DISPLAY" and offer.extension == "FICTIF_ALPHA"

    def test_defaults_apply_only_when_column_empty(self) -> None:
        mapping = mini_mapping(defaults={"currency": "EUR", "price_basis": "HT", "unit_basis": "UNIT", "language": "FR",
                                         "vat_rate": "0", "ship_from_country": "FR", "incoterm": "dap", "stock_pool_id": "P1"})
        result = run_rows([row(devise="", base="", unite="", langue="", tva="", pays="")], mapping)
        offer = result.offers[0]
        assert (offer.currency, offer.price_includes_vat, offer.language, offer.vat_rate) == ("EUR", False, "FR", Decimal(0))
        assert offer.incoterm == "DAP" and offer.stock_pool_id == "P1" and offer.ship_from_country == "FR"
        conflict = run_rows([row(devise="", prix="10 CHF")], mapping)
        assert "CURRENCY_CONFLICT" in reasons_of(conflict, "FICTIF-T-1")

    def test_currency_not_configured_is_critical(self) -> None:
        mapping = mini_mapping(columns={"currency": ["absente"]})
        result = run_rows([row(prix="10 €")], mapping)
        assert result.offers[0].currency == "EUR"
        assert any(a.code == "CURRENCY_NOT_CONFIGURED" and a.severity is Severity.CRITICAL for a in result.anomalies)
        assert result.escalation == "E1"

    def test_fictif_flag_rules(self) -> None:
        mapping = mini_mapping(columns={"fictif": ["fictif"]})
        header = [*MINI_ROW, "fictif"]
        assert run_import(mapping, csv_bytes([row(fictif="true")], header=header), now=NOW).status is ImportStatus.ACCEPTED
        bad = run_import(mapping, csv_bytes([row(fictif="false")], header=header), now=NOW)
        assert "FICTIF_FLAG_MISMATCH" in reasons_of(bad, "FICTIF-T-1")
        real = deep_merge(MINI_MAPPING, {"supplier_id": "reel_test", "fictif": False, "status": "TEMPLATE",
                                         "columns": {"fictif": ["fictif"]}})
        real.pop("extension_tables")
        table = ExtensionTable.load([ROOT / "data" / "extensions_aliases.yaml", SAMPLES / "FICTIF_extensions_aliases.yaml"])
        result = run_import(SupplierMapping.model_validate(real), csv_bytes([row(fictif="oui")], header=header), now=NOW,
                            extensions=table)
        assert "FICTIF_FLAG_MISMATCH" in reasons_of(result, "FICTIF-T-1")

    def test_duplicates(self) -> None:
        identical = run_rows([row(), row()])
        assert identical.status is ImportStatus.ACCEPTED and identical.accepted_count == 1
        assert identical.anomalies_with(AnomalyCode.DUPLICATE_IDENTICAL_ROW)[0].line_no == 3
        conflicting = run_rows([row(), row(prix="11.00"), row(sku="FICTIF-T-2", ean=g(902))])
        assert len(conflicting.quarantined_with(QuarantineReason.DUPLICATE_SKU)) == 2
        assert [o.supplier_sku for o in conflicting.offers] == ["FICTIF-T-2"]

    def test_malformed_row(self) -> None:
        content = csv_bytes([row()]) + b"FICTIF-T-9;12.00\n"
        result = run_import(mini_mapping(), content, now=NOW)
        q = result.quarantined_with(QuarantineReason.MALFORMED_ROW)
        assert len(q) == 1 and q[0].line_no == 3

    def test_blank_rows_ignored(self) -> None:
        content = csv_bytes([row()]) + b";;;;\n\n" + csv_bytes([row(sku="FICTIF-T-2", ean=g(902))]).split(b"\n", 1)[1]
        result = run_import(mini_mapping(), content, now=NOW)
        assert result.accepted_count == 2 and result.rows_read == 2
        assert result.anomalies_with(AnomalyCode.BLANK_ROWS)

    def test_unmapped_columns_reported(self) -> None:
        content = csv_bytes([row(colonne_secrete="x")], header=[*MINI_ROW, "colonne_secrete"])
        result = run_import(mini_mapping(), content, now=NOW)
        assert "colonne_secrete" in result.anomalies_with(AnomalyCode.UNMAPPED_COLUMNS)[0].detail


# -------------------------------------------------------- règles de niveau import


class TestImportRules:
    def baseline(self, rows: list[dict[str, str]] | None = None) -> ImportBaseline:
        rows = rows or [row(sku=f"FICTIF-T-{i}", ean=g(910 + i)) for i in range(10)]
        return ImportBaseline.from_result(run_rows(rows))

    def test_price_anomaly_both_directions_per_unit(self) -> None:
        base = self.baseline()
        rows = [row(sku=f"FICTIF-T-{i}", ean=g(910 + i)) for i in range(10)]
        rows[0]["prix"] = "100.00"
        rows[1]["prix"] = "1.00"
        rows[2]["prix"] = "16.00"
        rows[3] = row(sku="FICTIF-T-3", ean=g(913), unite="carton", unites="6", prix="60.00")  # même prix unitaire
        rows[4]["prix"] = "99.99"  # ×9,999 : sous le facteur 10 => signal seulement
        result = run_rows(rows, baseline=base)
        assert {q.supplier_sku for q in result.quarantined_with(QuarantineReason.PRICE_ANOMALY)} == {"FICTIF-T-0", "FICTIF-T-1"}
        assert "FICTIF-T-4" in result.offers_by_sku and "PRICE_CHANGE_LARGE" in anomaly_codes(result, "FICTIF-T-4")
        assert "PRICE_CHANGE_LARGE" in anomaly_codes(result, "FICTIF-T-2")
        assert "FICTIF-T-3" in result.offers_by_sku
        assert result.escalation == "E1"

    def test_factor_from_policy(self) -> None:
        base = self.baseline()
        rows = [row(sku=f"FICTIF-T-{i}", ean=g(910 + i)) for i in range(10)]
        rows[0]["prix"] = "50.00"
        strict = run_rows(rows, baseline=base, policy=ImportPolicy(price_anomaly_factor=Decimal(4)))
        assert "PRICE_ANOMALY" in reasons_of(strict, "FICTIF-T-0")
        lax = run_rows(rows, baseline=base)
        assert "PRICE_ANOMALY" not in reasons_of(lax, "FICTIF-T-0")

    def test_currency_and_vat_basis_change(self) -> None:
        base = self.baseline()
        rows = [row(sku=f"FICTIF-T-{i}", ean=g(910 + i)) for i in range(10)]
        rows[0]["devise"] = "CHF"
        rows[1]["base"] = "TTC"
        result = run_rows(rows, baseline=base)
        assert reasons_of(result, "FICTIF-T-0") == {"CURRENCY_CHANGED"}
        assert reasons_of(result, "FICTIF-T-1") == {"VAT_BASIS_CHANGED"}

    def test_incomplete_import_threshold(self) -> None:
        base = self.baseline()
        eight = [row(sku=f"FICTIF-T-{i}", ean=g(910 + i)) for i in range(8)]
        assert run_rows(eight, baseline=base).status is ImportStatus.ACCEPTED
        seven = eight[:7]
        result = run_rows(seven, baseline=base)
        assert result.status is ImportStatus.QUARANTINED and not result.offers
        assert all("INCOMPLETE_IMPORT" in q.reasons for q in result.quarantined)
        assert result.retain_previous_offers
        assert next_baseline(base, result) == base
        custom = run_rows(seven, baseline=base, policy=ImportPolicy(min_rows_ratio=Decimal("0.5")))
        assert custom.status is ImportStatus.ACCEPTED

    def test_empty_import(self) -> None:
        result = run_import(mini_mapping(), csv_bytes([]), now=NOW)
        assert result.status is ImportStatus.QUARANTINED and result.rows_read == 0
        assert result.anomalies_with(QuarantineReason.EMPTY_IMPORT)

    def test_stale_snapshot_from_connector_timestamp(self) -> None:
        result = run_import(mini_mapping(columns={"source_ts": ["absente"]}), csv_bytes([row()]), now=NOW,
                            source_ts=NOW - timedelta(hours=25))
        assert result.status is ImportStatus.QUARANTINED
        assert result.anomalies_with(QuarantineReason.STALE_SNAPSHOT)

    def test_exactly_24h_is_fresh(self) -> None:
        result = run_import(mini_mapping(columns={"source_ts": ["absente"]}), csv_bytes([row()]), now=NOW,
                            source_ts=NOW - timedelta(hours=24))
        assert result.status is ImportStatus.ACCEPTED

    def test_source_ts_assumed_when_absent(self) -> None:
        mapping = mini_mapping(columns={"source_ts": ["absente"]})
        result = run_import(mapping, csv_bytes([row()]), now=NOW)
        assert result.anomalies_with(AnomalyCode.SOURCE_TS_ASSUMED)
        assert result.effective_source_ts == NOW and result.offers[0].source_ts == NOW
        old_fetch = run_import(mapping, csv_bytes([row()]), now=NOW, fetched_at=NOW - timedelta(hours=30))
        assert old_fetch.status is ImportStatus.QUARANTINED

    def test_snapshot_ts_from_newest_row(self) -> None:
        rows = [row(), row(sku="FICTIF-T-2", ean=g(902), maj="2026-10-04T07:00:00+02:00")]
        assert run_rows(rows).effective_source_ts == datetime.fromisoformat("2026-10-04T07:00:00+02:00")

    def test_same_content_as_previous(self) -> None:
        content = csv_bytes([row()])
        first = run_import(mini_mapping(), content, now=NOW)
        second = run_import(mini_mapping(), content, now=NOW, baseline=ImportBaseline.from_result(first))
        assert second.anomalies_with(AnomalyCode.SAME_CONTENT_AS_PREVIOUS)

    def test_quarantine_ratio(self) -> None:
        rows = [row(sku=f"FICTIF-T-{i}", ean=g(910 + i)) for i in range(5)]
        rows[0]["prix"] = "0"
        ok = run_rows(rows)
        assert ok.quarantine_ratio == Decimal("0.2") and not ok.anomalies_with(AnomalyCode.QUARANTINE_RATIO_EXCEEDED)
        rows[1]["prix"] = "0"
        bad = run_rows(rows)
        assert bad.anomalies_with(AnomalyCode.QUARANTINE_RATIO_EXCEEDED) and bad.escalation == "E1"
        assert bad.reason_counts() == {"NON_POSITIVE_PRICE": 2}

    def test_baseline_rules(self) -> None:
        good = run_rows([row()])
        base = ImportBaseline.from_result(good)
        assert base.prices["FICTIF-T-1"].unit_price == Decimal("10.00")
        assert next_baseline(None, good) == base
        rejected = run_import(mini_mapping(), csv_bytes([]), now=NOW)
        assert next_baseline(None, rejected) is None
        with pytest.raises(ImporterError):
            ImportBaseline.from_result(rejected)
        updated = base.updated_with(run_rows([row(prix="11.00"), row(sku="FICTIF-T-2", ean=g(902))]))
        assert updated.prices["FICTIF-T-1"].unit_price == Decimal("11.00") and "FICTIF-T-2" in updated.prices
        other = base.model_copy(update={"supplier_id": "fictif_autre"})
        with pytest.raises(ImporterError):
            other.updated_with(good)
        with pytest.raises(ImporterError):
            run_rows([row()], baseline=other)

    def test_policy_from_rules(self) -> None:
        policy = ImportPolicy.from_rules(load_rules())
        assert policy.max_age == timedelta(hours=24) and policy.price_anomaly_factor == Decimal(10)
        assert policy.future_skew == timedelta(minutes=5)
        assert ImportPolicy.from_rules(load_rules(), min_rows_ratio=Decimal("0.9")).min_rows_ratio == Decimal("0.9")

    def test_now_must_be_aware(self) -> None:
        with pytest.raises(ImporterError):
            run_rows([row()], now=datetime(2026, 10, 4, 8, 0))


# ------------------------------------------------------- pannes et lecture


class _Boom(SupplierConnector):
    def fetch(self) -> RawSnapshot:
        raise TimeoutError("flux fournisseur muet")


class TestFailures:
    def test_missing_file_is_failed_not_exception(self, tmp_path: Path) -> None:
        result = run_import(mini_mapping(), tmp_path / "absent.csv", now=NOW)
        assert result.status is ImportStatus.FAILED and result.escalation == "E1"
        assert result.anomalies[0].code == "FETCH_FAILED" and not result.offers
        assert result.retain_previous_offers
        with pytest.raises(FetchError):
            FileConnector(tmp_path / "absent.csv", mini_mapping(), clock=lambda: NOW).fetch()

    def test_connector_exception_becomes_incident(self) -> None:
        result = _Boom(mini_mapping(), clock=lambda: NOW).run()
        assert result.status is ImportStatus.FAILED and "TimeoutError" in result.anomalies[0].detail
        assert result.snapshot.source_uri == "(indisponible)"

    def test_connector_supplier_must_match(self) -> None:
        other = mini_mapping(supplier_id="fictif_autre")
        with pytest.raises(ImporterError):
            run_import(mini_mapping(), BytesConnector(b"", other, clock=lambda: NOW), now=NOW)
        snap = RawSnapshot.capture(supplier_id="fictif_autre", content=csv_bytes([row()]), source_kind=SourceKind.CSV,
                                   source_uri="m", fetched_at=NOW)
        with pytest.raises(ImporterError):
            run_import(mini_mapping(), snap, now=NOW)

    def test_outage_never_touches_local_stock(self) -> None:
        """BP §12 : une panne de flux n'efface pas le stock local confirmé."""
        registry = StockRegistry(clock=lambda: NOW)
        registry.receive("FICTIF-T-1", 3, "reception-FICTIF", expected_version=0, at=NOW)
        movements_before = registry.movements()
        previous = run_rows([row(stock="40")])
        book = merge_offer_book({}, previous).offers
        later = NOW + timedelta(hours=30)
        outage = run_import(mini_mapping(), Path("/nonexistent/FICTIF.csv"), now=later)
        update = merge_offer_book(book, outage)
        assert update.import_rejected and update.offers == book and update.retained == ("FICTIF-T-1",)
        assert registry.movements() == movements_before and registry.sellable("FICTIF-T-1") == 3
        promise = availability_promise(local_sellable=registry.sellable("FICTIF-T-1"), offers=list(update.offers.values()), now=later)
        assert promise.kind is PromiseKind.LOCAL_STOCK and promise.local_qty == 3
        assert not promise.restock_signal  # offre amont périmée : aucun réassort promis

    @pytest.mark.parametrize(
        ("content", "fmt"),
        [
            ("é;prix\n".encode("latin-1"), "csv"),
            (b"", "csv"),
            (b"pas un zip", "xlsx"),
            (b"<a><b></a>", "xml"),
            (b'<?xml version="1.0"?><!DOCTYPE a [<!ENTITY x "y">]><a>&x;</a>', "xml"),
            (b"<catalogue><!ENTITY x></catalogue>", "xml"),
            (b'sku;prix\n"FICTIF;12\n', "csv"),
        ],
    )
    def test_unreadable_sources(self, content: bytes, fmt: str) -> None:
        mapping = mini_mapping(file_format=fmt, required_fields=["supplier_sku", "price"])
        result = run_import(mapping, content, now=NOW)
        assert result.status is ImportStatus.FAILED
        assert result.anomalies[0].code == "UNREADABLE_SOURCE"

    def test_missing_and_ambiguous_columns(self) -> None:
        missing = run_import(mini_mapping(required_fields=["supplier_sku", "price", "currency"]),
                             csv_bytes([row()], header=[c for c in MINI_ROW if c != "devise"]), now=NOW)
        assert missing.status is ImportStatus.FAILED and missing.anomalies[0].code == "MAPPING_MISMATCH"
        ambiguous_map = mini_mapping(columns={"price": ["prix", "prix_net"]})
        ambiguous = run_import(ambiguous_map, csv_bytes([row(prix_net="9")], header=[*MINI_ROW, "prix_net"]), now=NOW)
        assert ambiguous.status is ImportStatus.FAILED and "ambiguës" in ambiguous.anomalies[0].detail

    def test_duplicate_header_is_unreadable(self) -> None:
        result = run_import(mini_mapping(), b"sku;prix;prix\nA;1;2\n", now=NOW)
        assert result.status is ImportStatus.FAILED and "double" in result.anomalies[0].detail

    def test_reader_format_mismatch(self) -> None:
        from pokeshop.importers import CsvReader, XlsxReader, XmlReader

        snap = RawSnapshot.capture(supplier_id="fictif_test", content=b"x", source_kind=SourceKind.CSV, source_uri="m",
                                   fetched_at=NOW)
        for reader, fmt in ((CsvReader(), "xml"), (XlsxReader(), "csv"), (XmlReader(), "csv")):
            with pytest.raises(MappingError):
                reader.read(snap, mini_mapping(file_format=fmt))

    def test_simulation_only_mappings_refuse_real_runs(self) -> None:
        with pytest.raises(MappingError, match="simulation"):
            run_import(mini_mapping(), csv_bytes([row()]), now=NOW, dry_run=False)
        with pytest.raises(MappingError, match="simulation"):
            run_import("asmodee_fr", csv_bytes([row()]), now=NOW, dry_run=False)
        real = deep_merge(MINI_MAPPING, {"supplier_id": "reel_test", "fictif": False, "status": "VALIDE",
                                         "validated_by": "Personne FICTIVE", "validated_at": "2026-10-04",
                                         "validated_on_sample_sha256": "b" * 64})
        real.pop("extension_tables")
        table = ExtensionTable.load([ROOT / "data" / "extensions_aliases.yaml", SAMPLES / "FICTIF_extensions_aliases.yaml"])
        result = run_import(SupplierMapping.model_validate(real), csv_bytes([row()]), now=NOW, dry_run=False, extensions=table)
        assert result.dry_run is False and "MODE RÉEL" not in result.report_markdown()
        assert "mode RÉEL" in result.report_markdown()

    def test_template_mapping_runs_in_simulation_with_warning(self) -> None:
        header = ["sku", "ean", "designation", "langue", "extension", "format", "contenu", "prix", "devise",
                  "base_prix", "unite_prix", "maj"]
        content = csv_bytes([{"sku": "FICTIF-1", "ean": g(950), "designation": "Display", "langue": "FR",
                              "extension": "Nuit Noire", "format": "Display", "contenu": "36 boosters", "prix": "95,00",
                              "devise": "EUR", "base_prix": "HT", "unite_prix": "unité", "maj": J0_TS}], header=header)
        result = run_import("asmodee_fr", content, now=NOW)
        assert result.anomalies_with(AnomalyCode.MAPPING_TEMPLATE)
        assert result.status is ImportStatus.ACCEPTED and result.offers[0].extension == "ME05"
        assert result.offers[0].sealed is None  # aucun défaut supposé dans un TEMPLATE
        assert "Dictionnaire TEMPLATE" in result.report_markdown()


# --------------------------------------------------------- formats spécifiques


class TestFormats:
    def test_csv_skip_rows_quotes_and_line_numbers(self) -> None:
        mapping = mini_mapping(csv={"delimiter": ";", "skip_rows": 2})
        body = csv_bytes([row(designation='Display "spécial"; édition'), row(sku="FICTIF-T-2", ean=g(902))],
                         header=[*MINI_ROW, "designation"])
        content = b"FICTIF titre;\nexport du jour\n" + body
        result = run_import(mapping, content, now=NOW)
        assert result.accepted_count == 2
        assert [int(o.raw_ref.rsplit("#L", 1)[1]) for o in result.offers] == [4, 5]

    def test_csv_encoding(self) -> None:
        mapping = mini_mapping(csv={"delimiter": ";", "encoding": "cp1252"})
        content = csv_bytes([row(langue="Français")]).decode("utf-8").encode("cp1252")
        assert run_import(mapping, content, now=NOW).offers[0].language == "FR"

    def test_csv_import_helper(self) -> None:
        path = SAMPLES / "FICTIF_offres_grossiste_a_J-1.csv"
        j1 = datetime.fromisoformat(EXPECTED[path.name]["now"])
        assert import_csv(path, "fictif_grossiste_a", now=j1).status is ImportStatus.ACCEPTED
        assert import_csv(path, "fictif_grossiste_a", now=NOW).status is ImportStatus.QUARANTINED  # 26 h plus tard

    def _xlsx(self, tmp_path: Path, rows: list[list[Any]], *, sheet: str = "Tarif", header_row: int = 1) -> bytes:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = sheet
        for _ in range(header_row - 1):
            ws.append(["titre"])
        for r in rows:
            ws.append(r)
        path = tmp_path / "t.xlsx"
        wb.save(path)
        return path.read_bytes()

    def test_xlsx_typed_cells(self, tmp_path: Path) -> None:
        mapping = mini_mapping(file_format="xlsx", xlsx={"sheet": "Tarif", "header_row": 2})
        header = list(MINI_ROW)
        values = row(prix="", ean="", stock="", tva="")
        cells: list[Any] = [values[h] for h in header]
        cells[header.index("prix")] = 12.3
        cells[header.index("ean")] = int(g(960))
        cells[header.index("stock")] = 7.0
        cells[header.index("tva")] = 0.081
        cells[header.index("maj")] = datetime(2026, 10, 4, 6, 0)
        result = run_import(mapping, self._xlsx(tmp_path, [header, cells], header_row=2), now=NOW)
        offer = result.offers[0]
        assert offer.price == Decimal("12.3") and offer.gtin == g(960) and offer.available_qty == 7
        assert offer.vat_rate == Decimal("0.081") and offer.source_ts == datetime(2026, 10, 4, 6, 0, tzinfo=ZRH)

    def test_xlsx_errors(self, tmp_path: Path) -> None:
        wrong_sheet = mini_mapping(file_format="xlsx", xlsx={"sheet": "Absente"})
        content = self._xlsx(tmp_path, [list(MINI_ROW)])
        result = run_import(wrong_sheet, content, now=NOW)
        assert result.status is ImportStatus.FAILED and "Absente" in result.anomalies[0].detail
        beyond = mini_mapping(file_format="xlsx", xlsx={"header_row": 50})
        assert run_import(beyond, content, now=NOW).status is ImportStatus.FAILED
        empty_header = mini_mapping(file_format="xlsx")
        assert run_import(empty_header, self._xlsx(tmp_path, [[None, None], ["a", "b"]]), now=NOW).status is ImportStatus.FAILED
        dup = mini_mapping(file_format="xlsx")
        assert run_import(dup, self._xlsx(tmp_path, [["sku", "sku"], ["a", "b"]]), now=NOW).status is ImportStatus.FAILED

    def test_xlsx_export_timestamp_cell(self) -> None:
        result = import_xlsx(SAMPLES / "FICTIF_tarif_grossiste_b.xlsx", "fictif_grossiste_b", now=NOW)
        assert result.effective_source_ts == datetime(2026, 10, 4, 7, 30, tzinfo=ZRH)
        late = import_xlsx(SAMPLES / "FICTIF_tarif_grossiste_b.xlsx", "fictif_grossiste_b", now=NOW + timedelta(days=2))
        assert late.status is ImportStatus.QUARANTINED

    def test_xml_flatten(self) -> None:
        import xml.etree.ElementTree as ET

        el = ET.fromstring(
            '<offre sku="A"><prix devise="EUR">1.5</prix><paliers><palier qte="6"/><palier qte="12"/></paliers>'
            "<!-- commentaire --><vide/><n:x xmlns:n='urn:fictif'>ns</n:x></offre>"
        )
        flat = flatten_element(el)
        assert flat == {
            "@sku": "A", "prix": "1.5", "prix/@devise": "EUR", "paliers/palier/@qte": "6",
            "paliers/palier[2]/@qte": "12", "vide": None, "x": "ns",
        }

    def test_xml_namespaces_and_bad_path(self) -> None:
        xml = (
            '<c:catalogue xmlns:c="urn:fictif" exporte_le="2026-10-04T06:00:00+02:00">'
            f'<c:offre sku="FICTIF-N-1"><c:prix devise="EUR" base="HT">9.00</c:prix><c:ean>{g(970)}</c:ean>'
            '<c:langue>FR</c:langue><c:format>Display</c:format><c:extension>Fictive Alpha</c:extension>'
            '<c:contenu>36 boosters</c:contenu><c:cond unite="unité"/></c:offre></c:catalogue>'
        ).encode()
        mapping = mini_mapping(
            file_format="xml",
            xml={"record_path": "./c:offre", "source_ts_attribute": "exporte_le", "namespaces": {"c": "urn:fictif"}},
            columns={"supplier_sku": ["@sku"], "price": ["prix"], "currency": ["prix/@devise"], "price_basis": ["prix/@base"],
                     "unit_basis": ["cond/@unite"], "gtin": ["ean"], "language": ["langue"], "format": ["format"],
                     "extension": ["extension"], "content": ["contenu"], "source_ts": ["absent"]},
        )
        result = import_xml(xml, mapping, now=NOW)
        assert result.status is ImportStatus.ACCEPTED and result.offers[0].price == Decimal("9.00")
        bad = mapping.model_copy(update={"xml": mapping.xml.model_copy(update={"record_path": "./["})})
        assert run_import(bad, xml, now=NOW).anomalies[0].code == "MAPPING_MISMATCH"

    def test_xml_unparseable_export_timestamp(self) -> None:
        xml = b'<catalogue exporte_le="hier"><offre sku="X"><prix>1</prix></offre></catalogue>'
        mapping = mini_mapping(file_format="xml", xml={"record_path": "./offre", "source_ts_attribute": "exporte_le"},
                               columns={"supplier_sku": ["@sku"], "price": ["prix"]},
                               required_fields=["supplier_sku", "price"])
        result = run_import(mapping, xml, now=NOW)
        assert any(a.code == "UNPARSEABLE_VALUE" and a.severity is Severity.CRITICAL for a in result.anomalies)


# --------------------------------------------------- vue des offres et moteur


class TestIntegration:
    def test_merge_offer_book(self) -> None:
        first = run_rows([row(), row(sku="FICTIF-T-2", ean=g(902))])
        book = merge_offer_book({}, first)
        assert book.added == ("FICTIF-T-1", "FICTIF-T-2") and not book.import_rejected
        second = run_rows([row(prix="10.50"), row(sku="FICTIF-T-3", ean=g(903))])
        update = merge_offer_book(book.offers, second)
        assert update.updated == ("FICTIF-T-1",) and update.added == ("FICTIF-T-3",) and update.retained == ("FICTIF-T-2",)
        assert update.offers["FICTIF-T-1"].price == Decimal("10.50") and "FICTIF-T-2" in update.offers
        with pytest.raises(ImporterError):
            merge_offer_book({"AUTRE": first.offers[0]}, second)

    def test_pricing_blocks_wrong_language_and_drafts_unknowns(self) -> None:
        rules = load_rules()
        result = _import_sample("FICTIF_offres_grossiste_a.csv")
        offers = result.offers_by_sku
        fx = {"fx_rate_to_chf": Decimal("0.94"), "fx_source": "FICTIF", "fx_date": date(2026, 10, 4)}
        costs = {"inbound_freight_alloc": Decimal("1.00"), "customs_and_fees": Decimal("0.50")}
        jp = evaluate_offer(offers["FICTIF-A-038"], rules.pricing, now=NOW, **fx, **costs)  # type: ignore[arg-type]
        assert jp.status is DecisionStatus.BLOCKED and jp.has(Reason.LANGUAGE_MISMATCH)
        no_gtin = evaluate_offer(offers["FICTIF-A-025"], rules.pricing, now=NOW, **fx, **costs)  # type: ignore[arg-type]
        assert no_gtin.status is DecisionStatus.DRAFT and no_gtin.has(Reason.UNKNOWN_FIELDS)
        sleeves = offers["FICTIF-A-008"]
        acc = evaluate_offer(sleeves, rules.pricing, now=NOW, expected_language=expected_language_for(sleeves.format),
                             **fx, **costs)  # type: ignore[arg-type]
        assert not acc.has(Reason.LANGUAGE_MISMATCH) and acc.landed_cost is not None
        display = evaluate_offer(offers["FICTIF-A-001"], rules.pricing, now=NOW, **fx, **costs)  # type: ignore[arg-type]
        assert not display.has(Reason.UNKNOWN_FIELDS) and display.landed_cost is not None
        assert display.recommended_price is not None and display.restock_eligible

    def test_catalog_matching_of_imported_offers(self, fictif_table: ExtensionTable) -> None:
        previous = run_import("fictif_grossiste_a", SAMPLES / "FICTIF_offres_grossiste_a_J-1.csv",
                              now=datetime.fromisoformat(EXPECTED["FICTIF_offres_grossiste_a_J-1.csv"]["now"]))
        products = [
            CatalogProduct(product_id=f"P-{o.supplier_sku}", identity=o.identity())
            for o in previous.offers
            if o.supplier_sku in ("FICTIF-A-001", "FICTIF-A-002", "FICTIF-A-008")
        ]
        index = CatalogIndex(products, fictif_table)
        today = _import_sample("FICTIF_offres_grossiste_a.csv").offers_by_sku
        assert match_offer_to_product(today["FICTIF-A-001"], index).status is MatchStatus.MATCHED
        assert match_offer_to_product(today["FICTIF-A-027"], index).product_id == "P-FICTIF-A-001"  # carton du même display
        assert match_offer_to_product(today["FICTIF-A-008"], index).status is MatchStatus.MATCHED
        same_name = match_offer_to_product(today["FICTIF-A-033"], index)
        assert same_name.status is MatchStatus.AMBIGUOUS and "SAME_NAME_DIFFERENT_CONTENT" in same_name.reasons
        assert match_offer_to_product(today["FICTIF-A-025"], index).status is MatchStatus.NEW_DRAFT
        assert match_offer_to_product(today["FICTIF-A-030"], index).is_draft
        assert normalize_gtin(today["FICTIF-A-001"].gtin) == today["FICTIF-A-001"].gtin

    def test_snapshot_level_reasons_are_labelled(self) -> None:
        assert set(QuarantineReason) == set(QUARANTINE_LABELS_FR)
        assert SNAPSHOT_LEVEL_REASONS <= set(QuarantineReason)


# ------------------------------------------------------------- import assisté


def assisted_data() -> dict[str, Any]:
    return yaml.safe_load((SAMPLES / "FICTIF_tarif_email_assiste.yaml").read_text(encoding="utf-8"))


def fixed_assisted() -> AssistedPriceList:
    data = assisted_data()
    data["lines"][4]["line_total"] = "82,00"
    data["declared_total"] = "1 318,80"
    return AssistedPriceList.model_validate(data)


class TestAssisted:
    def test_sample_review_blocks_on_line_total(self) -> None:
        review = review_price_list(load_price_list(SAMPLES / "FICTIF_tarif_email_assiste.yaml"), now=NOW)
        assert review.status == "A_VALIDER_HUMAINEMENT" and review.requires_human_validation is True
        assert [(c.code, c.line_no) for c in review.blocking_checks] == [("LINE_TOTAL_MATCH", 5)]
        assert [(c.code, c.line_no) for c in review.warnings] == [("GTIN", 6)]
        assert review.computed_total == Decimal("1281.80") and review.declared_total == Decimal("1316.80")
        assert not review.is_approvable and review.fictif
        assert {o.supplier_sku for o in review.draft_offers} == {"FICTIF-D-001", "FICTIF-D-002", "FICTIF-D-003",
                                                                 "FICTIF-D-004", "FICTIF-D-006"}
        with pytest.raises(AssistedImportError, match="LINE_TOTAL_MATCH@L5"):
            approve_review(review, validated_by="Personne FICTIVE", validated_at=NOW)
        report = review.report_markdown()
        assert "A_VALIDER_HUMAINEMENT" in report and "## Validation humaine requise" in report and "ÉCHEC" in report

    def test_expected_file_matches(self) -> None:
        spec = EXPECTED["FICTIF_tarif_email_assiste.yaml"]
        review = review_price_list(load_price_list(SAMPLES / "FICTIF_tarif_email_assiste.yaml"),
                                   now=datetime.fromisoformat(spec["now"]))
        assert sorted(f"{c.code}@{c.line_no}" for c in review.blocking_checks) == sorted(spec["blocking"])
        assert sorted(f"{c.code}@{c.line_no}" for c in review.warnings) == sorted(spec["warnings"])

    def test_corrected_review_can_be_approved_by_a_person(self) -> None:
        review = review_price_list(fixed_assisted(), now=NOW)
        assert review.is_approvable and len(review.draft_offers) == 6
        result = approve_review(review, validated_by="  Personne FICTIVE  ", validated_at=NOW)
        assert result.status is ImportStatus.ACCEPTED and result.assisted and result.validated_by == "Personne FICTIVE"
        assert result.anomalies[0].code == "ASSISTED_IMPORT_VALIDATED" and "Personne FICTIVE" in result.anomalies[0].detail
        assert result.snapshot.source_kind is SourceKind.EMAIL and result.snapshot.checksum_sha256 == assisted_data()["document_sha256"]
        carton = result.offers_by_sku["FICTIF-D-004"]
        assert carton.units_per_pack == 24 and carton.carton_qty == 24 and carton.price == Decimal("268.80")
        sleeves = result.offers_by_sku["FICTIF-D-005"]
        assert sleeves.language == "NA" and sleeves.extension == "SANS_EXTENSION"
        assert result.offers_by_sku["FICTIF-D-006"].gtin is None
        assert "Import assisté validé par : Personne FICTIVE" in result.report_markdown()
        with pytest.raises(AssistedImportError, match="validated_by"):
            approve_review(review, validated_by="  ", validated_at=NOW)
        with pytest.raises(AssistedImportError, match="fuseau"):
            approve_review(review, validated_by="X", validated_at=datetime(2026, 10, 4))

    @pytest.mark.parametrize(
        ("mutate", "code"),
        [
            (lambda d: d.update(currency="$"), "CURRENCY"),
            (lambda d: d.update(currency=None), "CURRENCY"),
            (lambda d: d.update(currency="GBP"), "CURRENCY"),
            (lambda d: d.update(price_basis="net"), "PRICE_BASIS"),
            (lambda d: d.update(vat_rate="abc"), "VAT_RATE"),
            (lambda d: d.update(declared_total=None), "DOCUMENT_TOTAL"),
            (lambda d: d.update(declared_total="1 300,00"), "DOCUMENT_TOTAL"),
            (lambda d: d.update(declared_shipping="x"), "DOCUMENT_TOTAL"),
            (lambda d: d["lines"][0].update(supplier_sku=" "), "SKU"),
            (lambda d: d["lines"][1].update(supplier_sku="FICTIF-D-001"), "DUPLICATE_SKU"),
            (lambda d: d["lines"][0].update(quantity="deux"), "QUANTITY"),
            (lambda d: d["lines"][0].update(quantity="0"), "QUANTITY"),
            (lambda d: d["lines"][0].update(unit_price="0,00", line_total="0,00"), "UNIT_PRICE"),
            (lambda d: d["lines"][0].update(unit_price="95,00 CHF"), "UNIT_PRICE"),
            (lambda d: d["lines"][0].update(unit_price="9x"), "UNIT_PRICE"),
            (lambda d: d["lines"][0].update(line_total=None), "LINE_TOTAL"),
            (lambda d: d["lines"][0].update(line_total="1x"), "LINE_TOTAL"),
            (lambda d: d["lines"][0].update(unit_basis="palette"), "UNIT_BASIS"),
            (lambda d: d["lines"][3].update(units_per_pack=None), "UNIT_BASIS"),
            (lambda d: d["lines"][3].update(units_per_pack="vingt"), "UNITS_PER_PACK"),
            (lambda d: d["lines"][0].update(units_per_pack="6"), "UNIT_BASIS"),
            (lambda d: d["lines"][0].update(gtin=g(401)[:-1] + str((int(g(401)[-1]) + 1) % 10)), "GTIN"),
            (lambda d: d["lines"][0].update(language="VO"), "LANGUAGE"),
            (lambda d: d["lines"][0].update(vat_rate="200 %"), "VAT_RATE"),
        ],
    )
    def test_blocking_checks(self, mutate: Any, code: str) -> None:
        data = fixed_assisted().model_dump(mode="json")
        mutate(data)
        review = review_price_list(AssistedPriceList.model_validate(data), now=NOW)
        assert code in {c.code for c in review.blocking_checks}, [c for c in review.checks if not c.ok]
        with pytest.raises(AssistedImportError):
            approve_review(review, validated_by="Personne FICTIVE", validated_at=NOW)

    def test_stale_document_is_warning_only(self) -> None:
        review = review_price_list(fixed_assisted(), now=NOW + timedelta(days=3))
        assert review.is_approvable and ("FRESHNESS", None) in [(c.code, c.line_no) for c in review.warnings]

    def test_tolerance_of_one_cent(self) -> None:
        data = fixed_assisted().model_dump(mode="json")
        data["declared_total"] = "1 318,81"
        assert review_price_list(AssistedPriceList.model_validate(data), now=NOW).is_approvable
        data["declared_total"] = "1 318,82"
        assert not review_price_list(AssistedPriceList.model_validate(data), now=NOW).is_approvable

    def test_no_offer_when_document_level_fails(self) -> None:
        data = fixed_assisted().model_dump(mode="json")
        data["currency"] = None
        review = review_price_list(AssistedPriceList.model_validate(data), now=NOW)
        assert review.draft_offers == ()

    def test_approve_requires_offers(self) -> None:
        review = review_price_list(fixed_assisted(), now=NOW).model_copy(update={"draft_offers": ()})
        with pytest.raises(AssistedImportError, match="aucune offre"):
            approve_review(review, validated_by="X", validated_at=NOW)

    @pytest.mark.parametrize(
        "change",
        [
            {"document_sha256": "abc"},
            {"received_at": "2026-10-04T07:00:00"},
            {"ship_from_country": "France"},
            {"lines": []},
            {"fictif": False},
            {"source_kind": "FAX"},
        ],
    )
    def test_invalid_documents(self, change: dict[str, Any]) -> None:
        data = assisted_data()
        data.update(change)
        with pytest.raises(ValueError):
            AssistedPriceList.model_validate(data)

    def test_load_errors(self, tmp_path: Path) -> None:
        with pytest.raises(AssistedImportError):
            load_price_list(tmp_path / "absent.yaml")
        bad = tmp_path / "bad.yaml"
        bad.write_text("supplier_id: x\n", encoding="utf-8")
        with pytest.raises(AssistedImportError, match="invalide"):
            load_price_list(bad)

    def test_explicit_table(self, fictif_table: ExtensionTable) -> None:
        review = review_price_list(fixed_assisted(), now=NOW, table=fictif_table)
        assert review.is_approvable
