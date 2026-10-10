"""Tests du module pokeshop.forecast et du classeur modele_financier.xlsx (BP §3, §4, §10, étoile polaire)."""

from __future__ import annotations

import importlib.util
import sys
from decimal import Decimal
from pathlib import Path
from types import ModuleType

import pytest

from pokeshop import forecast as fc
from pokeshop.forecast import (
    CENT,
    CHF,
    CheckStatus,
    ForecastError,
    ScenarioAssumptions,
    WorkingCapitalAssumptions,
)

D = Decimal
ROOT = Path(__file__).resolve().parents[1]
FINANCE_DIR = ROOT / "docs" / "03-finance"


def r0(value: Decimal) -> Decimal:
    """Arrondi au franc (affichage BP)."""
    return fc.round_chf(value, CHF)


def r2(value: Decimal) -> Decimal:
    """Arrondi au centime."""
    return fc.round_chf(value, CENT)


# ---------------------------------------------------------------------------
# Utilitaires
# ---------------------------------------------------------------------------


class TestDecimalHelpers:
    def test_to_decimal_rejects_float(self) -> None:
        with pytest.raises(ForecastError, match="float"):
            fc.to_decimal(0.1)

    def test_to_decimal_rejects_bool_and_garbage(self) -> None:
        with pytest.raises(ForecastError):
            fc.to_decimal(True)
        with pytest.raises(ForecastError):
            fc.to_decimal("abc")
        with pytest.raises(ForecastError):
            fc.to_decimal(D("NaN"))
        with pytest.raises(ForecastError):
            fc.to_decimal("Infinity")
        with pytest.raises(ForecastError):
            fc.to_decimal([1])  # type: ignore[arg-type]

    def test_to_decimal_accepts_int_str_decimal(self) -> None:
        assert fc.to_decimal(3) == D(3)
        assert fc.to_decimal("0.081") == D("0.081")
        assert fc.to_decimal(D("1.5")) == D("1.5")

    def test_round_chf_half_up(self) -> None:
        assert fc.round_chf(D("5.2975")) == D("5.30")
        assert fc.round_chf(D("2.5"), CHF) == D(3)
        assert fc.round_chf(D("53.358"), CHF) == D(53)

    def test_ceil_int(self) -> None:
        assert fc.ceil_int(D("30.0001")) == 31
        assert fc.ceil_int(D("30")) == 30

    def test_basket_ht(self) -> None:
        assert r2(fc.basket_ht()) == D("87.88")
        assert fc.basket_ht(D(95), D(0)) == D(95)


# ---------------------------------------------------------------------------
# Scénarios BP §10 — reproduction exacte
# ---------------------------------------------------------------------------

BP_TABLE = {
    # nom : (commandes, TTC, HT, contribution, CAC, acquisition, fixes, résultat)
    "prudent": (40, "3800", "3515", "773", "8", "320", "400", "53"),
    "central": (100, "9500", "8788", "1933", "6", "600", "400", "933"),
    "developpement": (200, "19000", "17576", "3867", "5", "1000", "400", "2467"),
}


@pytest.mark.parametrize("name", list(BP_TABLE))
def test_bp_scenario_table(name: str) -> None:
    orders, ttc, ht, contrib, cac, acq, fixed, result = BP_TABLE[name]
    res = fc.run_bp_scenarios()[name]
    rounded = res.rounded()
    assert res.assumptions.orders_per_month == orders
    assert rounded["revenue_ttc"] == D(ttc)
    assert rounded["revenue_ht"] == D(ht)
    assert rounded["contribution_before_acquisition"] == D(contrib)
    assert res.assumptions.cac == D(cac)
    assert rounded["acquisition_total"] == D(acq)
    assert rounded["fixed_costs"] == D(fixed)
    assert rounded["result_before_remuneration"] == D(result)


def test_development_ht_differs_from_bp_by_one_franc() -> None:
    """BP affiche 17 577 ; 19 000 / 1,081 = 17 576,32 (écart documenté)."""
    res = fc.run_bp_scenarios()["developpement"]
    assert r2(res.revenue_ht) == D("17576.32")
    assert r0(res.revenue_ht) == D("17576") != D("17577")
    # contribution et résultat publiés restent cohérents avec la valeur exacte
    assert r0(res.contribution_before_acquisition) == D("3867")
    assert r0(res.result_before_remuneration) == D("2467")


def test_exact_values_are_not_prematurely_rounded() -> None:
    res = fc.run_bp_scenarios()["prudent"]
    assert res.result_before_remuneration == res.contribution_before_acquisition - D(320) - D(400)
    assert r2(res.result_before_remuneration) == D("53.36")
    assert res.vat_collected == res.revenue_ttc - res.revenue_ht


def test_central_per_order_contribution() -> None:
    cen = fc.run_bp_scenarios()["central"]
    assert r2(cen.contribution_per_order) == D("19.33")
    assert r2(cen.contribution_per_order_after_cac) == D("13.33")
    assert r2(cen.basket_ht) == D("87.88")
    # après CAC : 15,17 % du CA HT, sous la cible 20 % du BP §5 (écart de cohérence documenté)
    assert D("0.15") < cen.contribution_after_cac_pct < D("0.20")


def test_annualized_and_vat_threshold() -> None:
    results = fc.run_bp_scenarios()
    assert r0(results["central"].annual_revenue_ttc) == D("114000")
    assert fc.vat_threshold_check(results["central"].annual_revenue_ttc).exceeds
    assert fc.vat_threshold_check(results["central"].annual_revenue_ht).exceeds  # 105 458 HT
    prudent = fc.vat_threshold_check(results["prudent"].annual_revenue_ttc)
    assert not prudent.exceeds
    assert prudent.ratio == D("0.456")
    with pytest.raises(ForecastError):
        fc.vat_threshold_check(D(1), threshold=D(0))
    with pytest.raises(ForecastError):
        fc.vat_threshold_check(D(-1))


def test_remuneration_is_deducted() -> None:
    results = fc.run_bp_scenarios(remuneration=D(2000))
    assert r0(results["central"].result_after_remuneration) == D("-1067")
    assert r0(results["developpement"].result_after_remuneration) == D("467")
    rounded = results["central"].rounded()
    assert rounded["remuneration"] == D(2000)


def test_non_registered_variant_runs_with_zero_vat() -> None:
    a = ScenarioAssumptions("non_assujetti", 100, D(6), vat_rate=D(0), contribution_rate=D("0.20"))
    res = fc.run_scenario(a)
    assert res.revenue_ht == res.revenue_ttc == D(9500)
    assert res.vat_collected == 0
    assert res.contribution_before_acquisition == D(1900)


def test_zero_orders_scenario() -> None:
    res = fc.run_scenario(ScenarioAssumptions("vide", 0, D(6)))
    assert res.revenue_ttc == 0
    assert res.result_before_remuneration == D(-400)


@pytest.mark.parametrize(
    "kwargs",
    [
        {"orders_per_month": -1},
        {"orders_per_month": 1.5},
        {"orders_per_month": True},
        {"cac": D(-1)},
        {"basket_ttc": D(0)},
        {"vat_rate": D(1)},
        {"vat_rate": D("-0.01")},
        {"contribution_rate": D("1.01")},
        {"fixed_costs": D(-1)},
        {"remuneration": D(-1)},
        {"cac": 6.0},
        {"name": ""},
    ],
)
def test_scenario_validation(kwargs: dict) -> None:
    base = {"name": "x", "orders_per_month": 10, "cac": D(5)}
    base.update(kwargs)
    with pytest.raises(ForecastError):
        ScenarioAssumptions(**base)


def test_scenario_accepts_str_amounts() -> None:
    a = ScenarioAssumptions("x", 10, "6", basket_ttc="95")  # type: ignore[arg-type]
    assert a.cac == D(6) and a.basket_ttc == D(95)


# ---------------------------------------------------------------------------
# Seuil de rentabilité et sensibilité
# ---------------------------------------------------------------------------


def test_break_even_bp_method_gives_31() -> None:
    be = fc.break_even(unit_rounding=CENT)
    assert be.contribution_per_order_after_cac == D("13.33")
    assert be.orders_required == 31
    assert r2(be.orders_exact) == D("30.01")


def test_break_even_exact_gives_30_but_fragile() -> None:
    be = fc.break_even()
    assert be.orders_required == 30
    assert r2(be.orders_exact) == D("30.00")
    assert be.orders_exact < 30
    assert D(0) < be.surplus_at_required < D("0.02")
    assert be.is_fragile


def test_break_even_with_remuneration_181_vs_180() -> None:
    assert fc.break_even(remuneration=D(2000), unit_rounding=CENT).orders_required == 181
    exact = fc.break_even(remuneration=D(2000))
    assert exact.orders_required == 180
    assert exact.amount_to_cover == D(2400)


def test_sensitivity_contribution_15_pct() -> None:
    be = fc.break_even(contribution_rate=D("0.15"))
    assert r2(be.contribution_per_order) == D("13.18")
    assert r2(be.contribution_per_order_after_cac) == D("7.18")
    assert be.orders_required == 56
    assert fc.break_even(contribution_rate=D("0.15"), unit_rounding=CENT).orders_required == 56


def test_sensitivity_cac_15() -> None:
    be = fc.break_even(cac=D(15))
    assert r2(be.contribution_per_order_after_cac) == D("4.33")
    assert be.orders_required == 93
    assert fc.break_even(cac=D(15), unit_rounding=CENT).orders_required == 93


def test_break_even_impossible_raises() -> None:
    with pytest.raises(ForecastError, match="≤ 0"):
        fc.break_even(cac=D(20))
    with pytest.raises(ForecastError):
        fc.break_even(fixed_costs=D(-1))
    with pytest.raises(ForecastError):
        fc.break_even(contribution_rate=D("1.5"))


def test_sensitivity_grid_flags_hard_floor_and_unreachable() -> None:
    grid = fc.sensitivity_grid([D("0.15"), D("0.22")], [D(6), D(15), D(25)])
    assert len(grid) == 6
    by_key = {(c.contribution_rate, c.cac): c for c in grid}
    c15 = by_key[(D("0.15"), D(6))]
    assert c15.orders_required == 56 and c15.below_hard_floor  # 7,18 < 8 CHF (BP §5)
    cac15 = by_key[(D("0.22"), D(15))]
    assert cac15.orders_required == 93 and cac15.below_hard_floor
    central = by_key[(D("0.22"), D(6))]
    assert central.orders_required == 30 and not central.below_hard_floor
    never = by_key[(D("0.22"), D(25))]
    assert never.orders_required is None and not never.reachable
    rounded = fc.sensitivity_grid([D("0.22")], [D(6)], unit_rounding=CENT)
    assert rounded[0].orders_required == 31
    with_rem = fc.sensitivity_grid([D("0.22")], [D(6)], remuneration=D(2000))
    assert with_rem[0].orders_required == 180


def test_sensitivity_grid_validation() -> None:
    with pytest.raises(ForecastError):
        fc.sensitivity_grid([], [D(6)])
    with pytest.raises(ForecastError):
        fc.sensitivity_grid([D("0.22")], [D(-1)])
    with pytest.raises(ForecastError):
        fc.sensitivity_grid([D(2)], [D(6)])


# ---------------------------------------------------------------------------
# Stock, BFR, temps
# ---------------------------------------------------------------------------


def test_stock_need_central() -> None:
    need = fc.stock_need(100)
    assert r0(need.revenue_ht) == D(8788)
    assert r0(need.monthly_cogs_ht) == D(6152)
    assert r2(need.monthly_cogs_ht) == D("6151.71")
    # le stock initial de 3 000 CHF ne finance pas un mois central (BP §10)
    assert need.coverage_ratio < D("0.5")
    assert r2(need.coverage_days) == D("14.63")
    assert r0(need.shortfall_one_month) == D(3152)


def test_stock_need_prudent_is_covered() -> None:
    need = fc.stock_need(40)
    assert need.shortfall_one_month == 0
    assert need.coverage_ratio > 1


@pytest.mark.parametrize(
    "kwargs",
    [{"orders_per_month": 0}, {"product_cost_ratio": D("1.2")}, {"initial_stock": D(-1)}, {"days_per_month": 0},
     {"product_cost_ratio": D(0)}],
)
def test_stock_need_validation(kwargs: dict) -> None:
    params = {"orders_per_month": 100}
    params.update(kwargs)
    orders = params.pop("orders_per_month")
    with pytest.raises(ForecastError):
        fc.stock_need(orders, **params)


def test_working_capital_need_defaults() -> None:
    wc = fc.working_capital_need(100)
    daily_cogs = fc.stock_need(100).monthly_cogs_ht / 30
    assert wc.daily_cogs_ht == daily_cogs
    assert wc.stock_component == daily_cogs * 44
    assert wc.psp_component == D(9500) / 30 * 7
    assert wc.supplier_credit_component == 0
    assert r0(wc.total) == D(11239)


def test_working_capital_supplier_credit_reduces_need() -> None:
    base = fc.working_capital_need(100)
    credit = fc.working_capital_need(100, WorkingCapitalAssumptions(supplier_credit_days=30))
    assert credit.total == base.total - credit.supplier_credit_component
    assert credit.supplier_credit_component > 0


def test_working_capital_assumptions_validation() -> None:
    with pytest.raises(ForecastError):
        WorkingCapitalAssumptions(stock_cover_days=-1)
    with pytest.raises(ForecastError):
        WorkingCapitalAssumptions(psp_payout_delay_days=True)  # type: ignore[arg-type]


def test_time_valuation() -> None:
    cen = fc.run_bp_scenarios()["central"]
    tv = fc.time_valuation(cen.result_before_remuneration, supervision_hours_per_week=D(8), orders_per_month=100,
                           prep_minutes_per_order=D(15))
    assert r2(tv.hours_per_month) == D("59.67")
    assert tv.implicit_hourly_rate is not None
    assert r2(tv.implicit_hourly_rate) == D("15.64")
    prudent = fc.run_bp_scenarios()["prudent"]
    low = fc.time_valuation(prudent.result_before_remuneration, supervision_hours_per_week=D(6))
    assert r2(low.hours_per_month) == D("26.00")
    assert low.implicit_hourly_rate is not None and low.implicit_hourly_rate < D(3)
    zero = fc.time_valuation(D(100), supervision_hours_per_week=D(0))
    assert zero.implicit_hourly_rate is None
    with pytest.raises(ForecastError):
        fc.time_valuation(D(1), supervision_hours_per_week=D(-1))
    with pytest.raises(ForecastError):
        fc.time_valuation(D(1), supervision_hours_per_week=D(1), orders_per_month=-1)


# ---------------------------------------------------------------------------
# Prix plancher (contre-vérification BP §4)
# ---------------------------------------------------------------------------

EX = fc.BP_PRICE_EXAMPLE
COMMON = {k: EX[k] for k in ("payment_pct", "payment_fixed", "logistics", "after_sales", "acquisition", "vat_rate")}


def test_floor_price_registered() -> None:
    floor = fc.floor_price_crosscheck(EX["cost"], target_margin=EX["target_margin"], **COMMON)
    assert r2(floor) == D("208.79")
    # marge d'arrondi très faible : 208,79498
    assert D("208.794") < floor < D("208.795")
    assert fc.round_up_to_ending(floor) == D("209.90")


def test_floor_price_not_registered() -> None:
    cost = r2(EX["cost"] * D("1.081"))
    assert cost == D("151.34")
    params = dict(COMMON, vat_rate=D(0))
    floor = fc.floor_price_crosscheck(cost, target_margin=EX["target_margin"], **params)
    assert r2(floor) == D("207.28")


def test_floor_price_denominator_guard() -> None:
    with pytest.raises(ForecastError, match="dénominateur"):
        fc.floor_price_crosscheck(D(10), target_margin=D("0.98"), **COMMON)


def test_round_up_to_ending_variants() -> None:
    assert fc.round_up_to_ending(D("208.79"), step=D(1), ending=D("0.90")) == D("208.90")
    assert fc.round_up_to_ending(D("199.90")) == D("199.90")
    assert fc.round_up_to_ending(D("199.91")) == D("209.90")
    assert fc.round_up_to_ending(D("5")) == D("9.90")
    with pytest.raises(ForecastError):
        fc.round_up_to_ending(D(10), step=D(0))
    with pytest.raises(ForecastError):
        fc.round_up_to_ending(D(10), step=D(1), ending=D("1.5"))


def test_contribution_at_199_90() -> None:
    chk = fc.contribution_at_price(D("199.90"), EX["cost"], **COMMON)
    assert r2(chk.net_sales) == D("184.92")
    assert r2(chk.payment_fees) == D("5.30")
    assert r2(chk.contribution) == D("30.62")
    assert fc.round_chf(chk.contribution_pct, D("0.0001")) == D("0.1656")


def test_contribution_at_rounded_price_meets_target() -> None:
    chk = fc.contribution_at_price(D("209.90"), EX["cost"], **COMMON)
    assert chk.contribution_pct >= D("0.20")
    with pytest.raises(ForecastError):
        fc.contribution_at_price(D(0), EX["cost"], **COMMON)


# ---------------------------------------------------------------------------
# Étoile polaire : contribution nette cumulée hebdomadaire + stop-loss global
# ---------------------------------------------------------------------------


def nsw(week: int, net: str = "0", cost: str = "0", acq: str = "0", fixed: str = "0", orders: str = "0",
        pay: str = "0", log: str = "0", sav: str = "0") -> fc.NorthStarWeek:
    """Semaine FICTIVE de suivi étoile polaire."""
    return fc.NorthStarWeek(week, D(orders), D(net), D(cost), D(pay), D(log), D(sav), D(acq), D(fixed))


def test_north_star_week_components() -> None:
    w = nsw(1, net="1000", cost="700", pay="30", log="35", sav="10", acq="60", fixed="92", orders="11")
    assert w.contribution_before_acquisition == D(225)
    assert w.contribution_after_acquisition == D(165)
    assert w.net_contribution == D(73)


@pytest.mark.parametrize("name", list(BP_TABLE))
def test_projection_over_a_year_equals_twelve_monthly_results(name: str) -> None:
    a = next(s for s in fc.bp_scenarios() if s.name == name)
    weeks = fc.project_north_star(a)
    assert len(weeks) == fc.WEEKS_PER_YEAR == 52
    report = fc.north_star(weeks)
    assert r2(report.cumulative) == r2(fc.run_scenario(a).result_before_remuneration * 12)
    # contribution avant acquisition projetée = 22 % du CA HT (ventilation cohérente)
    w = weeks[0]
    assert r2(w.contribution_before_acquisition) == r2(w.net_sales_ht * D("0.22"))
    assert r2(w.logistics / w.orders) == D("3.36")  # emballage implicite (port refacturé)
    assert r2(w.payment_fees / w.orders) == D("2.68")  # 2,5 % × 95 + 0,30


def test_projection_central_with_opening_week_5() -> None:
    central = fc.bp_scenarios()[1]
    report = fc.north_star(fc.project_north_star(central, opening_week=fc.BP_SALES_OPENING_WEEK))
    assert report.rows[0].week.orders == 0 and r2(report.rows[0].net_contribution) == D("-92.31")
    assert r2(report.lowest.cumulative) == D("-369.23") and report.lowest.week.week == 4
    assert report.recovery_week == 6
    assert r2(report.rows[4].net_contribution) == D("215.40")
    assert r2(report.rows[4].net_contribution * fc.WEEKS_PER_MONTH) == D("933.40")
    assert not report.frozen and report.first_global_stoploss_week is None
    assert report.max_cumulative_loss == D(840)  # COH-01 : 20 % du point zéro de 4 200 (définition du moteur)
    check = report.validation()
    assert (check.first_week, check.last_week, check.complete) == (5, 12, True)
    assert r2(check.orders) == D("184.62") and check.passed


def test_projection_prudent_and_development_recovery() -> None:
    prudent, _, dev = fc.bp_scenarios()
    rp = fc.north_star(fc.project_north_star(prudent, opening_week=5))
    assert rp.recovery_week == 34
    assert r2(rp.rows[4].net_contribution) == D("12.31")
    rd = fc.north_star(fc.project_north_star(dev, opening_week=5))
    assert rd.recovery_week == 5


def test_pilot_pace_triggers_global_stoploss() -> None:
    """Au rythme du jalon BP §1 (≈ 15 commandes/mois, CAC 8), le gel global tombe en semaine 13 (seuil 840 CHF)."""
    pilot = ScenarioAssumptions("pilote_validation", 15, D(8))
    report = fc.north_star(fc.project_north_star(pilot, weeks=104, opening_week=5))
    assert report.first_global_stoploss_week == 13  # COH-01 : définition unique (point zéro 4 200 => 840 CHF)
    assert report.rows[11].cumulative > D(-840) >= report.rows[12].cumulative
    # L'ancienne approximation (8 000 CHF littéraux, seuil 1 600) gelait en semaine 28 : explicite seulement.
    legacy = fc.north_star(fc.project_north_star(pilot, weeks=104, opening_week=5), capital_engaged=fc.BP_CAPITAL_ENGAGED)
    assert legacy.first_global_stoploss_week == 28
    assert r2(report.rows[10].net_contribution) == D("-53.07")
    assert report.frozen and report.recovery_week is None
    assert not report.validation().orders_ok  # 27,7 commandes sur 8 semaines < 30


def test_global_stoploss_is_sticky_and_threshold_inclusive() -> None:
    weeks = [nsw(1, fixed="600"), nsw(2, fixed="240"), nsw(3, net="5000", cost="2000")]
    report = fc.north_star(weeks)
    assert [r.cumulative for r in report.rows] == [D(-600), D(-840), D(2160)]
    assert [r.global_stoploss for r in report.rows] == [False, True, True]  # −840 = seuil ⇒ gel
    assert report.first_global_stoploss_week == 2 and report.frozen
    assert report.recovery_week == 3
    lenient = fc.north_star(weeks, capital_engaged=D(10000), global_stoploss_pct=D("0.25"))
    assert not lenient.frozen and lenient.max_cumulative_loss == D(2500)


def test_return_week_with_negative_sales_is_accepted() -> None:
    """Une semaine de retours peut avoir des ventes nettes et un coût historique négatifs."""
    report = fc.north_star([nsw(1, net="-95", cost="-66.5", log="5")])
    assert report.cumulative == D("-33.5")


def test_validation_window_and_incomplete_data() -> None:
    weeks = [nsw(n, net="400", cost="280", acq="10", orders="4") for n in range(1, 8)]
    report = fc.north_star(weeks)
    check = report.validation(opening_week=2)
    assert (check.first_week, check.last_week) == (2, 9)
    assert not check.complete
    assert check.orders == D(24) and not check.orders_ok
    assert check.contribution_after_acquisition == D(6 * 110) and check.contribution_ok
    assert not check.passed
    assert report.validation(opening_week=1, days=14, target_orders=8).passed
    with pytest.raises(ForecastError):
        report.validation(opening_week=0)
    with pytest.raises(ForecastError):
        report.validation(days=6)


@pytest.mark.parametrize(
    "factory",
    [
        lambda: nsw(0),
        lambda: fc.NorthStarWeek(1.0, D(0), D(0), D(0), D(0), D(0), D(0), D(0), D(0)),  # type: ignore[arg-type]
        lambda: nsw(1, fixed="-1"),
        lambda: nsw(1, acq="-1"),
        lambda: nsw(1, orders="-1"),
        lambda: fc.NorthStarWeek(1, D(0), 10.5, D(0), D(0), D(0), D(0), D(0), D(0)),  # type: ignore[arg-type]
        lambda: fc.north_star([]),
        lambda: fc.north_star([nsw(2)]),
        lambda: fc.north_star([nsw(1), nsw(3)]),
        lambda: fc.north_star(["semaine"]),  # type: ignore[list-item]
        lambda: fc.north_star([nsw(1)], capital_engaged=D(0)),
        lambda: fc.north_star([nsw(1)], global_stoploss_pct=D("1.5")),
        lambda: fc.project_north_star(fc.bp_scenarios()[1], weeks=0),
        lambda: fc.project_north_star(fc.bp_scenarios()[1], opening_week=0),
        lambda: fc.project_north_star(fc.bp_scenarios()[1], payment_pct=D(-1)),
        # 22 % de contribution + 76 % de coût produit ne laissent rien pour paiement et SAV
        lambda: fc.project_north_star(fc.bp_scenarios()[1], product_cost_ratio=D("0.76")),
    ],
)
def test_north_star_validation(factory) -> None:
    with pytest.raises(ForecastError):
        factory()


def test_stock_need_and_time_valuation_reject_non_int_orders() -> None:
    with pytest.raises(ForecastError):
        fc.stock_need(100.0)  # type: ignore[arg-type]
    with pytest.raises(ForecastError):
        fc.time_valuation(D(1), supervision_hours_per_week=D(1), orders_per_month=True)  # type: ignore[arg-type]


# ---------------------------------------------------------------------------
# Vérification globale
# ---------------------------------------------------------------------------


def test_verify_against_bp_statuses() -> None:
    checks = fc.verify_against_bp()
    refs = [c.ref for c in checks]
    assert len(refs) == len(set(refs)) == 38
    summary = fc.summarize_checks(checks)
    assert summary[CheckStatus.OK] == 35
    assert summary[CheckStatus.ROUNDING_GAP] == 2
    assert summary[CheckStatus.GAP] == 1
    by_ref = {c.ref: c for c in checks}
    assert by_ref["T1"].status is CheckStatus.ROUNDING_GAP and by_ref["T1"].recomputed == 30
    assert by_ref["T2"].status is CheckStatus.ROUNDING_GAP and by_ref["T2"].recomputed == 180
    assert by_ref["S3.2"].status is CheckStatus.GAP and by_ref["S3.2"].delta == D(-1)
    assert all(c.comment for c in checks if c.status is not CheckStatus.OK)


def test_initial_budget_and_fixed_costs_constants() -> None:
    assert sum(a for _, a in fc.BP_INITIAL_BUDGET) == D(8000)
    assert sum(a for _, a in fc.BP_MONTHLY_FIXED) == D(400)


# ---------------------------------------------------------------------------
# Classeur modele_financier.xlsx (formules vivantes recalculées par LibreOffice)
# ---------------------------------------------------------------------------


def _load_generator() -> ModuleType:
    """Charge le générateur sans écrire de __pycache__ dans docs/."""
    path = FINANCE_DIR / "generer_classeurs.py"
    spec = importlib.util.spec_from_file_location("generer_classeurs", path)
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    sys.modules.setdefault("generer_classeurs", module)
    previous = sys.dont_write_bytecode
    sys.dont_write_bytecode = True
    try:
        spec.loader.exec_module(module)
    finally:
        sys.dont_write_bytecode = previous
    return module


@pytest.fixture(scope="module")
def generator() -> ModuleType:
    return _load_generator()


@pytest.fixture(scope="module")
def model_values(generator: ModuleType, tmp_path_factory: pytest.TempPathFactory):
    from openpyxl import load_workbook

    tmp = tmp_path_factory.mktemp("modele")
    src = tmp / generator.MODEL_FILE
    generator.build_financial_model().save(src)
    try:
        recalculated = generator.recalculate(src, tmp / "recalc")
    except generator.RecalcError as exc:
        pytest.skip(f"LibreOffice Calc indisponible : {exc}")
    assert generator.formula_errors(recalculated) == []
    wb = load_workbook(recalculated, data_only=True)
    yield wb
    wb.close()


def test_model_workbook_structure(generator: ModuleType, tmp_path: Path) -> None:
    from openpyxl import load_workbook

    path = tmp_path / "m.xlsx"
    generator.build_financial_model().save(path)
    wb = load_workbook(path)
    assert wb.sheetnames == ["Hypothèses", "Étoile polaire", "Budget initial", "Charges mensuelles", "Scénarios",
                             "Seuil & sensibilité", "Prix plancher", "Stock & BFR"]
    formulas = sum(
        1 for ws in wb.worksheets for row in ws.iter_rows() for c in row
        if isinstance(c.value, str) and c.value.startswith("=")
    )
    assert formulas > 250
    # les résultats clés sont des formules, pas des valeurs figées
    assert wb["Scénarios"]["C15"].value == "=C11-C13-C14"
    assert wb["Prix plancher"]["B17"].value.startswith("=IF(B16>0,B15/B16")
    assert wb["Hypothèses"]["B9"].fill.fgColor.rgb.endswith("FFF2CC")
    for required in ("TVA", "Panier_TTC", "Panier_HT", "Taux_contribution", "Charges_fixes", "Remuneration",
                     "Capital_engage", "Perte_max", "Reserve_cash", "Scenario_choisi", "Semaine_ouverture"):
        assert required in wb.defined_names
    ns = wb["Étoile polaire"]
    assert ns["K45"].value == "=I45-J45"
    assert ns["L45"].value == "=SUM($K$41:K45)"
    assert ns["V45"].value == '=IF(COUNT(M45:T45)=0,"",U45-T45)'
    assert ns["M45"].fill.fgColor.rgb.endswith("FFF2CC") and ns["M45"].value is None


def test_model_scenarios_match_bp(model_values) -> None:
    ws = model_values["Scénarios"]
    expected = {
        "B": (3800, 3515, 773, 320, 53),
        "C": (9500, 8788, 1933, 600, 933),
        "D": (19000, 17576, 3867, 1000, 2467),
    }
    for col, (ttc, ht, contrib, acq, result) in expected.items():
        assert round(ws[f"{col}8"].value) == ttc
        assert round(ws[f"{col}9"].value) == ht
        assert round(ws[f"{col}11"].value) == contrib
        assert round(ws[f"{col}13"].value) == acq
        assert round(ws[f"{col}15"].value) == result
    assert ws["B41"].value == "OK" and ws["C41"].value == "OK"
    assert ws["D41"].value == "ÉCART -1 CHF"
    assert ws["C23"].value.startswith("OUI")
    assert ws["B23"].value == "non"


def test_model_matches_python_engine(model_values) -> None:
    ws = model_values["Scénarios"]
    for col, res in zip("BCD", fc.run_bp_scenarios().values()):
        assert ws[f"{col}9"].value == pytest.approx(float(res.revenue_ht), abs=1e-6)
        assert ws[f"{col}15"].value == pytest.approx(float(res.result_before_remuneration), abs=1e-6)
        assert ws[f"{col}19"].value == pytest.approx(float(res.contribution_per_order_after_cac), abs=1e-9)


def test_model_break_even_and_sensitivity(model_values) -> None:
    ws = model_values["Seuil & sensibilité"]
    assert ws["B11"].value == 30 and ws["C11"].value == 31 and ws["D11"].value == 31
    assert ws["B13"].value == 180 and ws["C13"].value == 181
    assert ws["E11"].value == "OK méthode BP — exact : 30"
    assert ws["D17"].value == 56 and ws["D18"].value == 93
    assert ws["F17"].value.startswith("NON") and ws["F18"].value.startswith("NON")
    # grille : ligne 22 %, colonne CAC 6 (E) = 30 ; CAC 15 (I) = 93 ; 12 % / CAC 15 non rentable
    assert ws["E29"].value == 30
    assert ws["I29"].value == 93
    assert ws["I23"].value == "non rentable"
    rates = [D(r) for r in ("0.12", "0.14", "0.15", "0.16", "0.18", "0.20", "0.22", "0.24", "0.26")]
    cacs = [D(c) for c in (0, 4, 5, 6, 8, 10, 12, 15)]
    grid = fc.sensitivity_grid(rates, cacs)
    assert len(grid) == 72
    for idx, cell in enumerate(grid):
        row = 23 + idx // 8
        col = "BCDEFGHI"[idx % 8]
        expected = cell.orders_required if cell.orders_required is not None else "non rentable"
        assert ws[f"{col}{row}"].value == expected, (cell.contribution_rate, cell.cac)


def test_model_floor_price(model_values) -> None:
    ws = model_values["Prix plancher"]
    assert ws["B18"].value == pytest.approx(208.79)
    assert ws["C18"].value == pytest.approx(207.28)
    assert ws["C7"].value == pytest.approx(151.34)
    assert ws["B21"].value == pytest.approx(209.90)
    assert ws["B22"].value == pytest.approx(208.90)
    assert ws["B25"].value == "OK — conforme BP" and ws["C25"].value == "OK — conforme BP"
    for row in (30, 31, 32, 33):
        assert ws[f"E{row}"].value == "OK"
    assert ws["B34"].value.startswith("NON")  # 16,56 % < 20 %
    assert ws["C34"].value == "oui"


def test_model_budget_stock_and_bfr(model_values) -> None:
    budget = model_values["Budget initial"]
    assert budget["B12"].value == 8000 and budget["B14"].value == "OK — conforme BP"
    assert budget["C25"].value == 3000 and budget["C26"].value == 750
    fixed = model_values["Charges mensuelles"]
    assert fixed["B7"].value == 400 and fixed["B9"].value == "OK — conforme BP"
    stock = model_values["Stock & BFR"]
    assert round(stock["C9"].value) == 6152
    assert stock["D14"].value == "OK — conforme BP"
    assert stock["C12"].value == pytest.approx(float(fc.stock_need(100).coverage_days), abs=1e-9)
    wc = fc.working_capital_need(100)
    assert stock["C23"].value == pytest.approx(float(wc.total), abs=1e-6)
    assert stock["C25"].value > 0  # le BFR central dépasse stock + réserve


def test_model_pub_stoploss_row(model_values) -> None:
    ws = model_values["Scénarios"]
    assert [ws[f"{c}26"].value for c in "BCD"] == ["non", "non", "non"]


def test_model_north_star_projection_matches_engine(model_values) -> None:
    ws = model_values["Étoile polaire"]
    assert ws["E6"].value == "Central" and ws["E7"].value == 100 and ws["E8"].value == 6
    assert ws["E13"].value == -840 and ws["E14"].value == 1600  # COH-01 : seuil global 840, réserve cash 1 600
    assert ws["R12"].value == "OK — cohérent"
    central = fc.bp_scenarios()[1]
    report = fc.north_star(fc.project_north_star(central, opening_week=5))
    for row in report.rows:
        r = 40 + row.week.week
        assert ws[f"A{r}"].value == row.week.week
        assert ws[f"B{r}"].value == pytest.approx(float(row.week.orders), abs=1e-9)
        assert ws[f"F{r}"].value == pytest.approx(float(row.week.logistics), abs=1e-6)
        assert ws[f"K{r}"].value == pytest.approx(float(row.net_contribution), abs=1e-6)
        assert ws[f"L{r}"].value == pytest.approx(float(row.cumulative), abs=1e-6)
        assert ws[f"V{r}"].value in (None, "")
    assert ws["E18"].value == pytest.approx(float(report.cumulative), abs=1e-6)
    assert ws["F18"].value == "pas de saisie" and ws["F19"].value == 0
    assert ws["E20"].value == pytest.approx(-369.23, abs=0.005) and ws["E21"].value == 4
    assert ws["E22"].value == report.recovery_week == 6
    assert ws["E23"].value == "non"
    assert ws["E24"].value == pytest.approx(float(fc.run_bp_scenarios()["central"].result_before_remuneration))
    assert ws["G24"].value.startswith("OK")
    check = report.validation()
    assert ws["E25"].value == pytest.approx(float(check.orders), abs=1e-9)
    assert ws["E26"].value == "oui" and ws["E28"].value == "oui"
    assert ws["E27"].value == pytest.approx(float(check.contribution_after_acquisition), abs=1e-6)
    assert ws["G25"].value.startswith("Semaines 5 à 12 (56 jours")
    assert [ws[f"H{r}"].value for r in (33, 35, 36, 37)] == [750, 1600, -840, "60 jours"]


#: Saisie FICTIVE du réel (semaines 1 à 10) : colonnes M à T de la feuille Étoile polaire.
FICTIVE_ACTUALS: list[tuple[int, str, str, str, str, str, str, str]] = [
    # commandes, ventes nettes HT, coût historique, paiement, logistique, SAV, acquisition, charges fixes
    (0, "0", "0", "0", "0", "0", "0", "400"),
    (0, "0", "0", "0", "0", "0", "0", "0"),
    (0, "0", "0", "0", "0", "0", "0", "0"),
    (0, "0", "0", "0", "0", "0", "0", "0"),
    (0, "0", "0", "0", "0", "0", "0", "400"),
    (3, "258.74", "181.10", "7.89", "10.50", "3", "0", "0"),
    (5, "450.05", "320.40", "13.66", "17.50", "5", "150", "0"),
    (6, "519.33", "362.80", "15.84", "21", "6", "200", "0"),
    (8, "713.41", "500.10", "21.68", "28", "8", "150", "400"),
    (9, "779.46", "548.70", "23.77", "31.50", "9", "0", "0"),
]


def test_model_north_star_actuals_match_engine(generator: ModuleType, tmp_path: Path) -> None:
    """Les formules du réel (colonnes U à Y, synthèse) reproduisent pokeshop.forecast.north_star."""
    from openpyxl import load_workbook

    wb = generator.build_financial_model()
    ws = wb["Étoile polaire"]
    weeks = []
    for n, (orders, *amounts) in enumerate(FICTIVE_ACTUALS, start=1):
        r = 40 + n
        ws[f"M{r}"] = orders
        for col, amount in zip("NOPQRST", amounts):
            ws[f"{col}{r}"] = float(amount)
        net, cost, pay, log, sav, acq, fixed = (D(a) for a in amounts)
        weeks.append(fc.NorthStarWeek(n, D(orders), net, cost, pay, log, sav, acq, fixed))
    src = tmp_path / "ns.xlsx"
    wb.save(src)
    try:
        recalculated = generator.recalculate(src, tmp_path / "recalc")
    except generator.RecalcError as exc:
        pytest.skip(f"LibreOffice Calc indisponible : {exc}")
    assert generator.formula_errors(recalculated) == []
    values = load_workbook(recalculated, data_only=True)["Étoile polaire"]
    report = fc.north_star(weeks)
    for row in report.rows:
        r = 40 + row.week.week
        assert values[f"V{r}"].value == pytest.approx(float(row.net_contribution), abs=1e-6)
        assert values[f"W{r}"].value == pytest.approx(float(row.cumulative), abs=1e-6)
        expected = "GEL GLOBAL" if row.global_stoploss else ("semaine négative" if row.net_contribution < 0 else "OK")
        assert values[f"Y{r}"].value == expected
    assert values["W51"].value in (None, "")
    assert values["F18"].value == pytest.approx(float(report.cumulative), abs=1e-6)
    assert values["F19"].value == 10
    assert values["F20"].value == pytest.approx(float(report.lowest.cumulative), abs=1e-6)
    assert values["F21"].value == report.lowest.week.week
    assert values["F22"].value == "pas encore"
    # COH-01 : au seuil unique de 840 CHF, l'exemple FICTIF gèle en semaine 8 (l'ancien seuil 1 600 ne gelait pas).
    assert report.first_global_stoploss_week == 8
    assert values["F23"].value == "OUI — semaine 8 : TOUT GELER"
    check = report.validation()
    assert not check.complete  # 6 semaines saisies sur 8 : jalon provisoire
    assert values["F25"].value == float(check.orders) == 31
    assert values["F26"].value == "oui"
    assert values["F27"].value == pytest.approx(float(check.contribution_after_acquisition), abs=1e-6)
    assert values["F28"].value == ("oui" if check.contribution_ok else "NON")


def test_model_north_star_global_freeze(generator: ModuleType, tmp_path: Path) -> None:
    """Une perte cumulée de 840 CHF (20 % du point zéro de 4 200) déclenche le gel global, qui reste actif."""
    from openpyxl import load_workbook

    wb = generator.build_financial_model()
    ws = wb["Étoile polaire"]
    hyp = wb["Hypothèses"]
    assert hyp["B46"].value == "=Capital_engage*StopLoss_global_pct" and "point zéro" in hyp["A44"].value
    for r, fixed in ((41, 600), (42, 300), (43, 0)):
        ws[f"T{r}"] = fixed
    ws["N43"] = 2000
    src = tmp_path / "gel.xlsx"
    wb.save(src)
    try:
        recalculated = generator.recalculate(src, tmp_path / "recalc")
    except generator.RecalcError as exc:
        pytest.skip(f"LibreOffice Calc indisponible : {exc}")
    values = load_workbook(recalculated, data_only=True)["Étoile polaire"]
    assert [values[f"Y{r}"].value for r in (41, 42, 43)] == ["semaine négative", "GEL GLOBAL", "GEL GLOBAL"]
    assert values["F23"].value == "OUI — semaine 2 : TOUT GELER"
    assert values["F22"].value == 3


def test_delivered_model_file_is_recalculated() -> None:
    """Le fichier livré contient formules ET valeurs (recalcul LibreOffice)."""
    from openpyxl import load_workbook

    path = FINANCE_DIR / "modele_financier.xlsx"
    assert path.exists()
    values = load_workbook(path, data_only=True)
    formulas = load_workbook(path)
    assert formulas["Scénarios"]["C15"].value.startswith("=")
    assert round(values["Scénarios"]["C15"].value) == 933
    assert values["Seuil & sensibilité"]["C11"].value == 31
    assert formulas["Étoile polaire"]["L92"].value == "=SUM($K$41:K92)"
    assert values["Étoile polaire"]["L92"].value == pytest.approx(9969.91, abs=0.005)
    assert values["Étoile polaire"]["F18"].value == "pas de saisie"
