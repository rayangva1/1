"""Tests du prévisionnel de trésorerie 13 semaines (pokeshop.treasury) et du classeur associé."""

from __future__ import annotations

import importlib.util
import sys
from datetime import date, timedelta
from decimal import Decimal
from pathlib import Path
from types import ModuleType

import pytest

from pokeshop.treasury import (
    HORIZON_WEEKS,
    INITIAL_BUDGET_TOTAL,
    AlertCode,
    CashMovement,
    Flow,
    PreorderFulfilment,
    PspTerms,
    SaleBatch,
    TreasuryError,
    TreasuryPlan,
    WeeklyInput,
    build_forecast,
    initial_budget,
    initial_budget_movements,
    plan_from_weekly_inputs,
    vat_payable,
)

D = Decimal
START = date(2026, 10, 5)  # lundi suivant la date de référence du BP
ROOT = Path(__file__).resolve().parents[1]
FINANCE_DIR = ROOT / "docs" / "03-finance"


def day(week: int, offset: int = 0) -> date:
    """Jour ``offset`` (0 = lundi) de la semaine ``week`` (1-indexée)."""
    return START + timedelta(days=7 * (week - 1) + offset)


def plan(**kwargs) -> TreasuryPlan:
    base = {"start": START, "opening_balance": D(5000), "minimum_reserve": D(1000)}
    base.update(kwargs)
    return TreasuryPlan(**base)


# ---------------------------------------------------------------------------
# PSP et versements différés
# ---------------------------------------------------------------------------


class TestPsp:
    def test_fees_follow_bp_example(self) -> None:
        psp = PspTerms()
        assert psp.fees(D("199.90"), 1) == D("5.29750")
        assert psp.fees(D(950), 10) == D("26.750")

    def test_payout_day(self) -> None:
        assert PspTerms(payout_delay_days=3).payout_day(START) == START + timedelta(days=3)

    @pytest.mark.parametrize(
        "kwargs",
        [{"pct_fee": D(1)}, {"pct_fee": D("-0.01")}, {"fixed_fee": D(-1)}, {"payout_delay_days": -1},
         {"pct_fee": 0.025}, {"payout_delay_days": True}],
    )
    def test_validation(self, kwargs: dict) -> None:
        with pytest.raises(TreasuryError):
            PspTerms(**kwargs)


def test_payout_is_delayed_by_one_week() -> None:
    f = build_forecast(plan(sales=(SaleBatch(day(1), D(1000), 10),)))
    w1, w2 = f.weeks[0], f.weeks[1]
    assert w1.gross_sales_ttc == D(1000)
    assert w1.psp_fees == D("28.000")
    assert w1.flow(Flow.PSP_PAYOUT) == 0
    assert w1.closing_balance == D(5000)
    assert w2.flow(Flow.PSP_PAYOUT) == D("972.000")
    assert w2.closing_balance == D("5972.000")
    assert f.psp_in_transit_end == 0


def test_payout_delay_crossing_weeks_and_zero_delay() -> None:
    late = build_forecast(plan(psp=PspTerms(payout_delay_days=10), sales=(SaleBatch(day(1, 4), D(100), 1),)))
    # vendredi + 10 jours = lundi de la semaine 3
    assert late.weeks[2].flow(Flow.PSP_PAYOUT) == D("97.200")
    assert late.weeks[1].flow(Flow.PSP_PAYOUT) == 0
    same = build_forecast(plan(psp=PspTerms(payout_delay_days=0), sales=(SaleBatch(day(1), D(100), 1),)))
    assert same.weeks[0].flow(Flow.PSP_PAYOUT) == D("97.200")


def test_payout_after_horizon_is_in_transit() -> None:
    f = build_forecast(plan(sales=(SaleBatch(day(13, 2), D(500), 5),)))
    assert f.weeks[-1].gross_sales_ttc == D(500)
    assert sum(w.flow(Flow.PSP_PAYOUT) for w in f.weeks) == 0
    assert f.psp_in_transit_end == D("486.000")


def test_sale_before_start_contributes_only_its_payout() -> None:
    before = SaleBatch(START - timedelta(days=2), D(300), 3)
    f = build_forecast(plan(sales=(before,)))
    assert f.weeks[0].flow(Flow.PSP_PAYOUT) == D("291.600")
    assert all(w.gross_sales_ttc == 0 for w in f.weeks)
    already_paid = SaleBatch(START - timedelta(days=20), D(300), 3)
    assert build_forecast(plan(sales=(already_paid,))).weeks[0].flow(Flow.PSP_PAYOUT) == 0


def test_preorder_before_start_is_rejected() -> None:
    with pytest.raises(TreasuryError, match="opening_preorder_reserve"):
        build_forecast(plan(sales=(SaleBatch(START - timedelta(days=1), D(80), 1, preorder=True),)))


def test_sale_after_horizon_is_ignored_and_counted() -> None:
    f = build_forecast(plan(sales=(SaleBatch(day(14), D(100), 1),)))
    assert f.ignored_after_horizon == 1
    assert f.psp_in_transit_end == 0


# ---------------------------------------------------------------------------
# Alertes : réserve minimale, solde négatif, précommandes
# ---------------------------------------------------------------------------


def test_minimum_reserve_violation_raises_alert() -> None:
    f = build_forecast(plan(movements=(CashMovement(day(3), Flow.PURCHASE_COMMITTED, D(4500), "réassort"),)))
    w3 = f.weeks[2]
    assert w3.closing_balance == D(500)
    assert w3.status is AlertCode.BELOW_MINIMUM_RESERVE
    assert [a.week for a in f.alerts] == list(range(3, 14))
    first = f.alerts[0]
    assert first.code is AlertCode.BELOW_MINIMUM_RESERVE
    assert first.shortfall == D(500)
    assert first.week_start == day(3)
    assert f.has_alerts
    assert f.lowest_week.index == 3


def test_balance_exactly_at_reserve_is_ok() -> None:
    f = build_forecast(plan(movements=(CashMovement(day(1), Flow.ADVERTISING, D(4000)),)))
    assert f.weeks[0].closing_balance == D(1000)
    assert f.weeks[0].status is None
    assert not f.has_alerts


def test_negative_balance_is_most_severe() -> None:
    f = build_forecast(plan(movements=(CashMovement(day(2), Flow.PURCHASE_PLANNED, D(6000)),)))
    assert f.weeks[1].status is AlertCode.NEGATIVE_BALANCE
    assert f.alerts[0].shortfall == D(2000)  # 1 000 de découvert + 1 000 de réserve


def test_preorder_reserve_not_available_for_purchases() -> None:
    f = build_forecast(
        plan(
            opening_balance=D(1500),
            sales=(SaleBatch(day(1), D(800), 10, preorder=True),),
            movements=(CashMovement(day(2), Flow.PURCHASE_COMMITTED, D(700)),),
        )
    )
    w1, w2 = f.weeks[0], f.weeks[1]
    assert w1.preorder_reserve == D(800)
    assert w1.available_for_purchases == D(1500) - D(800) - D(1000)
    assert w1.status is AlertCode.PREORDER_RESERVE_UNCOVERED
    # semaine 2 : versement 800 − (2,5 % × 800 + 10 × 0,30) = 777 ; achat 700 ⇒ solde 1 577
    assert w2.closing_balance == D("1577.000")
    assert w2.status is AlertCode.PREORDER_RESERVE_UNCOVERED
    assert w2.preorder_sales_ttc == 0


def test_fulfilment_releases_preorder_reserve() -> None:
    f = build_forecast(
        plan(
            sales=(SaleBatch(day(1), D(600), 6, preorder=True),),
            fulfilments=(PreorderFulfilment(day(4), D(400)),),
        )
    )
    assert [w.preorder_reserve for w in f.weeks[:5]] == [D(600), D(600), D(600), D(200), D(200)]


def test_preorder_refund_releases_reserve_and_costs_cash() -> None:
    f = build_forecast(
        plan(
            sales=(SaleBatch(day(1), D(600), 6, preorder=True),),
            movements=(CashMovement(day(3), Flow.PREORDER_REFUND, D(100), "allocation réduite"),),
        )
    )
    w3 = f.weeks[2]
    assert w3.preorder_reserve == D(500)
    assert w3.flow(Flow.PREORDER_REFUND) == D(100)
    assert w3.total_outflows == D(100)
    assert w3.closing_balance == f.weeks[1].closing_balance - D(100)


def test_reserve_never_negative() -> None:
    f = build_forecast(plan(opening_preorder_reserve=D(50), fulfilments=(PreorderFulfilment(day(1), D(80)),)))
    assert f.weeks[0].preorder_reserve == 0


def test_customer_refund_reduces_balance() -> None:
    f = build_forecast(
        plan(
            sales=(SaleBatch(day(1), D(190), 2),),
            movements=(CashMovement(day(2), Flow.REFUND, D(95), "retour produit défectueux"),),
        )
    )
    w2 = f.weeks[1]
    assert w2.flow(Flow.REFUND) == D(95)
    # frais PSP non restitués : 190 − 5,35 versés, puis 95 remboursés
    assert w2.closing_balance == D(5000) + D("184.650") - D(95)
    assert w2.preorder_reserve == 0


def test_all_flows_aggregate_by_week() -> None:
    movements = tuple(
        CashMovement(day(2, 1), flow, D(10), flow.value) for flow in Flow if flow is not Flow.PSP_PAYOUT
    )
    f = build_forecast(plan(movements=movements))
    w2 = f.weeks[1]
    assert w2.total_inflows == D(10)  # autres encaissements
    assert w2.total_outflows == D(10) * (len(Flow) - 2)
    assert w2.flow(Flow.PURCHASE_COMMITTED) == w2.flow(Flow.PURCHASE_PLANNED) == D(10)
    assert w2.closing_balance == D(5000) + D(10) - D(100)


def test_movement_rules() -> None:
    with pytest.raises(TreasuryError, match="antérieur"):
        build_forecast(plan(movements=(CashMovement(START - timedelta(days=1), Flow.VAT, D(10)),)))
    with pytest.raises(TreasuryError, match="antérieure"):
        build_forecast(plan(fulfilments=(PreorderFulfilment(START - timedelta(days=1), D(10)),)))
    f = build_forecast(
        plan(
            movements=(CashMovement(day(14), Flow.VAT, D(10)),),
            fulfilments=(PreorderFulfilment(day(15), D(10)),),
        )
    )
    assert f.ignored_after_horizon == 2


@pytest.mark.parametrize(
    "factory",
    [
        lambda: CashMovement(START, Flow.VAT, D(-1)),
        lambda: CashMovement(START, Flow.VAT, 10.0),
        lambda: CashMovement(START, "tva", D(1)),  # type: ignore[arg-type]
        lambda: SaleBatch(START, D(10), 0),
        lambda: SaleBatch(START, D(-10), 1),
        lambda: SaleBatch(START, D(10), -1),
        lambda: PreorderFulfilment(START, D(-1)),
        lambda: TreasuryPlan(START, D(0), D(-1)),
        lambda: TreasuryPlan(START, D(0), D(0), weeks=0),
        lambda: TreasuryPlan(START, 1.5, D(0)),  # type: ignore[arg-type]
        lambda: TreasuryPlan(START, D(0), D(0), opening_preorder_reserve=D(-1)),
        lambda: TreasuryPlan(START, "abc", D(0)),  # type: ignore[arg-type]
        lambda: TreasuryPlan(START, D("NaN"), D(0)),
        lambda: TreasuryPlan(START, [1], D(0)),  # type: ignore[arg-type]
    ],
)
def test_input_validation(factory) -> None:
    with pytest.raises(TreasuryError):
        factory()


def test_plan_horizon_helpers() -> None:
    p = plan()
    assert p.weeks == HORIZON_WEEKS == 13
    assert p.end == START + timedelta(days=90)
    assert p.week_of(START) == 1 and p.week_of(p.end) == 13
    assert p.week_of(p.end + timedelta(days=1)) is None
    assert p.week_of(START - timedelta(days=1)) is None


def test_as_rows_flat_export() -> None:
    f = build_forecast(plan(sales=(SaleBatch(day(1), D("100.005"), 1),)))
    rows = f.as_rows()
    assert len(rows) == 13
    assert rows[0]["semaine"] == "1" and rows[0]["du"] == "2026-10-05" and rows[0]["au"] == "2026-10-11"
    assert rows[0]["ventes_ttc"] == "100.01"
    assert rows[1]["versements_psp"] == "97.20"
    assert rows[0]["alerte"] == "OK"
    assert set(Flow._value2member_map_) <= set(rows[0])
    assert f.closing_balance == f.weeks[-1].closing_balance


# ---------------------------------------------------------------------------
# Budget initial BP §3 et TVA
# ---------------------------------------------------------------------------


def test_initial_budget_matches_bp() -> None:
    lines = initial_budget()
    assert sum(line.amount for line in lines) == INITIAL_BUDGET_TOTAL == D(8000)
    reserve = [line for line in lines if line.flow is None]
    assert len(reserve) == 1 and reserve[0].amount == D(1600)


def test_initial_budget_spend_leaves_reserve_and_fixed_costs_eat_it() -> None:
    schedule = {"site": day(1), "admin": day(1), "da": day(2), "emballages": day(3), "stock": day(4),
                "acquisition": day(7)}
    movements = initial_budget_movements(schedule)
    assert sum(m.amount for m in movements) == D(6400)
    f = build_forecast(TreasuryPlan(START, D(8000), D(1600), movements=movements))
    assert f.closing_balance == D(1600)
    assert not f.has_alerts
    # les charges fixes des mois sans ventes ne sont pas dans le budget : elles entament la réserve
    fixed = tuple(CashMovement(day(w), Flow.FIXED_COSTS, D(400)) for w in (1, 5))
    g = build_forecast(TreasuryPlan(START, D(8000), D(1600), movements=movements + fixed))
    assert g.closing_balance == D(800)
    assert g.alerts[0].code is AlertCode.BELOW_MINIMUM_RESERVE


def test_initial_budget_movement_errors() -> None:
    with pytest.raises(TreasuryError, match="réserve"):
        initial_budget_movements({"reserve": START})
    with pytest.raises(TreasuryError, match="inconnue"):
        initial_budget_movements({"yacht": START})


def test_vat_payable() -> None:
    assert vat_payable(D("1081")) == D("81")
    assert vat_payable(D("1081"), deductible_input_vat=D(100)) == D(-19)
    assert vat_payable(D(100), vat_rate=D(0)) == 0
    with pytest.raises(TreasuryError):
        vat_payable(D(100), vat_rate=D(1))
    with pytest.raises(TreasuryError):
        vat_payable(D(-1))


# ---------------------------------------------------------------------------
# Saisie hebdomadaire (structure du classeur)
# ---------------------------------------------------------------------------


def test_plan_from_weekly_inputs_maps_every_field() -> None:
    week = WeeklyInput(
        sales_ttc=D(100), orders=1, preorder_sales_ttc=D(50), preorder_orders=1, psp_in_transit=D(5),
        other_inflows=D(6), purchases_committed=D(7), purchases_planned=D(8), vat=D(9), shipping=D(10),
        refunds=D(11), preorder_refunds=D(12), advertising=D(13), fixed_costs=D(14), setup_costs=D(15),
        other_outflows=D(16), preorders_fulfilled=D(17),
    )
    p = plan_from_weekly_inputs(start=START, opening_balance=D(1000), minimum_reserve=D(0), weeks=[week, WeeklyInput()],
                                payout_delay_weeks=0)
    assert p.weeks == 2
    assert len(p.sales) == 2 and sum(1 for s in p.sales if s.preorder) == 1
    assert len(p.movements) == 12
    f = build_forecast(p)
    w1 = f.weeks[0]
    net = D(150) - (D(150) * D("0.025") + D("0.60"))
    assert w1.flow(Flow.PSP_PAYOUT) == net + D(5)
    assert w1.flow(Flow.OTHER_INFLOW) == D(6)
    assert w1.total_outflows == sum(D(v) for v in range(7, 17))
    assert w1.preorder_reserve == D(50) - D(17) - D(12)


def test_plan_from_weekly_inputs_errors() -> None:
    with pytest.raises(TreasuryError):
        plan_from_weekly_inputs(start=START, opening_balance=D(0), minimum_reserve=D(0), weeks=[])
    with pytest.raises(TreasuryError):
        plan_from_weekly_inputs(start=START, opening_balance=D(0), minimum_reserve=D(0), weeks=[WeeklyInput()],
                                payout_delay_weeks=-1)
    with pytest.raises(TreasuryError):
        WeeklyInput(sales_ttc=D(10))
    with pytest.raises(TreasuryError):
        WeeklyInput(preorder_sales_ttc=D(10))
    with pytest.raises(TreasuryError):
        WeeklyInput(vat=D(-1))
    with pytest.raises(TreasuryError):
        WeeklyInput(orders=-1)


# ---------------------------------------------------------------------------
# Classeur tresorerie_13_semaines.xlsx
# ---------------------------------------------------------------------------

ALERT_TEXT = {
    None: "OK",
    AlertCode.NEGATIVE_BALANCE: "SOLDE NÉGATIF",
    AlertCode.BELOW_MINIMUM_RESERVE: "SOUS RÉSERVE MINI",
    AlertCode.PREORDER_RESERVE_UNCOVERED: "PRÉCOMMANDES NON COUVERTES",
}
WEEK_COLS = [chr(ord("C") + i) for i in range(13)]


def _load_generator() -> ModuleType:
    path = FINANCE_DIR / "generer_classeurs.py"
    spec = importlib.util.spec_from_file_location("generer_classeurs", path)
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    sys.modules.setdefault("generer_classeurs", module)
    spec.loader.exec_module(module)
    return module


@pytest.fixture(scope="module")
def generator() -> ModuleType:
    return _load_generator()


@pytest.fixture(scope="module")
def example_forecast(generator: ModuleType):
    params, weeks = generator.fictitious_example()
    p = plan_from_weekly_inputs(
        start=params.start, opening_balance=params.opening_balance, minimum_reserve=params.minimum_reserve,
        weeks=weeks, psp_pct_fee=params.psp_pct, psp_fixed_fee=params.psp_fixed,
        payout_delay_weeks=params.payout_delay_weeks, opening_preorder_reserve=params.opening_preorder_reserve,
    )
    return build_forecast(p)


@pytest.fixture(scope="module")
def treasury_values(generator: ModuleType, tmp_path_factory: pytest.TempPathFactory):
    from openpyxl import load_workbook

    tmp = tmp_path_factory.mktemp("tresorerie")
    src = tmp / generator.TREASURY_FILE
    generator.build_treasury_workbook().save(src)
    try:
        recalculated = generator.recalculate(src, tmp / "recalc")
    except generator.RecalcError as exc:
        pytest.skip(f"LibreOffice Calc indisponible : {exc}")
    assert generator.formula_errors(recalculated) == []
    wb = load_workbook(recalculated, data_only=True)
    yield wb
    wb.close()


def test_example_is_fictitious_and_shows_both_alert_types(example_forecast) -> None:
    codes = {a.code for a in example_forecast.alerts}
    assert codes == {AlertCode.BELOW_MINIMUM_RESERVE, AlertCode.PREORDER_RESERVE_UNCOVERED}
    assert example_forecast.weeks[-1].status is None
    assert example_forecast.psp_in_transit_end > 0


def test_workbook_sheets(generator: ModuleType, tmp_path: Path) -> None:
    from openpyxl import load_workbook

    path = tmp_path / "t.xlsx"
    generator.build_treasury_workbook().save(path)
    wb = load_workbook(path)
    assert wb.sheetnames == ["Mode d'emploi", "Exemple FICTIF", "À remplir"]
    ex = wb["Exemple FICTIF"]
    assert "FICTIF" in ex["A1"].value
    assert ex["D24"].value == "=IF(D$14-$B$10>=1,INDEX($C$22:$O$22,D$14-$B$10),0)"
    assert ex["D43"].value == "=D41+D42"
    blank = wb["À remplir"]
    assert blank["B6"].value is None and blank["C17"].value is None
    assert blank["C17"].fill.fgColor.rgb.endswith("FFF2CC")


def test_workbook_example_matches_engine(treasury_values, example_forecast) -> None:
    ws = treasury_values["Exemple FICTIF"]
    for week, col in zip(example_forecast.weeks, WEEK_COLS):
        assert ws[f"{col}43"].value == pytest.approx(float(week.closing_balance), abs=0.005), week.index
        assert ws[f"{col}24"].value == pytest.approx(float(week.flow(Flow.PSP_PAYOUT)), abs=0.005)
        assert ws[f"{col}21"].value == pytest.approx(float(week.psp_fees), abs=0.005)
        assert (ws[f"{col}46"].value or 0) == pytest.approx(float(week.preorder_reserve), abs=0.005)
        assert ws[f"{col}48"].value == pytest.approx(float(week.available_for_purchases), abs=0.005)
        assert ws[f"{col}49"].value == ALERT_TEXT[week.status], week.index
    assert ws["B52"].value == pytest.approx(float(example_forecast.lowest_week.closing_balance), abs=0.005)
    assert ws["B53"].value == example_forecast.lowest_week.index
    assert ws["B54"].value == len(example_forecast.alerts)
    assert ws["B55"].value == pytest.approx(float(example_forecast.psp_in_transit_end), abs=0.005)


def test_blank_template_is_clean(treasury_values) -> None:
    ws = treasury_values["À remplir"]
    assert all(ws[f"{c}49"].value == "OK" for c in WEEK_COLS)
    assert ws["B54"].value == 0
    assert ws["C15"].value in (None, "")


def test_delivered_treasury_file_is_recalculated() -> None:
    from openpyxl import load_workbook

    path = FINANCE_DIR / "tresorerie_13_semaines.xlsx"
    assert path.exists()
    values = load_workbook(path, data_only=True)["Exemple FICTIF"]
    formulas = load_workbook(path)["Exemple FICTIF"]
    assert formulas["O43"].value.startswith("=")
    assert isinstance(values["O43"].value, (int, float))
    assert values["B54"].value >= 1
