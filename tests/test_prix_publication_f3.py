"""Régressions du lot F3 (prix et publication) au niveau de l'API et du classeur de contrôle.

Chaque test reproduit un défaut confirmé par la revue adverse (identifiant dans la docstring) et
vérifie le comportement **fermé par défaut** corrigé. Jetons, montants et références FICTIFS.
"""

from __future__ import annotations

import json
from datetime import datetime, timedelta
from decimal import Decimal as D
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

import pytest
import jetons_roles as JR
from fastapi.testclient import TestClient
from pokeshop.api import API_TOKEN_HEADER, OWNER_TOKEN_HEADER, Services, create_app, guard_rules
from pokeshop.incidents import LogNotifier
from pokeshop.pricing import decide_price, small_product_min_order_required
from pokeshop.rules import load_rules
from pokeshop.settings import load_settings, sha256_hex
from pokeshop.stoploss import hash_owner_token

TZ = ZoneInfo("Europe/Zurich")
NOW = datetime(2026, 10, 4, 12, 0, tzinfo=TZ)
ROOT = Path(__file__).resolve().parents[1]
API_TOKEN = "FICTIF-jeton-api-0000000000000001"
OWNER_TOKEN = "FICTIF-jeton-proprietaire-tres-long-0001"
H = {API_TOKEN_HEADER: API_TOKEN}
HO = {API_TOKEN_HEADER: API_TOKEN, OWNER_TOKEN_HEADER: OWNER_TOKEN}
HBAD = {API_TOKEN_HEADER: API_TOKEN, OWNER_TOKEN_HEADER: "FICTIF-jeton-qui-n-est-pas-le-bon-0001"}

IDENTITY = {"gtin": "2000000001012", "language": "FR", "extension": "FICTIF_ALPHA", "format": "DISPLAY",
            "content": "36 BOOSTERS", "sealed": True}
LISTING: dict[str, Any] = {
    "product_key": "FICTIF-P1", "identity": IDENTITY, "public_sku": "DSP-FICTIF_ALPHA-FR",
    "description_html": "<p>Display de l'extension Extension Fictive Alpha, en français, neuf et scellé.</p>",
    "images": [{"url": "https://cdn.example.org/fictif.jpg", "alt": "Display face avant", "rights": "OWN_PHOTO"}],
    "stock_status": "stock_local", "content_text": "36 boosters", "fictif": True,
}


def reject_float(text: str) -> Any:
    raise AssertionError(f"nombre flottant dans une réponse JSON : {text}")


def body(resp: Any) -> Any:
    return json.loads(resp.content, parse_float=reject_float)


class Clock:
    """Horloge réglable du service (le même dossier d'état survit aux « redémarrages »)."""

    def __init__(self, now: datetime = NOW) -> None:
        self.now = now

    def __call__(self) -> datetime:
        return self.now


# Jetons nommés par rôle (matrice pokeshop.authz) ; le jeton commun H : lecture et aperçus seulement.
HF, HPHOTO, HOPS, HCAT, HSYNC, HDATA = JR.HF, JR.HPHOTO, JR.HOPS, JR.HCAT, JR.HSYNC, JR.HDATA


def approve(client: TestClient) -> None:
    """Fiche déposée par l'agent catalogue, validée par la propriétaire (registre, revue R3 R2-NEW-05)."""
    assert client.post("/catalog/items", headers=HCAT, json={"items": [{"product_id": "FICTIF-P1", "listing": LISTING}]}).status_code == 200
    assert client.post("/catalog/approvals", headers=HO, json={
        "product_id": "FICTIF-P1", "approved": True, "content_validated": True, "category_rule_validated": True,
        "reason": "fiche FICTIVE relue par la propriétaire"}).status_code == 201


def boot(clock: Clock | None = None, **extra: str) -> tuple[TestClient, Services]:
    env = {"POKESHOP_API_TOKEN_SHA256": sha256_hex(API_TOKEN), "POKESHOP_OWNER_TOKEN_SHA256": hash_owner_token(OWNER_TOKEN),
           "POKESHOP_AGENT_TOKENS_SHA256": JR.agent_tokens_env()}
    env.update(extra)
    svc = Services.build(load_settings(env), clock=clock or Clock(), notifier=LogNotifier())
    client = TestClient(create_app(services=svc))
    apport = {"movement_id": "FICTIF_APPORT", "at": (NOW - timedelta(days=30)).isoformat(), "kind": "CONTRIBUTION",
              "amount": "8000"}  # apport attesté par la propriétaire (SEC-06)  # fmt: skip
    assert client.post("/capital/movements", headers=HO, json=apport).status_code in (200, 201)
    return client, svc


def state_payload(at: datetime = NOW, **extra: Any) -> dict[str, Any]:
    base = {
        "as_of": at.isoformat(), "stock_budget_chf": "3000", "cash_available_chf": "5000", "ads_daily_cap_chf": "33",
        "net_worth": {"as_of": at.isoformat(), "cash_chf": "8000"},
    }
    base.update(extra)
    return base


def candidate(extension: str = "FICTIF_ALPHA") -> dict[str, Any]:
    return {
        "product_key": "FICTIF-P1", "extension": extension, "unit_cost_chf": "90.00", "sellable_qty": 0,
        "avg_daily_sales": "0.5", "lead_time_days": 7, "safety_stock": 1,
        "offer": {"supplier_id": "fictif_grossiste_a", "supplier_sku": "FICTIF-A-001", "availability_status": "IN_STOCK",
                  "available_qty": 24, "moq": 1, "carton_qty": 1, "source_ts": NOW.isoformat(), "raw_ref": "FICTIF:1"},
    }


def receive(client: TestClient, qty: int = 6, ref: str = "FICTIF-BL-0001") -> Any:
    return client.post("/stock/receive", headers=HOPS, json={"sku": "DSP-FICTIF_ALPHA-FR", "qty": qty, "ref": ref})


# ------------------------------------------------------------------ MOT-04 / SEC-03 : prix approuvés


def test_price_validation_is_no_longer_a_free_request_field() -> None:
    """MOT-04 / SEC-03 : ``price_validation(s)`` dans un corps de requête => 422 (champ inconnu)."""
    client, _ = boot()
    fake = {"price": "9.90", "validated_by": "proprietaire", "validated_at": NOW.isoformat()}
    preview = client.post("/publish/preview", headers=H, json={"listing": LISTING, "quote": {"cost_chf": "100.00"},
                                                               "price_validation": fake})
    assert preview.status_code == 422
    sync = client.post("/sync/run", headers=HSYNC, json={"supplier": "fictif_grossiste_a", "dry_run": False,
                                                     "source_path": "FICTIF_offres_grossiste_a.csv",
                                                     "price_validations": {"FICTIF-P1": fake}})
    assert sync.status_code == 422


def test_price_approval_requires_owner_token_and_is_audited() -> None:
    client, svc = boot()
    payload = {"product_key": "FICTIF-P1", "price": "149.90", "reason": "FICTIF alignement marché validé"}
    assert client.post("/pricing/approvals", headers=H, json=payload).status_code == 403
    assert client.post("/pricing/approvals", headers=HBAD, json=payload).status_code == 403
    assert svc.audit.events(action="pricing.approval.owner_token_refused")
    created = client.post("/pricing/approvals", headers=HO, json=payload)
    assert created.status_code == 201
    approval = body(created)["approval"]
    assert approval["approved_by"] == "propriétaire" and approval["price"] == "149.90"
    event = svc.audit.events(action="pricing.approval")[-1]
    assert event.actor == "propriétaire" and event.payload["approval_id"] == approval["approval_id"]
    listed = body(client.get("/pricing/approvals", headers=H))["approvals"]
    assert [a["approval_id"] for a in listed] == [approval["approval_id"]]
    assert client.post(f"/pricing/approvals/{approval['approval_id']}/revoke", headers=H,
                       json={"reason": "FICTIF erreur de saisie"}).status_code == 403  # jeton commun : jamais
    revoked = client.post(f"/pricing/approvals/{approval['approval_id']}/revoke", headers=HF,
                          json={"reason": "FICTIF erreur de saisie"})
    assert revoked.status_code == 200 and body(client.get("/pricing/approvals", headers=H))["approvals"] == []


def test_preview_uses_engine_approval_but_never_below_hard_floor() -> None:
    """SEC-03 (poc_api2 C) : coût 100, prix 9.90 « validé » => jamais publié sans exception C18."""
    client, _ = boot()
    receive(client)
    approve(client)
    quote = {"listing": LISTING, "quote": {"cost_chf": "100.00"}}
    base = body(client.post("/publish/preview", headers=H, json=quote))["plan"]
    assert base["price_source"] == "ENGINE" and base["price_chf"] == "154.90"
    client.post("/pricing/approvals", headers=HO, json={"product_key": "FICTIF-P1", "price": "9.90",
                                                         "reason": "FICTIF tentative sous le coût"})
    below = body(client.post("/publish/preview", headers=H, json=quote))["plan"]
    assert below["outcome"] == "NOT_SENT" and "APPROVED_PRICE_BELOW_FLOOR" in below["reviews"]
    client.post("/pricing/approvals", headers=HO, json={"product_key": "FICTIF-P1", "price": "149.90",
                                                         "reason": "FICTIF prix marché validé"})
    ok = body(client.post("/publish/preview", headers=H, json=quote))["plan"]
    assert ok["price_chf"] == "149.90" and ok["price_source"] == "HUMAN_VALIDATED" and ok["price_approval_id"]
    c18 = client.post("/pricing/approvals", headers=HO, json={
        "product_key": "FICTIF-P1", "price": "99.90", "reason": "FICTIF déstockage exceptionnel",
        "floor_exception_ref": "C18-FICTIF-2026-001"})
    assert c18.status_code == 201
    destock = body(client.post("/publish/preview", headers=H, json=quote))["plan"]
    assert destock["price_chf"] == "99.90" and destock["price_source"] == "HUMAN_VALIDATED"


def test_price_approvals_survive_restart() -> None:
    client, _ = boot()
    client.post("/pricing/approvals", headers=HO, json={"product_key": "FICTIF-P1", "price": "149.90",
                                                         "reason": "FICTIF prix marché validé"})
    _, svc2 = boot()
    assert svc2.price_approvals.active_for("FICTIF-P1", NOW) is not None and svc2.restore_errors == {}


# ------------------------------------------------------------------------------- SEC-18 : entrées


@pytest.mark.parametrize("value", ["NaN", "sNaN", "Infinity", "-Infinity", "1e999999", "1E+400000000", "-100",
                                   "0.0000001"])
def test_non_finite_extreme_or_negative_cost_is_422(value: str) -> None:
    client, _ = boot()
    resp = client.post("/pricing/quote", headers=H, json={"cost_chf": value})
    assert resp.status_code == 422, (value, resp.text)


@pytest.mark.parametrize(
    "change",
    [
        {"lines": [{"sku": "A", "qty": 1, "unit_price_ttc": "20.00", "unit_cost": "-50"}]},
        {"lines": [{"sku": "A", "qty": 1, "unit_price_ttc": "NaN", "unit_cost": "5"}]},
        {"lines": [{"sku": "A", "qty": 1, "unit_price_ttc": "0", "unit_cost": "5"}]},
        {"discount": {"kind": "PERCENT", "value": "1.5"}},
        {"discount": {"kind": "AMOUNT", "value": "-500"}},
        {"shipping_charged": "-1"},
        {"shipping_cost_actual": "Infinity"},
    ],
)
def test_basket_invalid_amounts_are_422_not_500(change: dict[str, Any]) -> None:
    client, _ = boot()
    payload = {"lines": [{"sku": "A", "qty": 1, "unit_price_ttc": "20.00", "unit_cost": "5"}], **change}
    resp = client.post("/pricing/basket", headers=H, json=payload)
    assert resp.status_code == 422, resp.text


def test_reorder_invalid_budget_is_422() -> None:
    client, _ = boot()
    for budget in ("NaN", "-1000000", "1e999999"):
        resp = client.post("/stock/reorder-proposal", headers=HF, json={"candidates": [], "budget_available": budget})
        assert resp.status_code == 422, budget


# ------------------------------------------------------------------- MOT-24 : réassort sans auto-déclaration


def test_reorder_proposal_refuses_self_declared_budget_and_cap_exceptions() -> None:
    client, _ = boot()
    req = {"candidates": [candidate()], "budget_available": "5000"}
    assert client.post("/stock/reorder-proposal", headers=HF, json=req).status_code == 409  # stop-loss inconnu
    assert client.post("/stoploss/state", headers=HPHOTO, json=state_payload()).status_code == 200
    assert client.post("/stock/reorder-proposal", headers=HF,
                       json={**req, "stock_budget_total": "100000"}).status_code == 422
    caps = {**req, "cap_exceptions": {"FICTIF_ALPHA": "1"}}
    assert client.post("/stock/reorder-proposal", headers=HF, json=caps).status_code == 403
    assert client.post("/stock/reorder-proposal", headers=HO, json=caps).status_code == 200


def test_reorder_uses_photo_exposure_caller_can_only_add() -> None:
    client, _ = boot()
    exposed = state_payload(extensions=[{"extension": "FICTIF_ALPHA", "stock_value_at_cost": "750"}])
    assert client.post("/stoploss/state", headers=HPHOTO, json=exposed).status_code == 200
    req = {"candidates": [candidate()], "budget_available": "5000", "extension_exposure": {"FICTIF_ALPHA": "0"}}
    data = body(client.post("/stock/reorder-proposal", headers=HF, json=req))
    assert data["proposal"]["lines"] == []  # 25 % × 3 000 = 750 déjà exposés : plafond atteint
    assert data["proposal"]["skipped"][0]["reason"] == "EXTENSION_CAP"


# ------------------------------------------------------------------------ E2E-17 : registre de stock


def test_stock_receive_feeds_the_registry_idempotently_and_survives_restart() -> None:
    client, svc = boot()
    first = receive(client, 6)
    assert first.status_code == 200 and body(first)["sellable"] == 6 and body(first)["replayed"] is False
    again = receive(client, 6)
    assert again.status_code == 200 and body(again)["replayed"] is True and body(again)["sellable"] == 6
    assert receive(client, 7).status_code == 409  # même bon de livraison, autre quantité
    assert body(client.post("/stock/sellable", headers=H, json={"sku": "DSP-FICTIF_ALPHA-FR"}))["sellable"] == 6
    assert svc.audit.events(action="stock.receive")
    client2, svc2 = boot()
    assert body(client2.post("/stock/sellable", headers=H, json={"sku": "DSP-FICTIF_ALPHA-FR"}))["sellable"] == 6
    assert svc2.restore_errors == {}
    assert client2.post("/stock/receive", headers=HOPS, json={"sku": "dsp minuscule", "qty": 1, "ref": "X-1"}).status_code == 422
    assert client2.post("/stock/receive", headers=H, json={"sku": "DSP-FICTIF_ALPHA-FR", "qty": 50,
                                                           "ref": "FICTIF-FANTOME"}).status_code == 403  # SEC-16


def test_preview_public_status_follows_registry() -> None:
    """E2E-12 : « stock local » déclaré sans stock => rupture (brouillon) ; après réception => stock local."""
    client, _ = boot()
    quote = {"listing": LISTING, "quote": {"cost_chf": "100.00"}}
    plan = body(client.post("/publish/preview", headers=H, json=quote))["plan"]
    assert plan["stock_status"] == "rupture" and plan["outcome"] == "SEND_DRAFT"
    receive(client)
    approve(client)
    plan = body(client.post("/publish/preview", headers=H, json=quote))["plan"]
    assert plan["stock_status"] == "stock_local" and plan["outcome"] == "SEND_ACTIVE"


# ------------------------------------------------------------- MOT-16 / E2E-04 : référence « dernier import »


def test_import_route_applies_and_persists_the_previous_import_baseline() -> None:
    clock = Clock(datetime(2026, 10, 3, 8, 0, tzinfo=TZ))
    client, svc = boot(clock)
    j1 = body(client.post("/imports/fictif_grossiste_a/run", headers=HDATA,
                          json={"source_path": "FICTIF_offres_grossiste_a_J-1.csv"}))
    assert j1["status"] in ("ACCEPTED", "PARTIAL") and j1["baseline_known"] is True
    clock.now = datetime(2026, 10, 4, 8, 0, tzinfo=TZ)
    client2, _ = boot(clock)  # redémarrage : la référence est relue
    incomplete = body(client2.post("/imports/fictif_grossiste_a/run", headers=HDATA,
                                   json={"source_path": "FICTIF_offres_grossiste_a_incomplet.csv"}))
    assert incomplete["status"] == "QUARANTINED" and "INCOMPLETE_IMPORT" in incomplete["reason_counts"]
    day = body(client2.post("/imports/fictif_grossiste_a/run", headers=HDATA,
                            json={"source_path": "FICTIF_offres_grossiste_a.csv"}))
    assert "PRICE_ANOMALY" in day["reason_counts"]  # ÷10 / ×10 vs J-1


# ------------------------------------------------------------------- MOT-10 / E2E-11 : petits produits


def test_small_product_rule_cannot_be_activated_by_unsigned_rules(tmp_path: Path) -> None:
    rules = load_rules()
    assert decide_price(D("3.79"), rules.pricing).recommended_price == D("23.90")  # règle inactive par défaut
    text = (ROOT / "config" / "pricing_rules.v1.yaml").read_text(encoding="utf-8")
    text = text.replace("small_product_min_order_ttc: null", 'small_product_min_order_ttc: "127.00"')
    text = text.replace("small_product_max_shipping_ttc: null", 'small_product_max_shipping_ttc: "0"')
    path = tmp_path / "pricing_rules.v1.yaml"
    path.write_text(text, encoding="utf-8")
    active = load_rules(path)
    assert small_product_min_order_required(active.pricing) == D("127.00")
    assert decide_price(D("3.79"), active.pricing).small_product
    unsigned, state, tightened = guard_rules(active, None)
    assert state == "UNSIGNED_STRICTEST" and "pricing.small_product_min_order_ttc" in tightened
    assert not decide_price(D("3.79"), unsigned.pricing).small_product
    signed, state, _ = guard_rules(active, active.content_sha256)
    assert state == "SIGNED" and decide_price(D("3.79"), signed.pricing).small_product


def test_quote_cannot_self_declare_a_small_product() -> None:
    client, _ = boot()
    data = body(client.post("/pricing/quote", headers=H, json={"cost_chf": "50", "small_product": True}))
    assert data["decision"]["small_product"] is False and "SMALL_PRODUCT_ADDON" not in data["decision"]["reasons"]


def test_free_shipping_basket_without_real_cost_is_not_allowed() -> None:
    """MOT-18 : port offert sans coût réel => REVIEW, ``allowed`` faux."""
    client, _ = boot()
    data = body(client.post("/pricing/basket", headers=H, json={
        "lines": [{"sku": "A", "qty": 1, "unit_price_ttc": "64.90", "unit_cost": "40"}], "shipping_charged": "0"}))
    assert data["basket"]["status"] == "REVIEW" and data["allowed"] is False
    assert "SHIPPING_COST_UNKNOWN" in data["basket"]["reasons"]


# --------------------------------------------------------------- COH-05 : classeur aligné sur le moteur


def _generator() -> Any:
    import importlib.util
    import sys

    path = ROOT / "docs" / "03-finance" / "generer_classeurs.py"
    spec = importlib.util.spec_from_file_location("generer_classeurs_f3", path)
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module  # dataclasses du générateur : module enregistré avant exécution
    previous = sys.dont_write_bytecode
    sys.dont_write_bytecode = True
    try:
        spec.loader.exec_module(module)
    finally:
        sys.dont_write_bytecode = previous
    return module


def test_delivered_workbook_shows_engine_rounding_row() -> None:
    from openpyxl import load_workbook

    path = ROOT / "docs" / "03-finance" / "modele_financier.xlsx"
    formulas = load_workbook(path)["Prix plancher"]
    values = load_workbook(path, data_only=True)["Prix plancher"]
    assert "moteur" in formulas["A26"].value and formulas["B26"].value.startswith("=IF(ISNUMBER(B17)")
    assert values["B26"].value == pytest.approx(209.90) and values["C29"].value == pytest.approx(209.90)
    assert "lecture BP" in formulas["A21"].value


def test_workbook_engine_rounding_matches_engine_for_several_costs(tmp_path: Path) -> None:
    """COH-05 : 50 => 83.90 (et non 89.90), 20 => 41.90, 140 => 209.90, recalculé par LibreOffice."""
    from openpyxl import load_workbook

    gen = _generator()
    rules = load_rules().pricing
    for cost in (140, 50, 20, 3):
        wb = gen.build_financial_model()
        wb["Prix plancher"]["B5"].value = cost
        src = tmp_path / f"m_{cost}.xlsx"
        wb.save(src)
        try:
            recalculated = gen.recalculate(src, tmp_path / f"recalc_{cost}")
        except gen.RecalcError as exc:
            pytest.skip(f"LibreOffice Calc indisponible : {exc}")
        assert gen.formula_errors(recalculated) == []
        sheet = load_workbook(recalculated, data_only=True)["Prix plancher"]
        engine = decide_price(D(cost), rules).profitable_price
        assert engine is not None
        from pokeshop.pricing import round_up_retail

        expected = round_up_retail(engine, tiers=rules.rounding_tiers)
        assert D(str(sheet["B26"].value)).quantize(D("0.01")) == expected, cost


# ------------------------------------------------------- persistance en base des trois nouveaux journaux


from test_persistance_postgres import pg  # noqa: E402, F401, F811  (base jetable ; ignoré sans PostgreSQL)


def test_f3_registries_survive_restart_in_postgres(pg: dict[str, Any]) -> None:  # noqa: F811
    """Approbations de prix, références « dernier import » et stock local : journal d'état en base."""
    env = {"POKESHOP_API_TOKEN_SHA256": sha256_hex(API_TOKEN), "POKESHOP_OWNER_TOKEN_SHA256": hash_owner_token(OWNER_TOKEN),
           "POKESHOP_AGENT_TOKENS_SHA256": JR.agent_tokens_env()}
    cfg = load_settings(env)

    def start() -> tuple[TestClient, Services]:
        svc = Services.build(cfg, clock=lambda: NOW, notifier=LogNotifier(), connect=pg["engine"])
        return TestClient(create_app(services=svc)), svc

    client, _ = start()
    assert client.post("/pricing/approvals", headers=HO, json={
        "product_key": "FICTIF-P1", "price": "149.90", "reason": "FICTIF prix marché validé"}).status_code == 201
    assert receive(client).status_code == 200
    assert client.post("/imports/fictif_grossiste_a/run", headers=HDATA,
                       json={"source_path": "FICTIF_offres_grossiste_a.csv"}).status_code == 200
    _, svc2 = start()
    assert svc2.restore_errors == {} and svc2.state_backend == "postgres"
    assert svc2.price_approvals.active_for("FICTIF-P1", NOW) is not None
    assert svc2.stock.sellable("DSP-FICTIF_ALPHA-FR") == 6
    assert svc2.baselines.get("fictif_grossiste_a") is not None
