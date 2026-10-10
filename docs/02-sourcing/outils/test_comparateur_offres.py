"""Tests du comparateur d'offres : structure, absence de prix réels, et égalité avec ``pokeshop.pricing``.

Les tests de valeurs recalculent le classeur avec LibreOffice ; ils sont ignorés si ``soffice`` est absent.
Lancer : ``python -m pytest -q docs/02-sourcing/outils``.
"""

from __future__ import annotations

from datetime import date
from decimal import Decimal
from pathlib import Path

import generer_comparateur as gc
import pytest
from openpyxl import load_workbook
from pokeshop.models import LandedCostInput, TierDiscount, VatMode
from pokeshop.pricing import contribution_breakdown, floor_price_exact, landed_unit_cost
from pokeshop.rules import load_rules

D = Decimal
needs_soffice = pytest.mark.skipif(gc.soffice_path() is None, reason="LibreOffice (soffice) absent")


def gtin_ok(code: str) -> bool:
    """Contrôle GS1 d'un GTIN-13."""
    if len(code) != 13 or not code.isdigit():
        return False
    return gc.gtin13(code[:12]) == code


def header_index(ws, row: int) -> dict[str, int]:
    return {str(c.value): c.column for c in ws[row] if c.value is not None}


# ---------------------------------------------------------------- fixtures
@pytest.fixture(scope="module")
def raw_book(tmp_path_factory):
    path = gc.build_workbook(tmp_path_factory.mktemp("raw") / "COMPARATEUR_OFFRES.xlsx")
    return load_workbook(path)


_CALC_CACHE: dict[str, tuple[str, Path, object]] = {}


def _calc(mode: str, tmp_path_factory) -> tuple[str, Path, object]:
    """Classeur généré dans le mode TVA ``mode`` puis recalculé par LibreOffice (mis en cache)."""
    if gc.soffice_path() is None:
        pytest.skip("LibreOffice (soffice) absent")
    if mode not in _CALC_CACHE:
        tmp = tmp_path_factory.mktemp(f"calc_{mode}")
        raw = gc.build_workbook(tmp / "COMPARATEUR_OFFRES.xlsx", vat_mode=mode)
        done = gc.recalculate(raw, tmp / "recalc")
        _CALC_CACHE[mode] = (mode, done, load_workbook(done, data_only=True))
    return _CALC_CACHE[mode]


@pytest.fixture(scope="module", params=["EFFECTIVE", "NOT_REGISTERED"])
def calc_book(request, tmp_path_factory):
    return _calc(request.param, tmp_path_factory)


@pytest.fixture(scope="module")
def calc_effective(tmp_path_factory):
    return _calc("EFFECTIVE", tmp_path_factory)


@pytest.fixture(scope="module")
def calc_not_registered(tmp_path_factory):
    return _calc("NOT_REGISTERED", tmp_path_factory)


# ---------------------------------------------------------------- structure
def test_sheets_present(raw_book):
    assert raw_book.sheetnames == [
        gc.HELP_SHEET,
        gc.PARAM_SHEET,
        gc.SUP_SHEET,
        gc.OFF_SHEET,
        gc.RANK_SHEET,
        gc.SYN_SHEET,
    ]


def test_help_sheet_has_human_validation_section(raw_book):
    texts = [c.value for c in raw_book[gc.HELP_SHEET]["A"] if c.value]
    assert any("VALIDATION HUMAINE REQUISE" in t for t in texts)
    assert any("AUCUN PRIX RÉEL" in t for t in texts)


def test_params_match_rules_file(raw_book):
    ws = raw_book[gc.PARAM_SHEET]
    rules = load_rules(gc.RULES, profile=VatMode.EFFECTIVE)
    p = rules.pricing
    expected = {
        "r": p.payment_pct,
        "b": p.payment_fixed,
        "L": p.logistics_cost,
        "R": p.after_sales_provision,
        "A": p.acquisition_cost,
        "m": p.target_margin,
        "seuil": p.market_review_threshold,
        "tva_ch": p.vat_rate_sales,
        "floor_pct": p.hard_floor_margin,
        "floor_chf": p.hard_floor_chf_per_order,
    }
    for key, value in expected.items():
        assert D(str(ws[f"B{gc.P_CELLS[key]}"].value)) == value, key
    assert ws[f"B{gc.P_CELLS['version']}"].value == rules.rules_version
    assert ws[f"B{gc.P_CELLS['mode']}"].value == "EFFECTIVE"


def test_supplier_rows(raw_book):
    ws = raw_book[gc.SUP_SHEET]
    names = [ws.cell(row=r, column=1).value for r in range(gc.SUP_FIRST, gc.SUP_LAST + 1)]
    for sup in gc.BP_SUPPLIERS:
        assert sup in names
    for r in range(gc.SUP_FIRST, gc.SUP_LAST + 1):
        name = ws.cell(row=r, column=1).value
        if name and not str(name).startswith("FICTIF"):
            # aucune donnée commerciale saisie pour un vrai fournisseur
            for col in range(2, 18):
                assert ws.cell(row=r, column=col).value in (None, ""), (name, col)
            assert ws.cell(row=r, column=18).value == "non contacté"
            assert ws.cell(row=r, column=19).value == "N"
        if name and str(name).startswith("FICTIF"):
            assert ws.cell(row=r, column=19).value == "O"
        assert str(ws.cell(row=r, column=20).value).startswith("=")


def test_offers_prefilled_without_real_prices(raw_book):
    ws = raw_book[gc.OFF_SHEET]
    panier = gc.panier_rows()
    filled = [r for r in range(gc.OFF_FIRST, gc.OFF_LAST + 1) if ws.cell(row=r, column=1).value]
    assert len(filled) == 2 + len(gc.PRIORITY_FR) * len(panier)
    for r in filled:
        fictif = ws.cell(row=r, column=19).value == "O"
        ref = ws.cell(row=r, column=1).value
        if fictif:
            assert ref == gc.FICTIF_REF
            assert str(ws.cell(row=r, column=3).value).startswith("FICTIF_FOURNISSEUR_")
            ean = str(ws.cell(row=r, column=5).value)
            assert ean.startswith("200") and gtin_ok(ean)
        else:
            # colonnes I à R (quantités, prix, marché) vides : aucun prix inventé
            for col in range(9, 19):
                assert ws.cell(row=r, column=col).value in (None, ""), (r, col)
            assert ws.cell(row=r, column=3).value in gc.PRIORITY_FR


def test_each_priority_supplier_gets_same_basket(raw_book):
    ws = raw_book[gc.OFF_SHEET]
    refs = [p["ref_id"] for p in gc.panier_rows()]
    for sup in gc.PRIORITY_FR:
        got = [
            ws.cell(row=r, column=1).value
            for r in range(gc.OFF_FIRST, gc.OFF_LAST + 1)
            if ws.cell(row=r, column=3).value == sup
        ]
        assert got == refs


def test_every_offer_row_has_formulas(raw_book):
    ws = raw_book[gc.OFF_SHEET]
    for r in range(gc.OFF_FIRST, gc.OFF_LAST + 1):
        for col in gc.offer_formulas(r):
            assert str(ws[f"{col}{r}"].value).startswith("="), f"{col}{r}"


def test_headers_cover_input_and_calc(raw_book):
    hdr = header_index(raw_book[gc.OFF_SHEET], gc.OFF_FIRST - 1)
    assert len(hdr) == len(gc.OFF_INPUT) + len(gc.OFF_CALC)
    assert hdr["COÛT RENDU C / unité (CHF)"] == 37  # colonne AK
    assert hdr["Statut"] == 42  # colonne AP


def test_ranking_lists_whole_basket(raw_book):
    ws = raw_book[gc.RANK_SHEET]
    refs = [ws.cell(row=r, column=1).value for r in range(5, 5 + len(gc.rank_refs()))]
    assert refs[0] == gc.FICTIF_REF
    assert refs[1:] == [p["ref_id"] for p in gc.panier_rows()]


def test_gtin13_helper():
    assert gc.gtin13("400638133393") == "4006381333931"  # exemple GS1 classique
    with pytest.raises(ValueError):
        gc.gtin13("123")


def test_unknown_vat_mode_rejected():
    with pytest.raises(ValueError):
        gc.load_params("FOO")


# ---------------------------------------------------------------- valeurs (LibreOffice)
def _row_values(ws, r: int) -> dict[str, object]:
    hdr = header_index(ws, gc.OFF_FIRST - 1)
    return {k: ws.cell(row=r, column=c).value for k, c in hdr.items()}


def _expected_fictif(mode: str) -> dict[str, dict[str, Decimal]]:
    """Recalcule les 2 lignes FICTIF avec pokeshop.pricing (allocation identique au classeur)."""
    vat_mode = VatMode(mode)
    params = load_rules(gc.RULES, profile=vat_mode).pricing
    rate_import = D("0.081")
    out: dict[str, dict[str, Decimal]] = {}
    sup = {s[0]: s for s in gc.FICTIF_SUPPLIERS}
    for row in gc.fictif_offers():
        name = row[2]
        s = sup[name]
        fx, vat_rate, vat_country, ship_country = D(str(s[3])), D(str(s[6])), s[7], s[1]
        port, fees, customs = D(str(s[9])), D(str(s[10])), D(str(s[11]))
        units, price, base, disc, qty = int(row[8]), D(str(row[9])), row[10], row[11], int(row[12])
        disc_d = D(str(disc)) if disc is not None else D(0)
        pack = price * (1 - disc_d)
        ht = pack / (1 + vat_rate) if base == "TTC" else pack
        # une seule ligne par fournisseur FICTIF : part de la commande = 100 %
        freight_unit = (port + fees) * fx / (units * qty)
        customs_unit = customs / (units * qty)
        purchase_unit = ht * fx / units
        import_vat = (purchase_unit + freight_unit + customs_unit) * rate_import if ship_country != "CH" else D(0)
        inp = LandedCostInput(
            purchase_net=price,
            currency=s[2],
            fx_rate_to_chf=fx,
            fx_source=s[4],
            fx_date=date(2026, 10, 4),
            price_includes_vat=(base == "TTC"),
            supplier_vat_rate=vat_rate,
            units_per_pack=units,
            inbound_freight_alloc=freight_unit,
            customs_and_fees=customs_unit,
            import_vat=import_vat,
            vat_mode=vat_mode,
            supplier_vat_country=vat_country,
            tier_discounts=(TierDiscount(min_qty=1, discount_pct=disc_d),) if disc_d else (),
            order_qty=units * qty,
        )
        market = D(str(row[17]))
        out[name] = {"C": landed_unit_cost(inp), "market": market, "params": params}  # type: ignore[dict-item]
    return out


@needs_soffice
def test_no_formula_error(calc_book):
    _, path, _ = calc_book
    assert gc.formula_errors(path) == []


@needs_soffice
def test_landed_cost_matches_engine(calc_book):
    mode, _, wb = calc_book
    ws = wb[gc.OFF_SHEET]
    exp = _expected_fictif(mode)
    for r in (gc.OFF_FIRST, gc.OFF_FIRST + 1):
        v = _row_values(ws, r)
        c_xlsx = D(str(v["COÛT RENDU C / unité (CHF)"]))
        assert abs(c_xlsx - exp[v["Fournisseur"]]["C"]) <= D("0.005"), (mode, v["Fournisseur"], c_xlsx)


@needs_soffice
def test_floor_and_contribution_match_engine(calc_book):
    mode, _, wb = calc_book
    ws = wb[gc.OFF_SHEET]
    params = load_rules(gc.RULES, profile=VatMode(mode)).pricing
    for r in (gc.OFF_FIRST, gc.OFF_FIRST + 1):
        v = _row_values(ws, r)
        c_xlsx = D(str(v["COÛT RENDU C / unité (CHF)"]))
        floor = floor_price_exact(c_xlsx, params)
        assert abs(D(str(v["Prix plancher TTC (CHF)"])) - floor) < D("0.000001")
        market = D(str(v["Prix marché observé TTC (CHF, médiane)"]))
        cb = contribution_breakdown(market, c_xlsx, params)
        assert abs(D(str(v["Contribution au prix marché (CHF)"])) - cb.contribution_chf) <= D("0.02")
        # statut selon la règle BP §5 recalculée ici
        gap = floor / market - 1
        contrib = D(str(v["Contribution au prix marché (CHF)"]))
        pct = D(str(v["Contribution au prix marché (% CA net)"]))
        if gap <= params.market_review_threshold:
            expected = "OK"
        elif pct >= params.hard_floor_margin and contrib >= params.hard_floor_chf_per_order:
            expected = "À REVOIR"
        else:
            expected = "NON RENTABLE"
        assert v["Statut"] == expected


@needs_soffice
def test_effective_mode_reference_values(calc_effective):
    """Valeurs de référence de l'exemple FICTIF en mode EFFECTIVE (calcul détaillé en commentaire)."""
    _, _, wb = calc_effective
    ws = wb[gc.OFF_SHEET]
    a, b = _row_values(ws, gc.OFF_FIRST), _row_values(ws, gc.OFF_FIRST + 1)
    # A : 40 EUR HT −5 % = 38 ; ×0,94 = 35,72 ; port 40 EUR ×0,94 /10 = 3,76 ; dédouanement 25/10 = 2,50
    assert D(str(a["COÛT RENDU C / unité (CHF)"])).quantize(D("0.01")) == D("41.98")
    # B : 49,90 TTC CH, TVA suisse récupérable : 49,90/1,081 + 15/10
    assert D(str(b["COÛT RENDU C / unité (CHF)"])).quantize(D("0.01")) == D("47.66")
    assert (a["Rang (1 = moins cher)"], b["Rang (1 = moins cher)"]) == (1, 2)
    assert (a["Statut"], b["Statut"]) == ("OK", "NON RENTABLE")
    assert "FICTIF" in a["Alertes"]


@needs_soffice
def test_not_registered_mode_adds_import_and_foreign_vat(calc_not_registered):
    _, _, wb = calc_not_registered
    ws = wb[gc.OFF_SHEET]
    a, b = _row_values(ws, gc.OFF_FIRST), _row_values(ws, gc.OFF_FIRST + 1)
    # A : + TVA import estimée (41,98 × 8,1 %) ; B : TVA suisse non récupérable → 49,90 + 1,50
    assert D(str(a["COÛT RENDU C / unité (CHF)"])).quantize(D("0.01")) == D("45.38")
    assert D(str(b["COÛT RENDU C / unité (CHF)"])).quantize(D("0.01")) == D("51.40")
    pm = wb[gc.PARAM_SHEET]
    assert pm[f"B{gc.P_CELLS['t']}"].value == 0


@needs_soffice
def test_ranking_and_synthesis(calc_book):
    _, _, wb = calc_book
    rk = wb[gc.RANK_SHEET]
    assert rk["A5"].value == gc.FICTIF_REF
    assert rk["C5"].value == 2
    assert rk["D5"].value == "FICTIF_FOURNISSEUR_A"
    assert rk["F5"].value == "FICTIF_FOURNISSEUR_B"
    assert rk["K6"].value == "AUCUNE OFFRE"
    sy = wb[gc.SYN_SHEET]
    rows = {
        sy.cell(row=r, column=1).value: [sy.cell(row=r, column=c).value for c in range(1, 12)]
        for r in range(5, 5 + gc.SUP_LAST - gc.SUP_FIRST + 1)
    }
    a = rows["FICTIF_FOURNISSEUR_A"]
    assert a[1] == 1 and a[3] == 1
    assert abs(D(str(a[4])) - D("357.20")) < D("0.0001")  # 38 × 0,94 × 10
    assert a[7] == "NON"  # 380 EUR < minimum FICTIF de 500 EUR
    assert rows["FICTIF_FOURNISSEUR_B"][7] == "OUI"
    assert rows["Asmodee France"][2] == len(gc.panier_rows())  # lignes incomplètes en attente de devis
    assert rows["Asmodee France"][10] == "non contacté"


def test_delivered_file_in_sync_and_recalculated():
    """Le fichier livré contient les formules du générateur et des valeurs calculées."""
    path = gc.OUTPUT
    assert path.exists()
    wb_f = load_workbook(path)
    ws = wb_f[gc.OFF_SHEET]
    assert str(ws[f"AK{gc.OFF_FIRST}"].value).startswith("=IF(")
    wb_v = load_workbook(path, data_only=True)
    ak = wb_v[gc.OFF_SHEET][f"AK{gc.OFF_FIRST}"].value
    assert ak is not None, "fichier livré non recalculé : lancer generer_comparateur.py"
    assert D(str(ak)).quantize(D("0.01")) == D("41.98")
    assert gc.formula_errors(path) == []
    assert wb_v[gc.PARAM_SHEET][f"B{gc.P_CELLS['mode']}"].value == "EFFECTIVE"
