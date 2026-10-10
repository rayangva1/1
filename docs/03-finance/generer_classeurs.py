"""Génère les classeurs finance avec formules Excel vivantes (BP §3, §4, §10).

Produit dans ``docs/03-finance/`` :

* ``modele_financier.xlsx`` — Hypothèses, Étoile polaire (contribution nette cumulée
  hebdomadaire + stop-loss global), Budget initial, Charges mensuelles, Scénarios,
  Seuil & sensibilité, Prix plancher, Stock & BFR ;
* ``tresorerie_13_semaines.xlsx`` — mode d'emploi, exemple FICTIF, modèle à remplir
  (réserve 1 600 CHF = seuil du stop-loss cash).

Les formules sont recalculées par LibreOffice headless (``soffice``) afin que les
fichiers livrés contiennent aussi les valeurs (lisibles par pandas, aperçus,
``openpyxl`` en ``data_only=True``). Excel recalcule de toute façon à l'ouverture.

Usage ::

    python docs/03-finance/generer_classeurs.py            # génère + recalcule + contrôle
    python docs/03-finance/generer_classeurs.py --sans-recalcul
"""

from __future__ import annotations

import argparse
import shutil
import subprocess
import sys
import tempfile
from dataclasses import dataclass
from datetime import date
from decimal import Decimal
from pathlib import Path
from typing import Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

ROOT = Path(__file__).resolve().parents[2]
if str(ROOT / "engine") not in sys.path:
    sys.path.insert(0, str(ROOT / "engine"))

from pokeshop.pricing import round_up_retail  # noqa: E402
from pokeshop.rules import load_rules  # noqa: E402
from pokeshop.treasury import WeeklyInput  # noqa: E402

OUT_DIR = Path(__file__).resolve().parent
MODEL_FILE = "modele_financier.xlsx"
TREASURY_FILE = "tresorerie_13_semaines.xlsx"

# ---------------------------------------------------------------------------
# Styles
# ---------------------------------------------------------------------------

FONT_NAME = "Arial"
F_BASE = Font(name=FONT_NAME, size=10)
F_INPUT = Font(name=FONT_NAME, size=10, color="0000FF")
F_LINK = Font(name=FONT_NAME, size=10, color="008000")
F_BP = Font(name=FONT_NAME, size=10, italic=True, color="595959")
F_BOLD = Font(name=FONT_NAME, size=10, bold=True)
F_TITLE = Font(name=FONT_NAME, size=14, bold=True, color="1F3A5F")
F_SUB = Font(name=FONT_NAME, size=9, italic=True, color="595959")
F_HEADER = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
F_SECTION = Font(name=FONT_NAME, size=11, bold=True, color="1F3A5F")
F_WARN = Font(name=FONT_NAME, size=10, bold=True, color="C00000")

FILL_INPUT = PatternFill("solid", fgColor="FFF2CC")
FILL_BP = PatternFill("solid", fgColor="EDEDED")
FILL_HEADER = PatternFill("solid", fgColor="1F3A5F")
FILL_SECTION = PatternFill("solid", fgColor="D9E1F2")
FILL_RESULT = PatternFill("solid", fgColor="E2EFDA")
FILL_ALERT = PatternFill("solid", fgColor="F8CBAD")
FILL_WARN = PatternFill("solid", fgColor="FCE4D6")

THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
TOP_RULE = Border(top=Side(style="thin", color="000000"))

NF_CHF = '#,##0;(#,##0);"-"'
NF_CHF2 = '#,##0.00;(#,##0.00);"-"'
NF_CHF2_RED = '#,##0.00;[Red](#,##0.00);"-"'
NF_PCT1 = "0.0%"
NF_PCT2 = "0.00%"
NF_INT = '#,##0;(#,##0);"0"'
NF_DEC1 = "0.0"
NF_DEC2 = "0.00"
NF_DATE = "DD.MM.YYYY"

WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def q(sheet: str, cell: str) -> str:
    """Référence inter-feuilles toujours entre apostrophes."""
    return f"'{sheet}'!{cell}"


def put(
    ws: Worksheet,
    cell: str,
    value: object,
    *,
    font: Font = F_BASE,
    fill: PatternFill | None = None,
    fmt: str | None = None,
    align: Alignment | None = None,
    border: Border | None = None,
    comment: str | None = None,
) -> None:
    """Écrit une cellule avec son style."""
    c = ws[cell]
    c.value = value
    c.font = font
    if fill is not None:
        c.fill = fill
    if fmt is not None:
        c.number_format = fmt
    if align is not None:
        c.alignment = align
    if border is not None:
        c.border = border
    if comment:
        c.comment = Comment(comment, "agent finance")


def inp(ws: Worksheet, cell: str, value: object, fmt: str | None = None, comment: str | None = None) -> None:
    """Cellule d'entrée : texte bleu sur fond jaune."""
    put(ws, cell, value, font=F_INPUT, fill=FILL_INPUT, fmt=fmt, border=BOX, comment=comment)


def bp(ws: Worksheet, cell: str, value: object, fmt: str | None = None) -> None:
    """Valeur publiée dans le BP (référence de contrôle, ne pas modifier)."""
    put(ws, cell, value, font=F_BP, fill=FILL_BP, fmt=fmt, border=BOX)


def header(ws: Worksheet, row: int, labels: Sequence[str], start_col: int = 1) -> None:
    """Ligne d'en-tête foncée."""
    for offset, label in enumerate(labels):
        c = ws.cell(row=row, column=start_col + offset, value=label)
        c.font = F_HEADER
        c.fill = FILL_HEADER
        c.alignment = CENTER
        c.border = BOX


def section(ws: Worksheet, row: int, text: str, width: int) -> None:
    """Bandeau de section sur ``width`` colonnes."""
    for col in range(1, width + 1):
        ws.cell(row=row, column=col).fill = FILL_SECTION
    c = ws.cell(row=row, column=1, value=text)
    c.font = F_SECTION


def title(ws: Worksheet, text: str, subtitle: str) -> None:
    """Titre + sous-titre de feuille."""
    put(ws, "A1", text, font=F_TITLE)
    put(ws, "A2", subtitle, font=F_SUB)


def widths(ws: Worksheet, values: dict[str, float]) -> None:
    """Largeurs de colonnes."""
    for col, width in values.items():
        ws.column_dimensions[col].width = width


def name(wb: Workbook, label: str, sheet: str, cell: str) -> None:
    """Nom défini pointant vers une cellule absolue."""
    col = "".join(ch for ch in cell if ch.isalpha())
    row = "".join(ch for ch in cell if ch.isdigit())
    wb.defined_names[label] = DefinedName(label, attr_text=f"'{sheet}'!${col}${row}")


def mput(ws: Worksheet, cells: str, value: object, **style: object) -> None:
    """Fusionne la plage ``cells`` et écrit ``value`` dans sa première cellule."""
    first = cells.split(":")[0]
    ws.merge_cells(cells)
    put(ws, first, value, **style)  # type: ignore[arg-type]


def page_setup(ws: Worksheet, landscape: bool = True) -> None:
    """Impression A4 ajustée en largeur."""
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


# ---------------------------------------------------------------------------
# Classeur 1 : modèle financier
# ---------------------------------------------------------------------------

S_HYP = "Hypothèses"
S_NS = "Étoile polaire"
S_BUD = "Budget initial"
S_FIX = "Charges mensuelles"
S_SCN = "Scénarios"
S_SEU = "Seuil & sensibilité"
S_PRX = "Prix plancher"
S_STK = "Stock & BFR"

SCENARIO_COLS = ("B", "C", "D")
SCENARIO_NAMES = ("Prudent", "Central", "Développement")

LEGEND = (
    "Légende : cellules jaunes à texte bleu = entrées à modifier ; texte noir = formules (ne pas écraser) ; "
    "texte vert = lien vers une autre feuille ; gris italique = valeur publiée dans le BP (référence de contrôle)."
)


def _sheet_hypotheses(wb: Workbook) -> None:
    ws = wb.active
    ws.title = S_HYP
    title(ws, "Modèle financier — {{NOM_BOUTIQUE}} · boutique Pokémon JCC FR, ventes Suisse",
          "Source : business plan du 4 octobre 2026 (BP). Toutes les valeurs sont des hypothèses de gestion, "
          "pas des devis. Montants en CHF.")
    put(ws, "A3", LEGEND, font=F_SUB, align=WRAP)
    ws.merge_cells("A3:E3")
    ws.row_dimensions[3].height = 28

    section(ws, 5, "Paramètres généraux", 5)
    header(ws, 6, ["Paramètre", "Valeur", "Unité", "Source / statut", "Nom défini"])
    rows: list[tuple[str, object, str, str, str, str, bool]] = [
        # libellé, valeur, format, unité, source, nom, lien?
        ("TVA sur ventes (taux normal)", 0.081, NF_PCT1, "%", "BP §4 [S6] — statut TVA à confirmer avec la fiduciaire",
         "TVA", False),
        ("Seuil d'assujettissement TVA", 100000, NF_CHF, "CHF/an", "BP §4 [S6] — CA déterminant de toute l'entité",
         "Seuil_TVA", False),
        ("Panier produits moyen TTC", 95, NF_CHF2, "CHF", "BP §10 — hypothèse", "Panier_TTC", False),
        ("Panier produits moyen HT", "=Panier_TTC/(1+TVA)", NF_CHF2, "CHF", "Calcul", "Panier_HT", False),
        ("Contribution produit avant acquisition", 0.22, NF_PCT1, "% CA HT",
         "BP §10 — après achat, paiement, emballage, provision SAV", "Taux_contribution", False),
        ("Charges fixes mensuelles", f"={q(S_FIX, 'B7')}", NF_CHF, "CHF/mois", "Lien feuille Charges mensuelles (BP §3)",
         "Charges_fixes", True),
        ("Rémunération budgétée", 2000, NF_CHF, "CHF/mois", "BP §10 — hors cotisations sociales", "Remuneration",
         False),
        ("Coût produit en % du CA net", 0.70, NF_PCT1, "%", "BP §10 — « par hypothèse »", "Cout_produit_pct", False),
        ("Stock initial acheté rendu Suisse", f"={q(S_BUD, 'B5')}", NF_CHF, "CHF", "Lien feuille Budget initial (BP §3)",
         "Stock_initial", True),
        ("Jours par mois (convention)", 30, "0", "jours", "Convention de calcul", "Jours_mois", False),
        ("Semaines par mois", "=52/12", NF_DEC2, "sem.", "Convention 52/12", "Semaines_mois", False),
        ("Plancher dur : contribution par commande", 8, NF_CHF2, "CHF", "BP §5", "Plancher_CHF", False),
        ("Plancher dur : contribution en % du CA net", 0.12, NF_PCT1, "%", "BP §5", "Plancher_pct", False),
        ("Supervision au lancement — minimum", 6, "0", "h/semaine", "BP §3", "Heures_min", False),
        ("Supervision au lancement — maximum", 10, "0", "h/semaine", "BP §3", "Heures_max", False),
        ("Préparation d'une commande (colis)", 15, "0", "minutes",
         "HYPOTHÈSE FICTIVE — à mesurer lors des commandes tests", "Minutes_prep", False),
    ]
    row = 7
    for label, value, fmt, unit, source, nm, link in rows:
        put(ws, f"A{row}", label, border=BOX)
        if isinstance(value, str) and value.startswith("="):
            put(ws, f"B{row}", value, font=F_LINK if link else F_BASE, fmt=fmt, border=BOX)
        else:
            inp(ws, f"B{row}", value, fmt=fmt)
        put(ws, f"C{row}", unit, border=BOX)
        put(ws, f"D{row}", source, border=BOX,
            font=F_WARN if "FICTIVE" in source else F_BASE)
        put(ws, f"E{row}", nm, font=F_SUB, border=BOX)
        name(wb, nm, S_HYP, f"B{row}")
        row += 1
    # row == 23
    section(ws, 25, "Scénarios mensuels (BP §10)", 5)
    header(ws, 26, ["Paramètre", *SCENARIO_NAMES, "Source"])
    put(ws, "A27", "Commandes payées par mois", border=BOX)
    put(ws, "A28", "CAC moyen mélangé (CHF/commande)", border=BOX)
    for col, orders, cac in zip(SCENARIO_COLS, (40, 100, 200), (8, 6, 5)):
        inp(ws, f"{col}27", orders, fmt="0")
        inp(ws, f"{col}28", cac, fmt=NF_CHF2)
    put(ws, "E27", "BP §10 — simulations, pas une prévision de demande", border=BOX)
    put(ws, "E28", "BP §10 — CAC sur commandes payées après annulations", border=BOX)

    section(ws, 30, "Sensibilité (BP §10)", 5)
    header(ws, 31, ["Paramètre", "Valeur", "Unité", "Source", "Nom défini"])
    sens = [
        ("Contribution dégradée", 0.15, NF_PCT1, "% CA HT", "BP §10", "Taux_sensibilite"),
        ("CAC élevé", 15, NF_CHF2, "CHF", "BP §10", "CAC_eleve"),
    ]
    row = 32
    for label, value, fmt, unit, source, nm in sens:
        put(ws, f"A{row}", label, border=BOX)
        inp(ws, f"B{row}", value, fmt=fmt)
        put(ws, f"C{row}", unit, border=BOX)
        put(ws, f"D{row}", source, border=BOX)
        put(ws, f"E{row}", nm, font=F_SUB, border=BOX)
        name(wb, nm, S_HYP, f"B{row}")
        row += 1

    section(ws, 35, "Délais du cycle cash — HYPOTHÈSES FICTIVES à remplacer (devis fournisseur, contrat PSP)", 5)
    header(ws, 36, ["Paramètre", "Valeur", "Unité", "Source / statut", "Nom défini"])
    delays = [
        ("Paiement fournisseur → réception (prépaiement)", 14, "Delai_prepaiement",
         "FICTIF — BP §4 : ne pas supposer de facture HT/à terme avant confirmation"),
        ("Jours de ventes détenus en stock", 30, "Jours_stock", "FICTIF — politique de stock à décider"),
        ("Crédit fournisseur après réception", 0, "Credit_fournisseur", "FICTIF — 0 = prépaiement (prudent)"),
        ("Délai de versement PSP", 7, "Delai_PSP", "FICTIF — BP §7 : tester les délais de versement"),
    ]
    row = 37
    for label, value, nm, source in delays:
        put(ws, f"A{row}", label, border=BOX)
        inp(ws, f"B{row}", value, fmt="0")
        put(ws, f"C{row}", "jours", border=BOX)
        put(ws, f"D{row}", source, font=F_WARN, border=BOX)
        put(ws, f"E{row}", nm, font=F_SUB, border=BOX)
        name(wb, nm, S_HYP, f"B{row}")
        row += 1

    section(ws, 42, "Étoile polaire et stop-loss du mandat — décisions de la propriétaire", 5)
    header(ws, 43, ["Paramètre", "Valeur", "Unité", "Source / statut", "Nom défini"])
    mandate: list[tuple[str, object, str, str, str, str, bool]] = [
        ("Capital engagé de référence = point zéro du stop-loss global",
         f"={q(S_BUD, 'B12')}-SUM({q(S_BUD, 'B6:B9')})-(1-0.7)*{q(S_BUD, 'B5')}", NF_CHF, "CHF",
         "HYPOTHÈSE — STOP_LOSS.md §5 option A (décision C03 à J3) : apports (Budget l. 12) − lancement assumé "
         "(l. 6 à 9) − décote prudente du stock de 30 % (ratio de liquidation 0,70) = 4 200 CHF ; remplacer par le "
         "point zéro réellement posé (POST /stoploss/baseline)", "Capital_engage", True),
        ("Stop-loss global : perte de valeur nette maximale", 0.20, NF_PCT1, "% référence",
         "Mandat propriétaire — gel total, réarmement par la propriétaire uniquement", "StopLoss_global_pct", False),
        ("Perte déclenchant le gel global (840 CHF au BP)", "=Capital_engage*StopLoss_global_pct", NF_CHF, "CHF",
         "Calcul — définition unique du moteur (pokeshop.stoploss) ; approchée ici par la contribution nette "
         "cumulée depuis le point zéro (l'ancien seuil de 1 600 CHF sur 8 000 est caduc)", "Perte_max", False),
        ("Réserve de trésorerie = seuil du stop-loss cash", f"={q(S_BUD, 'B11')}", NF_CHF, "CHF",
         "Lien Budget initial (BP §3) — sous ce seuil : plus d'achat ni de pub", "Reserve_cash", True),
        ("Semaine d'ouverture des ventes", 5, "0", "n° semaine",
         "BP §9 : ouverture douce aux jours 31 à 45 (semaine 5)", "Semaine_ouverture", False),
        ("Scénario projeté (1 prudent, 2 central, 3 développement)", 2, "0", "choix",
         "Feuille Étoile polaire", "Scenario_choisi", False),
        ("Frais de paiement proportionnels r", 0.025, NF_PCT1, "% du TTC", "BP §4 (hypothèse) — contrat PSP",
         "Frais_paiement_pct", False),
        ("Frais de paiement fixes b", 0.30, NF_CHF2, "CHF/commande", "BP §4 (hypothèse) — contrat PSP",
         "Frais_paiement_fixe", False),
        ("Provision SAV R", 1, NF_CHF2, "CHF/commande", "BP §4 (hypothèse)", "Provision_SAV", False),
        ("Fenêtre de validation", 60, "0", "jours",
         "BP §1 — le classeur retient les semaines pleines (60 j ⇒ 8 semaines, prudent)", "Jours_validation", False),
        ("Commandes payées visées sur la fenêtre", 30, "0", "commandes", "BP §1 — critère de test, pas une prévision",
         "Commandes_validation", False),
    ]
    row = 44
    for label, value, fmt, unit, source, nm, link in mandate:
        put(ws, f"A{row}", label, border=BOX)
        if isinstance(value, str) and value.startswith("="):
            put(ws, f"B{row}", value, font=F_LINK if link else F_BASE, fmt=fmt, border=BOX)
        else:
            inp(ws, f"B{row}", value, fmt=fmt)
        put(ws, f"C{row}", unit, border=BOX)
        put(ws, f"D{row}", source, border=BOX, font=F_WARN if "HYPOTHÈSE" in source else F_BASE)
        put(ws, f"E{row}", nm, font=F_SUB, border=BOX)
        name(wb, nm, S_HYP, f"B{row}")
        row += 1
    dv = DataValidation(type="whole", operator="between", formula1="1", formula2="3", allow_blank=False,
                        showErrorMessage=True, errorTitle="Scénario", error="1 = prudent, 2 = central, 3 = développement")
    ws.add_data_validation(dv)
    dv.add("B49")
    widths(ws, {"A": 46, "B": 16, "C": 16, "D": 62, "E": 18})
    ws.freeze_panes = "A5"
    page_setup(ws)


def _sheet_budget(wb: Workbook) -> None:
    ws = wb.create_sheet(S_BUD)
    title(ws, "Budget initial — décaissement pilote (BP §3)",
          "Aucune offre prestataire confirmée. Remplacer chaque ligne par un devis avant engagement.")
    header(ws, 4, ["Poste", "Budget pilote CHF", "Nature (BP)", "Traitement en trésorerie"])
    lines = [
        ("Stock acheté rendu Suisse", 3000, "Hypothèse de coût cash", "Achats de stock engagés"),
        ("Site et automatisation pilote", 1500, "Budget plafonné ; réalisation interne/assistée",
         "Dépenses de lancement"),
        ("DA et contenus de lancement", 400, "Dépenses cash hors temps interne", "Dépenses de lancement"),
        ("Administration et revue des documents", 700, "Provision à confirmer par devis", "Dépenses de lancement"),
        ("Emballages et matériel", 300, "Provision", "Dépenses de lancement"),
        ("Test acquisition", 500, "Plafond avant validation CAC", "Publicité (jours 46 à 60)"),
        ("Réserve de trésorerie", 1600, "Réassort, versements différés, remboursement", "Non décaissée"),
    ]
    for i, (label, amount, nature, treat) in enumerate(lines, start=5):
        put(ws, f"A{i}", label, border=BOX)
        inp(ws, f"B{i}", amount, fmt=NF_CHF)
        put(ws, f"C{i}", nature, border=BOX)
        put(ws, f"D{i}", treat, border=BOX)
    put(ws, "A12", "Total", font=F_BOLD, border=BOX)
    put(ws, "B12", "=SUM(B5:B11)", font=F_BOLD, fmt=NF_CHF, fill=FILL_RESULT, border=BOX)
    put(ws, "A13", "Total publié dans le BP", border=BOX)
    bp(ws, "B13", 8000, NF_CHF)
    put(ws, "A14", "Contrôle", border=BOX)
    put(ws, "B14", '=IF(B12=B13,"OK — conforme BP","ÉCART "&TEXT(B12-B13,"#,##0"))', border=BOX)
    put(ws, "A15", "Décaissé avant la première vente (hors réserve)", border=BOX)
    put(ws, "B15", "=SUM(B5:B10)", fmt=NF_CHF, border=BOX)
    put(ws, "A16", "Trésorerie conservée à l'ouverture (réserve)", border=BOX)
    put(ws, "B16", "=B11", fmt=NF_CHF, border=BOX)
    put(ws, "C16", "Attention : les charges fixes des mois sans ventes ne figurent pas dans ce budget "
                   "(voir NOTE_VERIFICATION_BP.md).", font=F_WARN, align=WRAP)
    ws.merge_cells("C16:D16")
    ws.row_dimensions[16].height = 28

    section(ws, 18, "Répartition indicative du budget stock (BP §1) — suit les marges et allocations, "
                    "pas un objectif de volume", 4)
    header(ws, 19, ["Catégorie", "Part du stock", "CHF", "Remarque"])
    split = [("Displays", 0.25), ("ETB", 0.25), ("Bundles et tripacks", 0.20), ("Coffrets", 0.20),
             ("Accessoires", 0.10)]
    for i, (label, pct) in enumerate(split, start=20):
        put(ws, f"A{i}", label, border=BOX)
        inp(ws, f"B{i}", pct, fmt=NF_PCT1)
        put(ws, f"C{i}", f"=$B$5*B{i}", fmt=NF_CHF, border=BOX)
        put(ws, f"D{i}", "", border=BOX)
    put(ws, "A25", "Total", font=F_BOLD, border=BOX)
    put(ws, "B25", "=SUM(B20:B24)", font=F_BOLD, fmt=NF_PCT1, border=BOX)
    put(ws, "C25", "=SUM(C20:C24)", font=F_BOLD, fmt=NF_CHF, border=BOX)
    put(ws, "D25", '=IF(ROUND(B25,4)=1,"OK — 100 %","ATTENTION : la répartition ne fait pas 100 %")', border=BOX)
    put(ws, "A26", "Plafond par extension au lancement", border=BOX)
    inp(ws, "B26", 0.25, fmt=NF_PCT1)
    put(ws, "C26", "=$B$5*B26", fmt=NF_CHF, border=BOX)
    put(ws, "D26", "BP §1 — exceptions à documenter", border=BOX)
    widths(ws, {"A": 46, "B": 18, "C": 46, "D": 40})
    page_setup(ws)


def _sheet_fixed(wb: Workbook) -> None:
    ws = wb.create_sheet(S_FIX)
    title(ws, "Charges mensuelles de simulation (BP §3)",
          "Enveloppes, pas les tarifs actuels des logiciels ou prestataires. Hors publicité et rémunération.")
    header(ws, 3, ["Poste", "CHF par mois", "Source / à remplacer par"])
    lines = [
        ("Site, apps, hébergement et automatisation", 180, "Abonnement boutique + apps + hébergement (devis)"),
        ("Comptabilité et administration", 120, "Offre fiduciaire"),
        ("Divers, assurance et petit stockage", 100, "Offres assurance RC / stockage"),
    ]
    for i, (label, amount, src) in enumerate(lines, start=4):
        put(ws, f"A{i}", label, border=BOX)
        inp(ws, f"B{i}", amount, fmt=NF_CHF)
        put(ws, f"C{i}", src, border=BOX)
    put(ws, "A7", "Total fixe hors publicité et rémunération", font=F_BOLD, border=BOX)
    put(ws, "B7", "=SUM(B4:B6)", font=F_BOLD, fmt=NF_CHF, fill=FILL_RESULT, border=BOX)
    put(ws, "A8", "Total publié dans le BP", border=BOX)
    bp(ws, "B8", 400, NF_CHF)
    put(ws, "A9", "Contrôle", border=BOX)
    put(ws, "B9", '=IF(B7=B8,"OK — conforme BP","ÉCART "&TEXT(B7-B8,"#,##0"))', border=BOX)

    section(ws, 11, "Temps de travail non rémunéré dans le résultat (BP §3 : le valoriser à part)", 3)
    header(ws, 12, ["Indicateur", "Heures par mois", "Calcul"])
    put(ws, "A13", "Supervision — minimum", border=BOX)
    put(ws, "B13", "=Heures_min*Semaines_mois", fmt=NF_DEC1, border=BOX)
    put(ws, "C13", "6 h/semaine × 52/12", border=BOX)
    put(ws, "A14", "Supervision — maximum", border=BOX)
    put(ws, "B14", "=Heures_max*Semaines_mois", fmt=NF_DEC1, border=BOX)
    put(ws, "C14", "10 h/semaine × 52/12", border=BOX)
    put(ws, "A15", "Préparation des colis", border=BOX)
    put(ws, "C15", "Voir feuille Scénarios (minutes par commande × commandes)", border=BOX)
    widths(ws, {"A": 46, "B": 18, "C": 58})
    page_setup(ws, landscape=False)


def _sheet_scenarios(wb: Workbook) -> None:
    ws = wb.create_sheet(S_SCN)
    title(ws, "Scénarios mensuels (BP §10)",
          "Entité assujettie, méthode effective. Aucun salaire, impôt, amortissement ni coût initial inclus. "
          "Valeurs exactes ; affichage arrondi.")
    header(ws, 4, ["Indicateur (par mois)", *SCENARIO_NAMES, "Unité / formule"])
    spec: list[tuple[int, str, str, str, str]] = [
        (5, "Commandes payées", f"={q(S_HYP, '{c}27')}", "0", "Hypothèses"),
        (6, "Panier produits TTC", "=Panier_TTC", NF_CHF2, "CHF"),
        (7, "Panier produits HT", "=Panier_HT", NF_CHF2, "CHF = TTC / (1 + TVA)"),
        (8, "CA produits TTC", "={c}5*{c}6", NF_CHF, "CHF"),
        (9, "CA produits HT", "={c}8/(1+TVA)", NF_CHF, "CHF"),
        (10, "TVA collectée (à reverser)", "={c}8-{c}9", NF_CHF, "CHF"),
        (11, "Contribution avant acquisition", "={c}9*Taux_contribution", NF_CHF, "CHF = CA HT × 22 %"),
        (12, "CAC moyen", f"={q(S_HYP, '{c}28')}", NF_CHF2, "CHF / commande"),
        (13, "Acquisition totale", "={c}5*{c}12", NF_CHF, "CHF"),
        (14, "Charges fixes", "=Charges_fixes", NF_CHF, "CHF"),
        (15, "Résultat avant rémunération", "={c}11-{c}13-{c}14", NF_CHF, "CHF"),
        (16, "Rémunération budgétée", "=Remuneration", NF_CHF, "CHF (hors cotisations)"),
        (17, "Résultat après rémunération", "={c}15-{c}16", NF_CHF, "CHF"),
        (18, "Contribution / commande avant acquisition", "={c}7*Taux_contribution", NF_CHF2, "CHF"),
        (19, "Contribution / commande après CAC", "={c}18-{c}12", NF_CHF2, "CHF"),
        (20, "Contribution après CAC en % du CA HT", "=IF({c}7>0,{c}19/{c}7,0)", NF_PCT1, "%"),
        (21, "Plancher dur BP §5 respecté ? (≥ 8 CHF et ≥ 12 %)",
         '=IF(AND({c}19>=Plancher_CHF,{c}20>=Plancher_pct),"oui","NON")', "@", "BP §5"),
        (22, "CA TTC annualisé", "={c}8*12", NF_CHF, "CHF/an"),
        (23, "Seuil TVA 100 000 CHF atteint ?", '=IF({c}22>=Seuil_TVA,"OUI — anticiper l\'examen TVA","non")',
         "@", "BP §4, §10"),
        (24, "Heures de travail / mois (supervision moyenne + colis)",
         "=(Heures_min+Heures_max)/2*Semaines_mois+{c}5*Minutes_prep/60", NF_DEC1, "h"),
        (25, "Rémunération implicite du travail", "=IF({c}24>0,{c}15/{c}24,0)", NF_CHF2,
         "CHF/h, avant impôts et charges sociales"),
        (26, "Stop-loss pub : CAC > contribution / commande ?", '=IF({c}12>{c}18,"OUI — campagne coupée","non")',
         "@", "Mandat (7 jours glissants en exploitation)"),
    ]
    for row, label, formula, fmt, unit in spec:
        bold = row in (15, 17)
        put(ws, f"A{row}", label, font=F_BOLD if bold else F_BASE, border=BOX)
        for c in SCENARIO_COLS:
            text = formula.replace("{c}", c)
            link = "Hypothèses" in text
            put(ws, f"{c}{row}", text, font=F_LINK if link else (F_BOLD if bold else F_BASE), fmt=fmt,
                border=BOX, fill=FILL_RESULT if bold else None)
        put(ws, f"E{row}", unit, font=F_SUB, border=BOX)
    ws.conditional_formatting.add("B21:D21", CellIsRule(operator="equal", formula=['"NON"'], fill=FILL_ALERT))
    ws.conditional_formatting.add("B26:D26", FormulaRule(formula=['LEFT(B26,3)="OUI"'], fill=FILL_ALERT))
    ws.conditional_formatting.add("B15:D17", CellIsRule(operator="lessThan", formula=["0"], fill=FILL_ALERT))

    section(ws, 27, "Contrôle vs valeurs publiées au BP §10 (arrondies au franc)", 5)
    header(ws, 28, ["Valeur BP", *SCENARIO_NAMES, "Écart = arrondi(recalcul) − BP"])
    checks = [
        (29, "CA produits TTC", 8, (3800, 9500, 19000)),
        (30, "CA produits HT", 9, (3515, 8788, 17577)),
        (31, "Contribution avant acquisition", 11, (773, 1933, 3867)),
        (32, "Acquisition totale", 13, (320, 600, 1000)),
        (33, "Résultat avant rémunération", 15, (53, 933, 2467)),
    ]
    for row, label, _src, values in checks:
        put(ws, f"A{row}", label, border=BOX)
        for c, v in zip(SCENARIO_COLS, values):
            bp(ws, f"{c}{row}", v, NF_CHF)
    header(ws, 35, ["Écart (CHF)", *SCENARIO_NAMES, "Commentaire"])
    for offset, (row, label, src, _values) in enumerate(checks):
        r = 36 + offset
        put(ws, f"A{r}", f"Écart {label}", border=BOX)
        for c in SCENARIO_COLS:
            put(ws, f"{c}{r}", f"=ROUND({c}{src},0)-{c}{row}", fmt='0;-0;"0"', border=BOX)
    put(ws, "E37", "Développement : 19 000 / 1,081 = 17 576,32 ⇒ 17 576 (BP 17 577, sans effet sur la suite)",
        font=F_SUB, border=BOX)
    put(ws, "A41", "Statut", font=F_BOLD, border=BOX)
    for c in SCENARIO_COLS:
        put(ws, f"{c}41", f'=IF(SUMPRODUCT(ABS({c}36:{c}40))=0,"OK","ÉCART "&SUMPRODUCT({c}36:{c}40)&" CHF")',
            font=F_BOLD, border=BOX)
    ws.conditional_formatting.add("B36:D40", CellIsRule(operator="notEqual", formula=["0"], fill=FILL_WARN))
    widths(ws, {"A": 50, "B": 17, "C": 17, "D": 17, "E": 52})
    ws.freeze_panes = "B5"
    page_setup(ws)


SENS_CACS = (0, 4, 5, 6, 8, 10, 12, 15)
SENS_RATES = (0.12, 0.14, 0.15, 0.16, 0.18, 0.20, 0.22, 0.24, 0.26)
GRID_COLS = ("B", "C", "D", "E", "F", "G", "H", "I")


def _grid(ws: Worksheet, top: int, label: str, amount_expr: str) -> None:
    """Grille seuil (commandes/mois) taux de contribution × CAC."""
    section(ws, top, label, 9)
    put(ws, f"A{top + 1}", "Contribution ↓ / CAC (CHF) →", font=F_HEADER, fill=FILL_HEADER, border=BOX,
        align=CENTER)
    for col, cac in zip(GRID_COLS, SENS_CACS):
        inp(ws, f"{col}{top + 1}", cac, fmt=NF_CHF2)
    for i, rate in enumerate(SENS_RATES):
        r = top + 2 + i
        inp(ws, f"A{r}", rate, fmt=NF_PCT1)
        for col in GRID_COLS:
            per_order = f"(Panier_HT*$A{r}-{col}${top + 1})"
            put(ws, f"{col}{r}", f'=IF({per_order}>0,ROUNDUP({amount_expr}/{per_order},0),"non rentable")',
                fmt="0", border=BOX, align=Alignment(horizontal="center"))
    first, last = top + 2, top + 1 + len(SENS_RATES)
    rng = f"B{first}:I{last}"
    ws.conditional_formatting.add(
        rng, FormulaRule(formula=[f"(Panier_HT*$A{first}-B${top + 1})<Plancher_CHF"], fill=FILL_WARN)
    )


def _sheet_breakeven(wb: Workbook) -> None:
    ws = wb.create_sheet(S_SEU)
    title(ws, "Seuil de rentabilité et sensibilité (BP §10)",
          "Seuil = charges à couvrir / contribution par commande après CAC, arrondi au-dessus. "
          "Le BP arrondit la contribution au centime avant de diviser.")
    section(ws, 4, "Seuil de rentabilité — scénario central", 6)
    header(ws, 5, ["Indicateur", "Calcul exact", "Méthode BP (centime)", "Valeur BP", "Statut", "Commentaire"])
    put(ws, "A6", "Contribution / commande avant acquisition", border=BOX)
    put(ws, "B6", "=Panier_HT*Taux_contribution", fmt="0.0000", border=BOX)
    put(ws, "C6", "=ROUND(B6,2)", fmt=NF_CHF2, border=BOX)
    bp(ws, "D6", 19.33, NF_CHF2)
    put(ws, "A7", "CAC moyen (central)", border=BOX)
    put(ws, "B7", f"={q(S_HYP, 'C28')}", font=F_LINK, fmt=NF_CHF2, border=BOX)
    put(ws, "C7", "=B7", fmt=NF_CHF2, border=BOX)
    bp(ws, "D7", 6, NF_CHF2)
    put(ws, "A8", "Contribution / commande après CAC", border=BOX)
    put(ws, "B8", "=B6-B7", fmt="0.0000", border=BOX)
    put(ws, "C8", "=ROUND(C6-C7,2)", fmt=NF_CHF2, border=BOX)
    bp(ws, "D8", 13.33, NF_CHF2)
    put(ws, "A9", "Charges fixes à couvrir", border=BOX)
    put(ws, "B9", "=Charges_fixes", fmt=NF_CHF, border=BOX)
    put(ws, "C9", "=Charges_fixes", fmt=NF_CHF, border=BOX)
    bp(ws, "D9", 400, NF_CHF)
    put(ws, "A10", "Seuil en commandes (décimal)", border=BOX)
    put(ws, "B10", "=B9/B8", fmt="0.0000", border=BOX)
    put(ws, "C10", "=C9/C8", fmt="0.0000", border=BOX)
    put(ws, "A11", "Seuil cash d'exploitation (commandes/mois)", font=F_BOLD, border=BOX)
    put(ws, "B11", "=ROUNDUP(B10,0)", font=F_BOLD, fmt="0", fill=FILL_RESULT, border=BOX)
    put(ws, "C11", "=ROUNDUP(C10,0)", font=F_BOLD, fmt="0", fill=FILL_RESULT, border=BOX)
    bp(ws, "D11", 31, "0")
    put(ws, "A12", "Surplus au seuil (CHF)", border=BOX)
    put(ws, "B12", "=B11*B8-B9", fmt=NF_CHF2, border=BOX)
    put(ws, "C12", "=C11*C8-C9", fmt=NF_CHF2, border=BOX)
    put(ws, "A13", "Seuil avec rémunération budgétée (commandes/mois)", font=F_BOLD, border=BOX)
    put(ws, "B13", "=ROUNDUP((B9+Remuneration)/B8,0)", font=F_BOLD, fmt="0", fill=FILL_RESULT, border=BOX)
    put(ws, "C13", "=ROUNDUP((C9+Remuneration)/C8,0)", font=F_BOLD, fmt="0", fill=FILL_RESULT, border=BOX)
    bp(ws, "D13", 181, "0")
    for r in (6, 8, 11, 13):
        put(ws, f"E{r}", f'=IF(C{r}=D{r},IF(ROUND(B{r},2)=D{r},"OK","OK méthode BP — exact : "&ROUND(B{r},2)),"ÉCART")',
            border=BOX)
    put(ws, "F11", "Exact 29,99 ⇒ 30 commandes, mais 30 × 13,334 ne dépasse 400 que de 0,02 CHF : garder 31.",
        font=F_SUB, border=BOX)
    put(ws, "F13", "Hors cotisations sociales et coûts de la rémunération (à ajouter).", font=F_SUB, border=BOX)

    section(ws, 15, "Sensibilité publiée au BP §10", 6)
    header(ws, 16, ["Cas", "Contribution / commande", "Après CAC", "Seuil (commandes)", "Valeur BP (seuil)",
                    "Plancher dur 8 CHF/commande (BP §5) ?"])
    put(ws, "A17", "Contribution 15 %, CAC central", border=BOX)
    put(ws, "B17", "=Panier_HT*Taux_sensibilite", fmt=NF_CHF2, border=BOX)
    put(ws, "C17", f"=B17-{q(S_HYP, 'C28')}", fmt=NF_CHF2, border=BOX)
    put(ws, "D17", '=IF(C17>0,ROUNDUP(Charges_fixes/C17,0),"non rentable")', fmt="0", fill=FILL_RESULT, border=BOX)
    bp(ws, "E17", 56, "0")
    put(ws, "F17", '=IF(C17>=Plancher_CHF,"respecté","NON — sous le plancher dur")', border=BOX)
    put(ws, "A18", "Contribution 22 %, CAC élevé", border=BOX)
    put(ws, "B18", "=Panier_HT*Taux_contribution", fmt=NF_CHF2, border=BOX)
    put(ws, "C18", "=B18-CAC_eleve", fmt=NF_CHF2, border=BOX)
    put(ws, "D18", '=IF(C18>0,ROUNDUP(Charges_fixes/C18,0),"non rentable")', fmt="0", fill=FILL_RESULT, border=BOX)
    bp(ws, "E18", 93, "0")
    put(ws, "F18", '=IF(C18>=Plancher_CHF,"respecté","NON — sous le plancher dur")', border=BOX)
    put(ws, "A19", "Valeurs BP par commande", font=F_SUB, border=BOX)
    bp(ws, "B19", "13,18 / 22 % : 19,33")
    bp(ws, "C19", "7,18 / 4,33")
    ws.conditional_formatting.add("F17:F18", FormulaRule(formula=['LEFT(F17,3)="NON"'], fill=FILL_ALERT))

    _grid(ws, 21, "Seuil (commandes/mois) — charges fixes seules ; orange = contribution après CAC "
                  "sous le plancher dur de 8 CHF", "Charges_fixes")
    _grid(ws, 33, "Seuil (commandes/mois) — charges fixes + rémunération budgétée", "(Charges_fixes+Remuneration)")
    widths(ws, {"A": 46, "B": 18, "C": 20, "D": 18, "E": 22, "F": 40, "G": 12, "H": 12, "I": 12})
    page_setup(ws)


def _tier_points(x: str, step: Decimal, endings: Sequence[Decimal]) -> str:
    """Plus petit point ``n × step + e`` ≥ x (n entier, e parmi ``endings``) en formule tableur."""
    parts = [f"-INT(-(({x})-{e})/{step})*{step}+{e}" for e in endings]
    return parts[0] if len(parts) == 1 else "MIN(" + ",".join(parts) + ")"


def engine_retail_formula(c: str) -> str:
    """Formule du prix public **du moteur** (COH-05) pour la colonne ``c`` de la feuille Prix plancher.

    Prix rentable = MAX(plancher cible m, plancher 8 CHF pour une commande d'une unité), puis grille
    ``rounding_tiers`` de ``config/pricing_rules.v1.yaml`` (même algorithme que
    :func:`pokeshop.pricing.round_up_retail` : point de la tranche ≥ prix, sinon premier point de la
    tranche suivante). La grille est lue dans le fichier de règles livré : classeur et moteur restent
    alignés à chaque régénération.
    """
    tiers = load_rules().pricing.rounding_tiers
    order_floor = f"({c}7+{c}9+{c}10+{c}11+{c}12+Plancher_CHF)/(1/(1+{c}14)-{c}8)"
    p = f"ROUND(MAX({c}17,{order_floor}),6)"

    def from_tier(i: int, x: str) -> str:
        tier = tiers[i]
        best = _tier_points(x, tier.step, tier.endings)
        if i + 1 == len(tiers):
            return best
        upper = tiers[i + 1].min_price
        overflow = round_up_retail(upper, tiers=tiers)  # premier point de la tranche suivante
        return f"IF({best}<{upper},{best},{overflow})"

    formula = from_tier(len(tiers) - 1, p)
    for i in range(len(tiers) - 2, -1, -1):
        formula = f"IF({p}<{tiers[i + 1].min_price},{from_tier(i, p)},{formula})"
    return f'=IF(ISNUMBER({c}17),{formula},"n.c.")'


def _sheet_floor_price(wb: Workbook) -> None:
    ws = wb.create_sheet(S_PRX)
    title(ws, "Prix plancher et contrôle de contribution (BP §4)",
          "P = (C + b + L + R + A) / ((1 − m) / (1 + t) − r). Exemple FICTIF du BP : aucun coût n'est un devis.")
    header(ws, 4, ["Paramètre", "Cas 1 — assujetti (méthode effective)", "Cas 2 — non assujetti (t = 0)",
                   "Unité", "Source"])
    rows: list[tuple[int, str, object, object, str, str, str]] = [
        (5, "Achat net hors TVA récupérable", 140, "=B5", NF_CHF2, "CHF", "BP §4 — exemple fictif"),
        (6, "TVA non récupérable incluse dans C", 0, "=TVA", NF_PCT1, "%", "Cas 2 : 8,1 % non récupérables"),
        (7, "Coût rendu C", "=B5*(1+B6)", "=C5*(1+C6)", NF_CHF2, "CHF", "BP §4 (151,34 dans le cas 2)"),
        (8, "Frais de paiement proportionnels r", 0.025, "=B8", NF_PCT1, "% du TTC", "BP §4"),
        (9, "Frais fixes de paiement b", 0.30, "=B9", NF_CHF2, "CHF", "BP §4"),
        (10, "Préparation/expédition supportée L", 3, "=B10", NF_CHF2, "CHF", "BP §4"),
        (11, "Provision SAV R", 1, "=B11", NF_CHF2, "CHF", "BP §4"),
        (12, "Acquisition attribuée A", 5, "=B12", NF_CHF2, "CHF", "BP §4"),
        (13, "Contribution cible m", 0.20, "=B13", NF_PCT1, "% CA net", "BP §4-§5"),
        (14, "TVA sur ventes t", "=TVA", 0, NF_PCT1, "%", "BP §4"),
    ]
    for row, label, v1, v2, fmt, unit, src in rows:
        put(ws, f"A{row}", label, border=BOX)
        for col, value in (("B", v1), ("C", v2)):
            if isinstance(value, str) and value.startswith("="):
                put(ws, f"{col}{row}", value, fmt=fmt, border=BOX)
            else:
                inp(ws, f"{col}{row}", value, fmt=fmt)
        put(ws, f"D{row}", unit, border=BOX)
        put(ws, f"E{row}", src, border=BOX)
    calc = [
        (15, "Numérateur C + b + L + R + A", "={c}7+{c}9+{c}10+{c}11+{c}12", NF_CHF2),
        (16, "Dénominateur (1 − m)/(1 + t) − r", "=(1-{c}13)/(1+{c}14)-{c}8", "0.000000"),
        (17, "Prix plancher P (exact)", '=IF({c}16>0,{c}15/{c}16,"dénominateur ≤ 0 : marge impossible")', "0.0000"),
        (18, "Prix plancher arrondi au centime", "=ROUND({c}17,2)", NF_CHF2),
    ]
    for row, label, formula, fmt in calc:
        put(ws, f"A{row}", label, font=F_BOLD if row == 18 else F_BASE, border=BOX)
        for c in ("B", "C"):
            put(ws, f"{c}{row}", formula.replace("{c}", c), fmt=fmt, border=BOX,
                font=F_BOLD if row == 18 else F_BASE, fill=FILL_RESULT if row == 18 else None)
    put(ws, "A19", "Lecture BP — pas d'arrondi du prix public", border=BOX)
    inp(ws, "B19", 10, fmt=NF_CHF2, comment="BP : 208,79 → 209,90 implique un pas de 10 CHF et une terminaison 9,90. "
        "Lecture du BP seulement : le moteur applique la grille par tranches (ligne 26).")
    put(ws, "C19", "=B19", fmt=NF_CHF2, border=BOX)
    put(ws, "A20", "Lecture BP — terminaison du prix public", border=BOX)
    inp(ws, "B20", 9.90, fmt=NF_CHF2)
    put(ws, "C20", "=B20", fmt=NF_CHF2, border=BOX)
    put(ws, "A21", "Prix public — lecture BP (pas 10 / 9,90), différente du moteur v1", font=F_BOLD, border=BOX)
    for c in ("B", "C"):
        put(ws, f"{c}21", f"=ROUNDUP(({c}17-{c}20)/{c}19,0)*{c}19+{c}20", font=F_BOLD, fmt=NF_CHF2,
            fill=FILL_RESULT, border=BOX)
    put(ws, "E21", "Concorde avec le moteur sur le cas BP seulement (coût 50 : 89,90 ici, 83,90 moteur)",
        font=F_WARN, border=BOX)
    put(ws, "A22", "Lecture alternative « prochain X,90 » (pas de 1 CHF)", border=BOX)
    for c in ("B", "C"):
        put(ws, f"{c}22", f"=ROUNDUP(({c}17-0.9)/1,0)*1+0.9", fmt=NF_CHF2, border=BOX)
    put(ws, "E22", "Lectures du BP ; règle appliquée par le moteur : ligne 26", font=F_WARN, border=BOX)
    put(ws, "A23", "Valeur BP — prix plancher", border=BOX)
    bp(ws, "B23", 208.79, NF_CHF2)
    bp(ws, "C23", 207.28, NF_CHF2)
    put(ws, "A24", "Valeur BP — prix public arrondi", border=BOX)
    bp(ws, "B24", 209.90, NF_CHF2)
    put(ws, "C24", "n.c.", font=F_SUB, border=BOX)
    put(ws, "A25", "Statut", font=F_BOLD, border=BOX)
    put(ws, "B25", '=IF(AND(B18=B23,ROUND(B21,2)=B24,ROUND(B26,2)=B24),"OK — conforme BP","ÉCART")', font=F_BOLD,
        border=BOX)
    put(ws, "C25", '=IF(C18=C23,"OK — conforme BP","ÉCART")', font=F_BOLD, border=BOX)
    put(ws, "A26", "Prix public moteur v1 (grille rounding_tiers, prix rentable ≥ 8 CHF par commande)", font=F_BOLD,
        border=BOX)
    for c in ("B", "C"):
        put(ws, f"{c}26", engine_retail_formula(c), font=F_BOLD, fmt=NF_CHF2, fill=FILL_RESULT, border=BOX)
    put(ws, "D26", "CHF", border=BOX)
    put(ws, "E26", "pokeshop.pricing (fait foi) : config/pricing_rules.v1.yaml, avant revérification", font=F_SUB,
        border=BOX)

    section(ws, 27, "Contribution d'une vente unitaire à un prix donné (cas 1)", 5)
    header(ws, 28, ["Indicateur", "Prix testé", "Prix public arrondi", "BP à 199,90", "Statut (prix testé)"])
    put(ws, "A29", "Prix public TTC", border=BOX)
    inp(ws, "B29", 199.90, fmt=NF_CHF2)
    put(ws, "C29", "=B26", fmt=NF_CHF2, border=BOX)
    bp(ws, "D29", 199.90, NF_CHF2)
    lines = [
        (30, "Vente nette = P / (1 + t)", "={c}29/(1+$B$14)", NF_CHF2, 184.92),
        (31, "Frais de paiement = r × P + b", "=$B$8*{c}29+$B$9", NF_CHF2, 5.30),
        (32, "Contribution = net − C − paiement − L − R − A", "={c}30-$B$7-{c}31-$B$10-$B$11-$B$12", NF_CHF2,
         30.62),
        (33, "Contribution en % du CA net", "={c}32/{c}30", NF_PCT2, 0.1656),
    ]
    for row, label, formula, fmt, bp_value in lines:
        put(ws, f"A{row}", label, border=BOX)
        for c in ("B", "C"):
            put(ws, f"{c}{row}", formula.replace("{c}", c), fmt=fmt, border=BOX)
        bp(ws, f"D{row}", bp_value, fmt)
        digits = 4 if row == 33 else 2
        put(ws, f"E{row}", f'=IF(ROUND(B{row},{digits})=D{row},"OK","ÉCART")', border=BOX)
    put(ws, "A34", "Cible m atteinte ?", border=BOX)
    for c in ("B", "C"):
        put(ws, f"{c}34", f'=IF({c}33>=$B$13,"oui","NON — sous la cible")', border=BOX)
    put(ws, "A35", "Plancher dur BP §5 respecté (≥ 12 % et ≥ 8 CHF) ?", border=BOX)
    for c in ("B", "C"):
        put(ws, f"{c}35", f'=IF(AND({c}33>=Plancher_pct,{c}32>=Plancher_CHF),"oui","NON — blocage")', border=BOX)
    ws.conditional_formatting.add("B34:C35", FormulaRule(formula=['LEFT(B34,3)="NON"'], fill=FILL_ALERT))
    widths(ws, {"A": 50, "B": 24, "C": 24, "D": 16, "E": 40})
    page_setup(ws)


def _sheet_stock(wb: Workbook) -> None:
    ws = wb.create_sheet(S_STK)
    title(ws, "Stock nécessaire et besoin en fonds de roulement (BP §10)",
          "Les délais du cycle cash sont des HYPOTHÈSES FICTIVES (feuille Hypothèses) à remplacer par les "
          "conditions réelles fournisseur et PSP.")
    section(ws, 4, "Stock consommé par mois", 5)
    header(ws, 5, ["Indicateur", *SCENARIO_NAMES, "Unité"])
    spec = [
        (6, "Commandes par mois", f"={q(S_HYP, '{c}27')}", "0", "commandes"),
        (7, "CA net HT", "={c}6*Panier_HT", NF_CHF, "CHF/mois"),
        (8, "Coût produit en % du CA net", "=Cout_produit_pct", NF_PCT1, "%"),
        (9, "Achats consommés HT", "={c}7*{c}8", NF_CHF, "CHF/mois"),
        (10, "Stock initial (budget BP §3)", "=Stock_initial", NF_CHF, "CHF"),
        (11, "Part d'un mois couverte par le stock initial", "=IF({c}9>0,{c}10/{c}9,0)", NF_PCT1, "%"),
        (12, "Couverture du stock initial", "={c}11*Jours_mois", NF_DEC1, "jours de ventes"),
        (13, "Manque pour financer un mois d'achats", "=MAX(0,{c}9-{c}10)", NF_CHF, "CHF"),
    ]
    for row, label, formula, fmt, unit in spec:
        put(ws, f"A{row}", label, border=BOX)
        for c in SCENARIO_COLS:
            text = formula.replace("{c}", c)
            put(ws, f"{c}{row}", text, font=F_LINK if "Hypothèses" in text else F_BASE, fmt=fmt, border=BOX)
        put(ws, f"E{row}", unit, font=F_SUB, border=BOX)
    put(ws, "A14", "Valeurs BP (central) : CA net 8 788 ; achats 6 152", border=BOX)
    bp(ws, "C14", 6152, NF_CHF)
    put(ws, "D14", '=IF(ROUND(C9,0)=C14,"OK — conforme BP","ÉCART")', border=BOX)

    section(ws, 16, "Besoin en fonds de roulement (BFR) d'exploitation", 5)
    header(ws, 17, ["Indicateur", *SCENARIO_NAMES, "Calcul"])
    spec2 = [
        (18, "Achats consommés par jour", "={c}9/Jours_mois", NF_CHF2, "achats / jours par mois"),
        (19, "CA TTC encaissé par jour", "={c}6*Panier_TTC/Jours_mois", NF_CHF2, "commandes × panier TTC / jours"),
        (20, "Stock immobilisé (prépaiement + jours de stock)", "={c}18*(Delai_prepaiement+Jours_stock)", NF_CHF,
         "achats/jour × (délai + jours de stock)"),
        (21, "Encaissements PSP en attente de versement", "={c}19*Delai_PSP", NF_CHF, "CA/jour × délai PSP"),
        (22, "Crédit fournisseur (ressource)", "={c}18*Credit_fournisseur", NF_CHF, "achats/jour × jours de crédit"),
        (23, "BFR total", "={c}20+{c}21-{c}22", NF_CHF, "stock + PSP − crédit"),
        (24, "Trésorerie mobilisable au lancement (stock + réserve)", f"=Stock_initial+{q(S_BUD, 'B11')}", NF_CHF,
         "budget BP §3"),
        (25, "Écart à financer (BFR − mobilisable)", "=MAX(0,{c}23-{c}24)", NF_CHF, "apport, réassort plus fréquent"),
    ]
    for row, label, formula, fmt, unit in spec2:
        bold = row in (23, 25)
        put(ws, f"A{row}", label, font=F_BOLD if bold else F_BASE, border=BOX)
        for c in SCENARIO_COLS:
            text = formula.replace("{c}", c)
            put(ws, f"{c}{row}", text, font=F_LINK if "Budget" in text else (F_BOLD if bold else F_BASE), fmt=fmt,
                border=BOX, fill=FILL_RESULT if bold else None)
        put(ws, f"E{row}", unit, font=F_SUB, border=BOX)
    ws.conditional_formatting.add("B25:D25", CellIsRule(operator="greaterThan", formula=["0"], fill=FILL_ALERT))
    put(ws, "A27", "Non comptés (prudence) : TVA collectée conservée jusqu'au décompte, TVA d'achat/import "
                   "préfinancée. À affiner avec la fiduciaire.", font=F_SUB, align=WRAP)
    ws.merge_cells("A27:E27")
    ws.row_dimensions[27].height = 28
    widths(ws, {"A": 52, "B": 17, "C": 17, "D": 17, "E": 40})
    page_setup(ws)


NS_WEEKS = 52
NS_FIRST = 41  # première ligne de semaine
NS_LAST = NS_FIRST + NS_WEEKS - 1
#: Colonnes du tableau hebdomadaire (projection, réel saisi, formules du réel, aides).
NS_PROJ = ("B", "C", "D", "E", "F", "G", "H", "I", "J", "K", "L")
NS_REAL_INPUTS = ("M", "N", "O", "P", "Q", "R", "S", "T")
NS_HEADERS = {
    "A": "Semaine",
    "B": "Commandes", "C": "Ventes nettes HT", "D": "Coût historique", "E": "Paiement", "F": "Logistique nette",
    "G": "SAV", "H": "Acquisition", "I": "Contribution après pub", "J": "Charges fixes",
    "K": "Contribution nette", "L": "Cumul projeté",
    "M": "Commandes payées", "N": "Ventes nettes HT", "O": "Coût historique", "P": "Paiement",
    "Q": "Logistique nette", "R": "SAV", "S": "Acquisition", "T": "Charges fixes",
    "U": "Contribution après pub", "V": "Contribution nette", "W": "CUMUL RÉEL (étoile polaire)",
    "X": "Écart cumul réel − projeté", "Y": "Statut stop-loss global",
    "Z": "aide : reprise projetée", "AA": "aide : gel projeté", "AB": "aide : reprise réelle", "AC": "aide : gel réel",
}


def _rng(col: str) -> str:
    """Plage absolue d'une colonne du tableau hebdomadaire."""
    return f"${col}${NS_FIRST}:${col}${NS_LAST}"


def _sheet_north_star(wb: Workbook) -> None:
    """Feuille « Étoile polaire » : contribution nette cumulée hebdomadaire, projetée et réelle."""
    ws = wb.create_sheet(S_NS)
    title(ws, "Étoile polaire — contribution nette cumulée (métrique unique de pilotage)",
          "Contribution nette = ventes nettes HT − coût historique − paiement − logistique − SAV − acquisition − "
          "charges fixes. Ni chiffre d'affaires ni followers.")
    mput(ws, "A3:Y3", LEGEND + " Colonnes M à T : saisir le réel chaque lundi (montants HT, CHF).", font=F_SUB,
         align=WRAP)
    ws.row_dimensions[3].height = 26

    # --- Paramètres -------------------------------------------------------------------------
    section(ws, 5, "Paramètres (à modifier dans la feuille Hypothèses, lignes 44 à 54)", 25)
    params: list[tuple[int, str, str, str, str]] = [
        (6, "Scénario projeté", '=CHOOSE(Scenario_choisi,"Prudent","Central","Développement")', "@",
         "Choix en Hypothèses!B49 (1, 2 ou 3)"),
        (7, "Commandes par mois du scénario", f"=INDEX({q(S_HYP, '$B$27:$D$27')},Scenario_choisi)", "0",
         "BP §10 — simulation, pas une prévision de demande"),
        (8, "CAC moyen du scénario (CHF/commande)", f"=INDEX({q(S_HYP, '$B$28:$D$28')},Scenario_choisi)", NF_CHF2,
         "BP §10"),
        (9, "Commandes par semaine (× 12/52)", "=E7*12/52", NF_DEC2, "Volumes mensuels ramenés à la semaine"),
        (10, "Semaine d'ouverture des ventes", "=Semaine_ouverture", "0",
         "BP §9 — charges fixes dues dès la semaine 1"),
        (11, "Charges fixes par semaine (× 12/52)", "=Charges_fixes*12/52", NF_CHF2, "BP §3 : 400 CHF/mois"),
        (12, "Logistique nette implicite par commande (emballage)",
         "=Panier_HT*(1-Taux_contribution-Cout_produit_pct)-(Frais_paiement_pct*Panier_TTC+Frais_paiement_fixe)"
         "-Provision_SAV", NF_CHF2,
         "Solde des 8 % de CA HT restant après coût produit 70 % et contribution 22 % (BP §10), moins paiement et SAV"),
        (13, "Seuil du stop-loss global (perte ≥ 20 % de la référence du point zéro)", "=-Perte_max", NF_CHF2_RED,
         "Perte de valeur nette ≥ 20 % du capital engagé de référence (point zéro 4 200 ⇒ 840 CHF), approchée par "
         "la contribution nette cumulée ⇒ tout gelé, retour au niveau d'autonomie 1"),
        (14, "Réserve de trésorerie = seuil du stop-loss cash", "=Reserve_cash", NF_CHF,
         "Suivi dans tresorerie_13_semaines.xlsx (lignes 50-51)"),
    ]
    for row, label, formula, fmt, note in params:
        mput(ws, f"A{row}:D{row}", label, border=BOX)
        put(ws, f"E{row}", formula, fmt=fmt, border=BOX, font=F_BOLD)
        mput(ws, f"F{row}:Q{row}", note, font=F_SUB)
    put(ws, "R12", '=IF(E12<0,"INCOHÉRENT : hypothèses incompatibles","OK — cohérent")', font=F_BOLD)

    # --- Synthèse ---------------------------------------------------------------------------
    section(ws, 16, "Synthèse", 25)
    mput(ws, "A17:D17", "Indicateur", font=F_HEADER, fill=FILL_HEADER, border=BOX, align=CENTER)
    put(ws, "E17", "Projection", font=F_HEADER, fill=FILL_HEADER, border=BOX, align=CENTER)
    put(ws, "F17", "Réel", font=F_HEADER, fill=FILL_HEADER, border=BOX, align=CENTER)
    mput(ws, "G17:Q17", "Commentaire / contrôle", font=F_HEADER, fill=FILL_HEADER, border=BOX, align=CENTER)
    real_none = f'COUNT({_rng("V")})=0'
    window_last = "Semaine_ouverture+INT(Jours_validation/7)-1"

    def window_sum(col: str) -> str:
        return f'SUMIFS({_rng(col)},{_rng("A")},">="&Semaine_ouverture,{_rng("A")},"<="&{window_last})'

    summary: list[tuple[int, str, str, str, str, str]] = [
        (18, "Contribution nette cumulée (fin de période / à date)", f"=L{NS_LAST}",
         f'=IF({real_none},"pas de saisie",SUM({_rng("V")}))', NF_CHF2_RED,
         "La métrique étoile polaire : elle doit croître semaine après semaine"),
        (19, "Semaines couvertes", f"=COUNT({_rng('K')})", f"=COUNT({_rng('V')})", "0", ""),
        (20, "Cumul le plus bas", f"=MIN({_rng('L')})", f'=IF({real_none},"",MIN({_rng("W")}))', NF_CHF2_RED,
         "Creux de trésorerie d'exploitation à financer"),
        (21, "Semaine du cumul le plus bas", f"=INDEX({_rng('A')},MATCH(E20,{_rng('L')},0))",
         f'=IF({real_none},"",INDEX({_rng("A")},MATCH(F20,{_rng("W")},0)))', "0", ""),
        (22, "Retour à un cumul positif (semaine)",
         f'=IF(MIN({_rng("Z")})=0,"au-delà de {NS_WEEKS} sem.",MIN({_rng("Z")}))',
         f'=IF({real_none},"",IF(MIN({_rng("AB")})=0,"pas encore",MIN({_rng("AB")})))', "0",
         "Première semaine après le creux où le cumul redevient > 0"),
        (23, "Stop-loss global déclenché ?",
         f'=IF(MIN({_rng("AA")})=0,"non","OUI — semaine "&MIN({_rng("AA")}))',
         f'=IF({real_none},"",IF(MIN({_rng("AC")})=0,"non","OUI — semaine "&MIN({_rng("AC")})&" : TOUT GELER"))',
         "@", "Perte cumulée ≥ seuil (ligne 13). Réarmement : propriétaire uniquement"),
        (24, "Équivalent mensuel en régime (dernière semaine × 52/12)", f"=K{NS_LAST}*52/12", '=""', NF_CHF2,
         ""),
        (25, "Jalon BP §1 : commandes payées sur la fenêtre de validation", f"={window_sum('B')}",
         f'=IF({real_none},"",{window_sum("M")})', NF_DEC1, ""),
        (26, "Jalon commandes atteint ?", '=IF(E25>=Commandes_validation,"oui","NON")',
         '=IF(F25="","",IF(F25>=Commandes_validation,"oui","NON"))', "@", ""),
        (27, "Jalon BP §1 : contribution après publicité sur la fenêtre", f"={window_sum('I')}",
         f'=IF({real_none},"",{window_sum("U")})', NF_CHF2_RED, "Avant charges fixes"),
        (28, "Jalon contribution positive ?", '=IF(E27>0,"oui","NON")', '=IF(F27="","",IF(F27>0,"oui","NON"))',
         "@", "Survente et écoulement de 50 % du stock pilote : suivis hors de ce classeur"),
    ]
    for row, label, proj, real, fmt, note in summary:
        mput(ws, f"A{row}:D{row}", label, border=BOX, font=F_BOLD if row == 18 else F_BASE)
        put(ws, f"E{row}", proj, fmt=fmt, border=BOX, font=F_BOLD, fill=FILL_RESULT if row == 18 else None)
        put(ws, f"F{row}", real, fmt=fmt, border=BOX, font=F_BOLD, fill=FILL_RESULT if row == 18 else None)
        mput(ws, f"G{row}:Q{row}", note, font=F_SUB, border=BOX)
    put(ws, "G24", f'=IF(ROUND(E24-INDEX({q(S_SCN, "$B$15:$D$15")},Scenario_choisi),6)=0,'
                   '"OK — égal au résultat avant rémunération du scénario (feuille Scénarios)","ÉCART")',
        font=F_SUB, border=BOX)
    put(ws, "G25", f'="Semaines "&Semaine_ouverture&" à "&({window_last})&" ("&INT(Jours_validation/7)*7&'
                   '" jours pleins ≈ 60 j du BP) ; cible : "&Commandes_validation&" commandes"',
        font=F_SUB, border=BOX)
    ws.conditional_formatting.add("E23:F23", FormulaRule(formula=['LEFT(E23,3)="OUI"'], fill=FILL_ALERT))
    ws.conditional_formatting.add("E26:F28", FormulaRule(formula=['E26="NON"'], fill=FILL_ALERT))

    # --- Rappel des stop-loss ---------------------------------------------------------------
    section(ws, 30, "Stop-loss du mandat — rappel (pokeshop.stoploss, agent gouvernance, fait foi)", 25)
    for cells, text in (("A31:B31", "Niveau"), ("C31:G31", "Déclencheur"), ("H31:I31", "Seuil"),
                        ("J31:O31", "Effet automatique"), ("P31:U31", "Où le suivre")):
        mput(ws, cells, text, font=F_HEADER, fill=FILL_HEADER, border=BOX, align=CENTER)
    stoploss: list[tuple[str, str, str, str, str, str]] = [
        ("Produit", "Contribution < 12 % du CA net ou < 8 CHF par commande",
         '=TEXT(Plancher_pct,"0%")&" / "&TEXT(Plancher_CHF,"0")&" CHF"', "@", "Vente ou promotion bloquée",
         "Prix plancher l. 35 ; Scénarios l. 21"),
        ("Extension", "> 25 % du budget stock, ou 45 jours sans vente",
         f"={q(S_BUD, 'C26')}", NF_CHF, "Plus de réassort + proposition de démarque", "Budget initial l. 26"),
        ("Publicité", "CAC > contribution par commande sur 7 jours glissants, ou plafond journalier",
         f"=INDEX({q(S_SCN, '$B$18:$D$18')},Scenario_choisi)", NF_CHF2, "Campagne coupée",
         "Scénarios l. 26 ; tableau de bord acquisition"),
        ("Cash", "Cash disponible (solde − précommandes) < réserve", "=Reserve_cash", NF_CHF,
         "Plus d'achat ni de publicité", "tresorerie_13_semaines.xlsx l. 49 à 51"),
        ("Global", "Perte de valeur nette ≥ 20 % du capital engagé de référence (point zéro)", "=-Perte_max", NF_CHF2_RED,
         "TOUT gelé, retour au niveau d'autonomie 1, alerte ; réarmement par la propriétaire", "Cette feuille, colonne Y"),
        ("Temps", "60 jours sans atteindre les seuils de validation", '=Jours_validation&" jours"', "@",
         "Dossier continuer / ajuster / arrêter", "Cette feuille, lignes 25 à 28"),
    ]
    for i, (level, trigger, value, fmt, effect, where) in enumerate(stoploss, start=32):
        mput(ws, f"A{i}:B{i}", level, font=F_BOLD, border=BOX)
        mput(ws, f"C{i}:G{i}", trigger, border=BOX)
        mput(ws, f"H{i}:I{i}", value, fmt=fmt, border=BOX, font=F_BOLD)
        mput(ws, f"J{i}:O{i}", effect, border=BOX)
        mput(ws, f"P{i}:U{i}", where, font=F_SUB, border=BOX)

    # --- Tableau hebdomadaire ---------------------------------------------------------------
    mput(ws, "B39:L39", '="PROJECTION — scénario "&E6&" (formules, ne pas saisir)"', font=F_HEADER,
         fill=FILL_HEADER, align=CENTER, border=BOX)
    mput(ws, "M39:T39", "RÉEL — à saisir chaque lundi (CHF HT, coûts en positif)", font=F_HEADER,
         fill=PatternFill("solid", fgColor="7F6000"), align=CENTER, border=BOX)
    mput(ws, "U39:Y39", "RÉEL — calculé", font=F_HEADER, fill=FILL_HEADER, align=CENTER, border=BOX)
    mput(ws, "Z39:AC39", "Aides de calcul", font=F_SUB, align=CENTER)
    header(ws, 40, [NS_HEADERS[c] for c in ("A", *NS_PROJ, *NS_REAL_INPUTS, "U", "V", "W", "X", "Y")])
    for col in ("Z", "AA", "AB", "AC"):
        put(ws, f"{col}40", NS_HEADERS[col], font=F_SUB, align=CENTER)
    ws.row_dimensions[40].height = 42
    for n in range(1, NS_WEEKS + 1):
        r = NS_FIRST + n - 1
        put(ws, f"A{r}", n, fmt="0", border=BOX, align=Alignment(horizontal="center"))
        proj_cells = {
            "B": f"=IF(A{r}>=Semaine_ouverture,$E$9,0)",
            "C": f"=B{r}*Panier_HT",
            "D": f"=C{r}*Cout_produit_pct",
            "E": f"=B{r}*(Frais_paiement_pct*Panier_TTC+Frais_paiement_fixe)",
            "F": f"=B{r}*$E$12",
            "G": f"=B{r}*Provision_SAV",
            "H": f"=B{r}*$E$8",
            "I": f"=C{r}-D{r}-E{r}-F{r}-G{r}-H{r}",
            "J": "=$E$11",
            "K": f"=I{r}-J{r}",
            "L": f"=SUM($K${NS_FIRST}:K{r})",
        }
        for col, formula in proj_cells.items():
            put(ws, f"{col}{r}", formula, fmt=NF_DEC2 if col == "B" else NF_CHF2_RED if col in "IKL" else NF_CHF2,
                border=BOX, font=F_BOLD if col == "L" else F_BASE)
        for col in NS_REAL_INPUTS:
            inp(ws, f"{col}{r}", None, fmt="0" if col == "M" else NF_CHF2)
        blank = f"COUNT(M{r}:T{r})=0"
        real_cells = {
            "U": f'=IF({blank},"",N{r}-O{r}-P{r}-Q{r}-R{r}-S{r})',
            "V": f'=IF({blank},"",U{r}-T{r})',
            "W": f'=IF(V{r}="","",SUM($V${NS_FIRST}:V{r}))',
            "X": f'=IF(W{r}="","",W{r}-L{r})',
            "Y": (f'=IF(W{r}="","",IF(MIN($W${NS_FIRST}:W{r})<=-Perte_max,"GEL GLOBAL",'
                  f'IF(V{r}<0,"semaine négative","OK")))'),
        }
        for col, formula in real_cells.items():
            put(ws, f"{col}{r}", formula, fmt=None if col == "Y" else NF_CHF2_RED, border=BOX,
                font=F_BOLD if col in "WY" else F_BASE, fill=FILL_RESULT if col == "W" else None,
                align=Alignment(horizontal="center") if col == "Y" else None)
        helpers = {
            "Z": f'=IF(AND(A{r}>$E$21,L{r}>0),A{r},"")',
            "AA": f'=IF(L{r}<=-Perte_max,A{r},"")',
            "AB": f'=IF(W{r}="","",IF(AND(A{r}>$F$21,W{r}>0),A{r},""))',
            "AC": f'=IF(W{r}="","",IF(W{r}<=-Perte_max,A{r},""))',
        }
        for col, formula in helpers.items():
            put(ws, f"{col}{r}", formula, font=F_SUB, fmt="0")
    ws.conditional_formatting.add(f"Y{NS_FIRST}:Y{NS_LAST}",
                                  CellIsRule(operator="equal", formula=['"GEL GLOBAL"'], fill=FILL_ALERT))
    ws.conditional_formatting.add(f"Y{NS_FIRST}:Y{NS_LAST}",
                                  CellIsRule(operator="equal", formula=['"semaine négative"'], fill=FILL_WARN))
    ws.conditional_formatting.add(f"L{NS_FIRST}:L{NS_LAST}",
                                  CellIsRule(operator="lessThanOrEqual", formula=["-Perte_max"], fill=FILL_ALERT))
    widths(ws, {"A": 9, **{c: 12.5 for c in (*NS_PROJ, *NS_REAL_INPUTS, "U", "V", "X")}, "W": 15, "Y": 17,
                "Z": 8, "AA": 8, "AB": 8, "AC": 8})
    ws.freeze_panes = f"B{NS_FIRST}"
    page_setup(ws)


def build_financial_model() -> Workbook:
    """Construit le classeur ``modele_financier.xlsx`` (formules vivantes)."""
    wb = Workbook()
    _sheet_hypotheses(wb)
    _sheet_north_star(wb)
    _sheet_budget(wb)
    _sheet_fixed(wb)
    _sheet_scenarios(wb)
    _sheet_breakeven(wb)
    _sheet_floor_price(wb)
    _sheet_stock(wb)
    wb.calculation.fullCalcOnLoad = True
    return wb


# ---------------------------------------------------------------------------
# Classeur 2 : trésorerie 13 semaines
# ---------------------------------------------------------------------------

WEEK_COLS = tuple(chr(ord("C") + i) for i in range(13))  # C..O
TOTAL_COL = "P"


@dataclass(frozen=True)
class TreasuryParams:
    """Paramètres d'une feuille 13 semaines."""

    start: date | None
    opening_balance: Decimal | None
    minimum_reserve: Decimal | None
    psp_pct: Decimal = Decimal("0.025")
    psp_fixed: Decimal = Decimal("0.30")
    payout_delay_weeks: int = 1
    vat_rate: Decimal = Decimal("0.081")
    opening_preorder_reserve: Decimal = Decimal(0)


#: Ligne du classeur → champ de WeeklyInput (même ordre que la saisie).
INPUT_ROWS: dict[int, str] = {
    17: "sales_ttc",
    18: "orders",
    19: "preorder_sales_ttc",
    20: "preorder_orders",
    25: "psp_in_transit",
    26: "other_inflows",
    29: "purchases_committed",
    30: "purchases_planned",
    31: "vat",
    32: "shipping",
    33: "refunds",
    34: "preorder_refunds",
    35: "advertising",
    36: "fixed_costs",
    37: "setup_costs",
    38: "other_outflows",
    45: "preorders_fulfilled",
}
ROW_LABELS: dict[int, str] = {
    17: "Ventes encaissées TTC (hors précommandes)",
    18: "Nombre de commandes",
    19: "Précommandes encaissées TTC (allocation ferme uniquement)",
    20: "Nombre de précommandes",
    21: "Frais PSP",
    22: "Net à verser par le PSP",
    24: "Versements PSP des ventes (décalés)",
    25: "Versements PSP en transit d'avant la semaine 1",
    26: "Autres encaissements (apport, prêt)",
    27: "Total encaissements",
    29: "Achats de stock engagés (commandes fermes)",
    30: "Achats de stock prévus (non engagés)",
    31: "TVA (décompte AFC)",
    32: "Livraisons (transporteur, étiquettes)",
    33: "Remboursements clients",
    34: "Remboursements de précommandes",
    35: "Publicité",
    36: "Charges fixes (site, compta, divers)",
    37: "Dépenses de lancement (budget initial)",
    38: "Autres décaissements",
    39: "Total décaissements",
    41: "Solde d'ouverture",
    42: "Variation nette de la semaine",
    43: "Solde de clôture",
    45: "Précommandes livrées (réserve libérée, TTC)",
    46: "Réserve précommandes (fin de semaine)",
    47: "Réserve minimale = seuil du stop-loss cash",
    48: "Disponible pour de nouveaux achats",
    49: "Alerte (clôture)",
    50: "Stop-loss cash actif en début de semaine ?",
    51: "Achats prévus + publicité à bloquer (mandat)",
}
ALERT_TEXT = {
    None: "OK",
    "SOLDE_NEGATIF": "SOLDE NÉGATIF",
    "SOUS_RESERVE_MINI": "SOUS RÉSERVE MINI",
    "PRECOMMANDES_NON_COUVERTES": "PRÉCOMMANDES NON COUVERTES",
}


def _d(value: str) -> Decimal:
    return Decimal(value)


def fictitious_example() -> tuple[TreasuryParams, list[WeeklyInput]]:
    """Exemple FICTIF calé sur le budget BP §3 et le plan 90 jours BP §9 (aucune donnée réelle).

    Il illustre le constat de la note de vérification : les charges fixes d'avant
    l'ouverture ne sont pas budgétées, si bien que la réserve de 1 600 CHF (seuil du
    stop-loss cash) est entamée dès la semaine 5. Le test publicitaire prévu en semaine 7
    tombe en stop-loss ; en semaine 9, les précommandes encaissées ne sont pas du cash
    disponible (alerte), ce qui bloque la publicité de la semaine 10.
    """
    params = TreasuryParams(
        start=date(2026, 10, 5),
        opening_balance=_d("8000"),
        minimum_reserve=_d("1600"),
    )
    w = WeeklyInput
    weeks = [
        # J1–15 : étude, dossier B2B, page de présentation (budget BP §3 : site, admin, DA)
        w(fixed_costs=_d("400"), setup_costs=_d("1100")),
        w(setup_costs=_d("1150")),
        # J16–30 : identité, site test, recette ; stock pilote commandé (3 000 CHF)
        w(setup_costs=_d("650")),
        w(purchases_committed=_d("3000")),
        # J31–45 : réception du stock, ouverture douce — 2e mois de charges fixes
        w(fixed_costs=_d("400")),
        w(sales_ttc=_d("279.70"), orders=3, shipping=_d("27")),
        # J46–60 : test publicité prévu… mais la semaine s'ouvre en stop-loss cash
        w(sales_ttc=_d("486.50"), orders=5, shipping=_d("45"), advertising=_d("150")),
        w(sales_ttc=_d("561.40"), orders=6, shipping=_d("54")),
        w(sales_ttc=_d("771.20"), orders=8, preorder_sales_ttc=_d("639.20"), preorder_orders=8, shipping=_d("72"),
          refunds=_d("94.90")),
        # J61–90 : réassort des ventes prouvées, une fois la réserve reconstituée
        w(sales_ttc=_d("842.60"), orders=9, shipping=_d("81"), advertising=_d("150"), fixed_costs=_d("400")),
        w(sales_ttc=_d("968.30"), orders=10, shipping=_d("90"), purchases_committed=_d("900"),
          preorder_refunds=_d("79.90")),
        w(sales_ttc=_d("931.80"), orders=10, shipping=_d("90"), advertising=_d("200"),
          purchases_planned=_d("650")),
        w(sales_ttc=_d("1149.50"), orders=12, shipping=_d("108"), vat=_d("150"), preorders_fulfilled=_d("559.30")),
    ]
    return params, weeks


def _plan_sheet(ws: Worksheet, params: TreasuryParams, weeks: Sequence[WeeklyInput] | None, heading: str,
                note: str) -> None:
    """Feuille 13 semaines (exemple ou modèle vierge) avec formules vivantes."""
    title(ws, heading, note)
    put(ws, "A3", LEGEND.replace("gris italique = valeur publiée dans le BP (référence de contrôle)",
                                 "orange = alerte de trésorerie"), font=F_SUB, align=WRAP)
    ws.merge_cells("A3:P3")
    ws.row_dimensions[3].height = 26
    section(ws, 4, "Paramètres", 16)
    plist: list[tuple[int, str, object, str, str]] = [
        (5, "Lundi de la semaine 1", params.start, NF_DATE, "Date de départ du prévisionnel glissant"),
        (6, "Solde bancaire d'ouverture (CHF)", params.opening_balance, NF_CHF2, "Relevé bancaire du jour"),
        (7, "Réserve minimale = seuil du stop-loss cash (CHF)", params.minimum_reserve, NF_CHF2,
         "BP §3 : réserve 1 600 CHF. Sous ce seuil (solde − précommandes) : plus d'achat prévu ni de pub"),
        (8, "Frais PSP proportionnels", params.psp_pct, NF_PCT2, "BP §4 : 2,5 % (hypothèse) — contrat PSP"),
        (9, "Frais PSP fixes par commande (CHF)", params.psp_fixed, NF_CHF2, "BP §4 : 0,30 (hypothèse)"),
        (10, "Délai de versement PSP (semaines, 0 à 4)", params.payout_delay_weeks, "0",
         "HYPOTHÈSE — à mesurer lors des paiements tests"),
        (11, "Taux TVA ventes (estimation du décompte)", params.vat_rate, NF_PCT1, "BP §4 : 8,1 %"),
        (12, "Réserve précommandes à l'ouverture (CHF)", params.opening_preorder_reserve, NF_CHF2,
         "Précommandes encaissées non livrées au départ"),
    ]
    for row, label, value, fmt, src in plist:
        put(ws, f"A{row}", label, border=BOX)
        inp(ws, f"B{row}", float(value) if isinstance(value, Decimal) else value, fmt=fmt)
        put(ws, f"C{row}", src, font=F_SUB)
    dv = DataValidation(type="whole", operator="between", formula1="0", formula2="4", allow_blank=False,
                        showErrorMessage=True, errorTitle="Délai PSP", error="Entier entre 0 et 4 semaines")
    ws.add_data_validation(dv)
    dv.add("B10")

    header(ws, 14, ["Semaine", "", *[str(i) for i in range(1, 14)], "Total 13 sem."])
    for i, col in enumerate(WEEK_COLS, start=1):
        c = ws[f"{col}14"]
        c.value = i
    put(ws, "A15", "Du", font=F_SUB)
    put(ws, "A16", "Au", font=F_SUB)
    for i, col in enumerate(WEEK_COLS):
        put(ws, f"{col}15", f"=IF($B$5=\"\",\"\",$B$5+7*({col}$14-1))", fmt=NF_DATE, font=F_SUB)
        put(ws, f"{col}16", f"=IF($B$5=\"\",\"\",{col}15+6)", fmt=NF_DATE, font=F_SUB)

    sections = {23: "ENCAISSEMENTS", 28: "DÉCAISSEMENTS", 40: "SOLDE", 44: "RÉSERVES, ALERTE ET STOP-LOSS CASH"}
    section(ws, 13, "", 16)
    ws["A13"].value = "VENTES (saisie)"
    ws["A13"].font = F_SECTION
    for row, text in sections.items():
        section(ws, row, text, 16)

    for row, label in ROW_LABELS.items():
        bold = row in (27, 39, 43, 48, 49, 50, 51)
        put(ws, f"A{row}", label, font=F_BOLD if bold else F_BASE, border=BOX)
    first, last = WEEK_COLS[0], WEEK_COLS[-1]
    for idx, col in enumerate(WEEK_COLS):
        prev = WEEK_COLS[idx - 1] if idx else None
        data = weeks[idx] if weeks is not None and idx < len(weeks) else None
        for row, attr in INPUT_ROWS.items():
            raw = getattr(data, attr) if data is not None else None
            cell_value: object = None
            if raw is not None and raw != 0:
                cell_value = float(raw) if isinstance(raw, Decimal) else raw
            fmt = "0" if attr in ("orders", "preorder_orders") else NF_CHF2
            inp(ws, f"{col}{row}", cell_value, fmt=fmt)
        formulas = {
            21: f"=({col}17+{col}19)*$B$8+({col}18+{col}20)*$B$9",
            22: f"={col}17+{col}19-{col}21",
            24: f"=IF({col}$14-$B$10>=1,INDEX(${first}$22:${last}$22,{col}$14-$B$10),0)",
            27: f"={col}24+{col}25+{col}26",
            39: f"=SUM({col}29:{col}38)",
            41: "=$B$6" if prev is None else f"={prev}43",
            42: f"={col}27-{col}39",
            43: f"={col}41+{col}42",
            46: (f"=MAX(0,$B$12+{col}19-{col}45-{col}34)" if prev is None
                 else f"=MAX(0,{prev}46+{col}19-{col}45-{col}34)"),
            47: "=$B$7",
            48: f"={col}43-{col}46-{col}47",
            49: (f'=IF($B$6="","à saisir",IF({col}43<0,"SOLDE NÉGATIF",IF({col}43<{col}47,"SOUS RÉSERVE MINI",'
                 f'IF({col}48<0,"PRÉCOMMANDES NON COUVERTES","OK"))))'),
            50: ('=IF($B$6="","à saisir",IF($B$6-$B$12<$B$7,"OUI","non"))' if prev is None
                 else f'=IF($B$6="","à saisir",IF({prev}48<0,"OUI","non"))'),
            51: f'=IF({col}50="OUI",{col}30+{col}35,0)',
        }
        for row, formula in formulas.items():
            bold = row in (27, 39, 43, 48, 49, 50, 51)
            fmt = NF_CHF2_RED if row in (42, 43, 48) else NF_CHF2
            text_row = row in (49, 50)
            put(ws, f"{col}{row}", formula, fmt=None if text_row else fmt, border=BOX,
                font=F_BOLD if bold else F_BASE, fill=FILL_RESULT if row == 43 else None,
                align=Alignment(horizontal="center") if text_row else None)
    for row in (17, 18, 19, 20, 21, 22, 24, 25, 26, 27, *range(29, 40), 42, 45, 51):
        fmt = "0" if row in (18, 20) else NF_CHF2
        put(ws, f"{TOTAL_COL}{row}", f"=SUM({first}{row}:{last}{row})", fmt=fmt, font=F_BOLD, border=BOX)
    put(ws, f"{TOTAL_COL}41", f"={first}41", fmt=NF_CHF2, font=F_BOLD, border=BOX)
    put(ws, f"{TOTAL_COL}43", f"={last}43", fmt=NF_CHF2_RED, font=F_BOLD, border=BOX)
    alert_range = f"{first}49:{last}49"
    ws.conditional_formatting.add(
        alert_range, FormulaRule(formula=[f'AND({first}49<>"OK",{first}49<>"à saisir")'], fill=FILL_ALERT)
    )
    ws.conditional_formatting.add(f"{first}50:{last}50",
                                  CellIsRule(operator="equal", formula=['"OUI"'], fill=FILL_ALERT))
    ws.conditional_formatting.add(f"{first}51:{last}51",
                                  CellIsRule(operator="greaterThan", formula=["0"], fill=FILL_ALERT))
    ws.conditional_formatting.add(f"{first}43:{last}43", FormulaRule(formula=[f'AND($B$6<>"",{first}43<{first}47)'],
                                                                      fill=FILL_ALERT))

    section(ws, 53, "SYNTHÈSE", 16)
    alerts = f'COUNTIF({first}49:{last}49,"<>OK")-COUNTIF({first}49:{last}49,"à saisir")'
    summary = [
        (54, "Solde de clôture le plus bas", f"=MIN({first}43:{last}43)", NF_CHF2_RED),
        (55, "Semaine du solde le plus bas", f"=INDEX({first}14:{last}14,MATCH(B54,{first}43:{last}43,0))", "0"),
        (56, "Nombre de semaines en alerte (clôture)", f"={alerts}", "0"),
        (57, "Versements PSP en transit à la fin de l'horizon", f"=SUM({first}22:{last}22)-SUM({first}24:{last}24)",
         NF_CHF2),
        (58, "Frais PSP sur 13 semaines", f"={TOTAL_COL}21", NF_CHF2),
        (59, "TVA contenue dans les ventes encaissées (estimation)",
         f"=({TOTAL_COL}17+{TOTAL_COL}19)*$B$11/(1+$B$11)", NF_CHF2),
        (60, "Semaines ouvertes en stop-loss cash", f'=COUNTIF({first}50:{last}50,"OUI")', "0"),
        (61, "Première semaine en stop-loss cash",
         f'=IF(B60=0,"aucune",INDEX({first}14:{last}14,MATCH("OUI",{first}50:{last}50,0)))', "0"),
        (62, "Achats prévus + publicité à bloquer (13 semaines)", f"={TOTAL_COL}51", NF_CHF2),
    ]
    for row, label, formula, fmt in summary:
        put(ws, f"A{row}", label, font=F_BOLD, border=BOX)
        put(ws, f"B{row}", formula, fmt=fmt, font=F_BOLD, border=BOX, fill=FILL_RESULT)
    put(ws, "C59", "Hors impôt préalable déductible : le décompte réel est établi avec la fiduciaire.", font=F_SUB)
    put(ws, "C62", "Montants à supprimer ou à décaler : le mandat interdit tout achat prévu et toute publicité "
                   "tant que le stop-loss cash est actif. Les achats déjà engagés restent dus.", font=F_SUB)

    widths(ws, {"A": 52, "B": 14, **{c: 12 for c in WEEK_COLS}, TOTAL_COL: 14})
    ws.freeze_panes = "C17"
    page_setup(ws)


def _sheet_howto(ws: Worksheet) -> None:
    title(ws, "Trésorerie glissante 13 semaines — mode d'emploi",
          "BP §3 : distinguer ventes, résultat et cash ; garder une réserve pour les précommandes.")
    lines = [
        ("1.", "Chaque lundi, dupliquer la feuille « À remplir » (ou repartir de la précédente), mettre la date "
               "du lundi en B5 et le solde bancaire réel en B6."),
        ("2.", "Saisir uniquement les cellules jaunes : ventes encaissées et nombre de commandes, précommandes "
               "(seulement sur allocation ferme), achats engagés (commandes fournisseur signées) et prévus, TVA, "
               "livraisons, remboursements, publicité, charges fixes, dépenses de lancement."),
        ("3.", "Les versements PSP sont calculés : ventes TTC − frais (r × montant + b × commandes), décalés du "
               "délai B10. Saisir en ligne 25 les versements attendus des ventes faites avant la semaine 1."),
        ("4.", "Réserve précommandes = précommandes encaissées − livrées − remboursées. Ce montant n'est pas "
               "disponible pour un nouvel achat (BP §3)."),
        ("5.", "Alerte de clôture (ligne 49) : SOLDE NÉGATIF, SOUS RÉSERVE MINI (solde < B7), PRÉCOMMANDES NON "
               "COUVERTES (solde − réserve précommandes − réserve mini < 0)."),
        ("6.", "STOP-LOSS CASH (lignes 50-51) : la réserve minimale B7 vaut 1 600 CHF, la réserve de trésorerie du "
               "budget BP §3. Une semaine qui s'ouvre après une clôture en alerte est en stop-loss : plus aucun achat "
               "prévu ni aucune publicité (ligne 51 = montants à supprimer ou décaler). Les achats engagés restent "
               "dus. Seule la propriétaire peut modifier B7."),
        ("7.", "La feuille « Exemple FICTIF » applique le budget BP §3 et le plan 90 jours BP §9 avec des montants "
               "inventés : ne jamais l'utiliser comme prévision. Elle montre la réserve entamée dès la semaine 5 "
               "(charges fixes non budgétées)."),
        ("8.", "Le moteur Python pokeshop.treasury (plan_from_weekly_inputs + build_forecast) reproduit "
               "exactement ces calculs pour le tableau de bord et les tests ; enforce_cash_stoploss=True y simule "
               "en plus l'application du blocage."),
    ]
    for i, (num, text) in enumerate(lines, start=4):
        put(ws, f"A{i}", num, font=F_BOLD, align=Alignment(vertical="top"))
        put(ws, f"B{i}", text, align=WRAP)
        ws.row_dimensions[i].height = 42
    put(ws, "A13", "Légende", font=F_SECTION)
    inp(ws, "A14", "")
    put(ws, "B14", "Cellule d'entrée (texte bleu sur fond jaune)")
    put(ws, "A15", 1234.5, fmt=NF_CHF2)
    put(ws, "B15", "Formule (texte noir) — ne pas écraser")
    put(ws, "A16", "", fill=FILL_ALERT)
    put(ws, "B16", "Alerte de trésorerie ou stop-loss cash")
    put(ws, "A18", "Validation humaine requise", font=F_SECTION)
    checks = [
        "Seuil du stop-loss cash (B7, 1 600 CHF par défaut) : toute modification est une décision de la propriétaire.",
        "Financement des charges fixes d'avant l'ouverture (≈ 400 à 600 CHF), absentes du budget de 8 000 CHF.",
        "Délai réel de versement et frais du PSP (contrat, paiements tests carte et TWINT).",
        "Périodicité et échéances du décompte TVA (fiduciaire).",
        "Conditions de paiement fournisseur (prépaiement ou terme) avant tout achat engagé.",
    ]
    for i, text in enumerate(checks, start=19):
        put(ws, f"A{i}", "☐", font=F_BOLD)
        put(ws, f"B{i}", text, align=WRAP)
    widths(ws, {"A": 8, "B": 110})


def build_treasury_workbook() -> Workbook:
    """Construit ``tresorerie_13_semaines.xlsx`` (mode d'emploi, exemple FICTIF, modèle vierge)."""
    wb = Workbook()
    _sheet_howto(wb.active)
    wb.active.title = "Mode d'emploi"
    params, weeks = fictitious_example()
    _plan_sheet(wb.create_sheet("Exemple FICTIF"), params, weeks,
                "EXEMPLE FICTIF — trésorerie 13 semaines (aucun chiffre réel)",
                "Montants inventés pour illustrer le plan 90 jours du BP §9. Ne pas utiliser comme prévision.")
    blank = TreasuryParams(start=None, opening_balance=None, minimum_reserve=Decimal("1600"))
    _plan_sheet(wb.create_sheet("À remplir"), blank, None,
                "Trésorerie 13 semaines — {{NOM_BOUTIQUE}}",
                "Saisir les cellules jaunes chaque lundi. Montants en CHF.")
    wb.calculation.fullCalcOnLoad = True
    return wb


# ---------------------------------------------------------------------------
# Recalcul LibreOffice et contrôles
# ---------------------------------------------------------------------------

EXCEL_ERRORS = ("#VALUE!", "#DIV/0!", "#REF!", "#NAME?", "#NULL!", "#NUM!", "#N/A", "Err:")


class RecalcError(RuntimeError):
    """LibreOffice indisponible ou conversion échouée."""


def recalculate(src: Path, dest_dir: Path, timeout: int = 180) -> Path:
    """Recalcule ``src`` avec LibreOffice headless ; écrit la copie recalculée dans ``dest_dir``."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if soffice is None:
        raise RecalcError("soffice introuvable")
    dest_dir.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="lo-finance-") as tmp:
        tmp_path = Path(tmp)
        out = tmp_path / "out"
        cmd = [
            soffice, "--headless", "--norestore", "--nologo",
            f"-env:UserInstallation={(tmp_path / 'profile').as_uri()}",
            "--convert-to", "xlsx:Calc Office Open XML", "--outdir", str(out), str(src.resolve()),
        ]
        try:
            proc = subprocess.run(cmd, capture_output=True, text=True, timeout=timeout, check=False)
        except subprocess.TimeoutExpired as exc:
            raise RecalcError(f"LibreOffice : délai dépassé ({timeout} s)") from exc
        produced = out / src.name
        if proc.returncode != 0 or not produced.exists():
            raise RecalcError(f"LibreOffice n'a pas converti {src.name} : {proc.stdout} {proc.stderr}".strip())
        target = dest_dir / src.name
        shutil.copyfile(produced, target)
    return target


def formula_errors(path: Path) -> list[str]:
    """Cellules en erreur après recalcul (doit être vide)."""
    wb = load_workbook(path, data_only=True)
    errors: list[str] = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and any(cell.value.startswith(e) for e in EXCEL_ERRORS):
                    errors.append(f"{ws.title}!{cell.coordinate}={cell.value}")
    wb.close()
    return errors


def generate(out_dir: Path = OUT_DIR, recalc: bool = True) -> list[Path]:
    """Écrit les deux classeurs ; si ``recalc``, remplace par la version recalculée (formules + valeurs)."""
    out_dir.mkdir(parents=True, exist_ok=True)
    paths = []
    for builder, filename in ((build_financial_model, MODEL_FILE), (build_treasury_workbook, TREASURY_FILE)):
        path = out_dir / filename
        builder().save(path)
        if recalc:
            with tempfile.TemporaryDirectory(prefix="finance-recalc-") as tmp:
                recalculated = recalculate(path, Path(tmp))
                errors = formula_errors(recalculated)
                if errors:
                    raise RecalcError(f"{filename} : erreurs de formule {errors[:10]}")
                shutil.copyfile(recalculated, path)
        paths.append(path)
    return paths


def main(argv: Sequence[str] | None = None) -> int:
    """Point d'entrée CLI."""
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--sortie", type=Path, default=OUT_DIR, help="dossier de sortie")
    parser.add_argument("--sans-recalcul", action="store_true", help="ne pas recalculer avec LibreOffice")
    args = parser.parse_args(argv)
    for path in generate(args.sortie, recalc=not args.sans_recalcul):
        print(f"écrit : {path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
