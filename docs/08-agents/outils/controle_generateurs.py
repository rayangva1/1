"""Contrôle en lecture seule des classeurs générés (agent 12 QA, COH-08).

Les générateurs ``docs/03-finance/generer_classeurs.py`` et ``docs/02-sourcing/outils/generer_comparateur.py``
n'ont pas de mode « contrôle » : par défaut, ils **réécrivent** les classeurs du dépôt. Le QA n'écrit jamais dans
le dépôt ; il passe donc par cet outil, qui :

1. exécute chaque générateur avec ``--sortie`` vers un **dossier temporaire hors du dépôt** (répertoire de travail
   = ce dossier, sans fichier de bytecode) ;
2. compare, feuille par feuille et cellule par cellule, le classeur produit au classeur du dépôt : formules
   (``data_only=False``) et, après recalcul LibreOffice, valeurs (``data_only=True``) ;
3. vérifie que les classeurs du dépôt n'ont **pas changé** pendant le contrôle (empreinte sha256 avant/après) :
   un générateur qui écrirait dans le dépôt est signalé comme une erreur (fermé par défaut).

Code de sortie 0 : classeurs du dépôt identiques à une régénération ; 1 : écart, erreur du générateur ou écriture
dans le dépôt. Aucun fichier du dépôt n'est modifié.

Usage ::

    python docs/08-agents/outils/controle_generateurs.py                 # les deux générateurs, avec recalcul
    python docs/08-agents/outils/controle_generateurs.py --sans-recalcul # formules seulement (sans LibreOffice)
    python docs/08-agents/outils/controle_generateurs.py finance         # un seul générateur
"""

from __future__ import annotations

import argparse
import hashlib
import os
import re
import subprocess
import sys
import tempfile
from collections.abc import Callable, Sequence
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

REPO = Path(__file__).resolve().parents[3]
MAX_DIFFS_PER_FILE = 20
TIMEOUT_SECONDS = 600


@dataclass(frozen=True)
class Generator:
    """Un générateur du dépôt : script, arguments de sortie (dossier temporaire) et classeurs produits."""

    script: Path
    output_args: Callable[[Path], list[str]]
    outputs: tuple[Path, ...]


GENERATORS: dict[str, Generator] = {
    "finance": Generator(
        script=REPO / "docs" / "03-finance" / "generer_classeurs.py",
        output_args=lambda tmp: ["--sortie", str(tmp)],
        outputs=(
            REPO / "docs" / "03-finance" / "modele_financier.xlsx",
            REPO / "docs" / "03-finance" / "tresorerie_13_semaines.xlsx",
        ),
    ),
    "comparateur": Generator(
        script=REPO / "docs" / "02-sourcing" / "outils" / "generer_comparateur.py",
        output_args=lambda tmp: ["--sortie", str(tmp / "COMPARATEUR_OFFRES.xlsx")],
        outputs=(REPO / "docs" / "02-sourcing" / "COMPARATEUR_OFFRES.xlsx",),
    ),
}


@dataclass
class Result:
    """Résultat du contrôle d'un générateur."""

    name: str
    errors: list[str] = field(default_factory=list)
    differences: list[str] = field(default_factory=list)

    @property
    def ok(self) -> bool:
        return not self.errors and not self.differences


def sha256_of(path: Path) -> str | None:
    """Empreinte d'un fichier (``None`` s'il n'existe pas)."""
    if not path.is_file():
        return None
    return hashlib.sha256(path.read_bytes()).hexdigest()


_QUOTED_SHEET_RE = re.compile(r"'([^']+)'!")
_DECIMAL_LITERAL_RE = re.compile(r"(?<![\w.$])(\d+)\.(\d+)(?![\w.])")


def _strip_zeros(match: re.Match[str]) -> str:
    whole, frac = match.group(1), match.group(2).rstrip("0")
    return f"{whole}.{frac}" if frac else whole


def normalize(value: object) -> object:
    """Neutralise la réécriture des formules par LibreOffice au recalcul (même formule, autre graphie).

    ``'Feuille'!A1`` ≡ ``Feuille!A1`` et ``0.50`` ≡ ``0.5`` (``1.00`` ≡ ``1``) ; les valeurs ne sont pas touchées.
    """
    if isinstance(value, str) and value.startswith("="):
        return _DECIMAL_LITERAL_RE.sub(_strip_zeros, _QUOTED_SHEET_RE.sub(r"\1!", value))
    return value


def compare_workbooks(expected: Path, produced: Path, *, data_only: bool) -> list[str]:
    """Écarts entre deux classeurs : feuilles, puis cellules (valeurs si ``data_only``, sinon formules)."""
    mode = "valeur" if data_only else "formule"
    wb_expected = load_workbook(expected, data_only=data_only)
    wb_produced = load_workbook(produced, data_only=data_only)
    diffs: list[str] = []
    if wb_expected.sheetnames != wb_produced.sheetnames:
        diffs.append(f"{expected.name} : feuilles {wb_expected.sheetnames} ≠ {wb_produced.sheetnames}")
    for title in wb_expected.sheetnames:
        if title not in wb_produced.sheetnames:
            continue
        a, b = wb_expected[title], wb_produced[title]
        for row in range(1, max(a.max_row, b.max_row) + 1):
            for col in range(1, max(a.max_column, b.max_column) + 1):
                va, vb = a.cell(row, col).value, b.cell(row, col).value
                if normalize(va) != normalize(vb):
                    diffs.append(f"{expected.name} [{title}!{a.cell(row, col).coordinate}] {mode} : {va!r} ≠ {vb!r}")
    return diffs


def _outside_repo(path: Path) -> bool:
    return not path.resolve().is_relative_to(REPO.resolve())


def _label(path: Path) -> str:
    return str(path.relative_to(REPO)) if path.is_relative_to(REPO) else str(path)


def check_generator(name: str, generator: Generator, *, recalc: bool = True) -> Result:
    """Exécute un générateur dans un dossier temporaire et compare ses classeurs à ceux du dépôt."""
    result = Result(name)
    if not generator.script.is_file():
        result.errors.append(f"{name} : générateur introuvable {generator.script}")
        return result
    before = {path: sha256_of(path) for path in generator.outputs}
    with tempfile.TemporaryDirectory(prefix="qa-generateurs-") as tmp_name:
        tmp = Path(tmp_name)
        if not _outside_repo(tmp):
            result.errors.append(f"{name} : dossier temporaire dans le dépôt ({tmp}), contrôle refusé")
            return result
        command = [sys.executable, str(generator.script), *generator.output_args(tmp)]
        if not recalc:
            command.append("--sans-recalcul")
        env = {**os.environ, "PYTHONDONTWRITEBYTECODE": "1"}
        try:
            proc = subprocess.run(
                command, cwd=tmp, env=env, capture_output=True, text=True, timeout=TIMEOUT_SECONDS, check=False
            )
        except subprocess.TimeoutExpired:
            result.errors.append(f"{name} : délai dépassé ({TIMEOUT_SECONDS} s)")
            proc = None
        if proc is not None and proc.returncode != 0:
            result.errors.append(
                f"{name} : le générateur a échoué ({proc.returncode}) {proc.stdout} {proc.stderr}".strip()
            )
        after = {path: sha256_of(path) for path in generator.outputs}
        for path in generator.outputs:
            if before[path] != after[path]:
                result.errors.append(f"{name} : {_label(path)} a été modifié pendant le contrôle")
        if result.errors:
            return result
        for expected in generator.outputs:
            produced = tmp / expected.name
            if not expected.is_file():
                result.differences.append(f"{expected.name} : absent du dépôt")
                continue
            if not produced.is_file():
                result.errors.append(f"{name} : {expected.name} non produit dans le dossier temporaire")
                continue
            diffs = compare_workbooks(expected, produced, data_only=False)
            if recalc:
                diffs += compare_workbooks(expected, produced, data_only=True)
            result.differences.extend(diffs[:MAX_DIFFS_PER_FILE])
            if len(diffs) > MAX_DIFFS_PER_FILE:
                result.differences.append(f"{expected.name} : … {len(diffs) - MAX_DIFFS_PER_FILE} écart(s) de plus")
    return result


def main(argv: Sequence[str] | None = None) -> int:
    """Point d'entrée : rapport par générateur, code 1 au moindre écart ou à la moindre erreur."""
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("noms", nargs="*", help=f"générateurs à contrôler parmi {sorted(GENERATORS)} (défaut : tous)")
    parser.add_argument("--sans-recalcul", action="store_true", help="comparer les formules seulement")
    args = parser.parse_args(argv)
    unknown = sorted(set(args.noms) - set(GENERATORS))
    if unknown:
        parser.error(f"générateur inconnu {unknown} (attendu : {sorted(GENERATORS)})")
    names = args.noms or list(GENERATORS)
    status = 0
    for name in names:
        result = check_generator(name, GENERATORS[name], recalc=not args.sans_recalcul)
        if result.ok:
            print(f"OK      {name} : classeurs du dépôt identiques à une régénération en dossier temporaire")
            continue
        status = 1
        for line in result.errors:
            print(f"ERREUR  {line}")
        for line in result.differences:
            print(f"ÉCART   {line}")
    return status


if __name__ == "__main__":
    sys.exit(main())
