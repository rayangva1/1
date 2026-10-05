"""Tests de l'outil de contrôle des générateurs (QA en lecture seule, COH-08).

Lancer : ``python -m pytest -q docs/08-agents/outils``.
"""

from __future__ import annotations

import hashlib
import textwrap
from pathlib import Path

import controle_generateurs as cg
import pytest
from openpyxl import Workbook

FAKE_GENERATOR = """
import sys
from pathlib import Path
from openpyxl import Workbook

args = sys.argv[1:]
target = Path(args[args.index("--sortie") + 1])
wb = Workbook()
ws = wb.active
ws.title = "Feuille"
ws["A1"] = "titre"
ws["B2"] = {value!r}
ws["C3"] = "=B2*2"
wb.save(target)
{extra}
"""


def _write_book(path: Path, value: object) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Feuille"
    ws["A1"] = "titre"
    ws["B2"] = value
    ws["C3"] = "=B2*2"
    wb.save(path)
    return path


def _generator(tmp_path: Path, *, value: object = 42, extra: str = "", committed_value: object = 42) -> cg.Generator:
    """Générateur factice : écrit ``CLASSEUR.xlsx`` dans le dossier passé par ``--sortie``."""
    repo_copy = tmp_path / "depot"
    repo_copy.mkdir(parents=True)
    committed = _write_book(repo_copy / "CLASSEUR.xlsx", committed_value)
    script = tmp_path / "generer_factice.py"
    script.write_text(FAKE_GENERATOR.format(value=value, extra=textwrap.dedent(extra)), encoding="utf-8")
    return cg.Generator(
        script=script,
        output_args=lambda tmp: ["--sortie", str(tmp / "CLASSEUR.xlsx")],
        outputs=(committed,),
    )


def _digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


# ------------------------------------------------------------------------------------------- comparaison
def test_identical_regeneration_is_ok(tmp_path):
    result = cg.check_generator("factice", _generator(tmp_path), recalc=False)
    assert result.ok, (result.errors, result.differences)


def test_cell_difference_reported(tmp_path):
    result = cg.check_generator("factice", _generator(tmp_path, value=43), recalc=False)
    assert not result.errors
    assert any("[Feuille!B2]" in d and "42 ≠ 43" in d for d in result.differences)


def test_generator_writing_into_repo_is_an_error(tmp_path):
    """Un générateur qui réécrit le classeur du dépôt (au lieu de --sortie) est refusé, fermé par défaut."""
    gen = _generator(tmp_path)
    committed = gen.outputs[0]
    extra = f"Path({str(committed)!r}).write_bytes(b'ecrase')"
    gen = cg.Generator(
        script=_generator(tmp_path / "bis", extra=extra).script,
        output_args=gen.output_args,
        outputs=gen.outputs,
    )
    result = cg.check_generator("factice", gen, recalc=False)
    assert any("modifié pendant le contrôle" in e for e in result.errors)


def test_failing_generator_is_an_error(tmp_path):
    gen = _generator(tmp_path, extra="sys.exit(3)")
    result = cg.check_generator("factice", gen, recalc=False)
    assert any("a échoué (3)" in e for e in result.errors)


def test_missing_generator_is_an_error(tmp_path):
    gen = cg.Generator(script=tmp_path / "absent.py", output_args=lambda tmp: ["--sortie", str(tmp)], outputs=())
    assert "introuvable" in cg.check_generator("absent", gen).errors[0]


def test_normalize_libreoffice_rewrites():
    assert cg.normalize("='Paramètres'!$B$4*0.50") == cg.normalize("=Paramètres!$B$4*0.5")
    assert cg.normalize("=-INT(-(B1-0.90)/1.00)*10.00+104.90") == "=-INT(-(B1-0.9)/1)*10+104.9"
    assert cg.normalize("=B17") == "=B17"
    assert cg.normalize(0.50) == 0.5
    assert cg.normalize("Lecture 1.00") == "Lecture 1.00"  # texte, pas une formule


def test_main_exit_codes(tmp_path, monkeypatch, capsys):
    monkeypatch.setattr(cg, "GENERATORS", {"factice": _generator(tmp_path)})
    assert cg.main(["--sans-recalcul"]) == 0
    assert "OK      factice" in capsys.readouterr().out
    monkeypatch.setattr(cg, "GENERATORS", {"factice": _generator(tmp_path / "ecart", value=7)})
    assert cg.main(["--sans-recalcul", "factice"]) == 1
    assert "ÉCART" in capsys.readouterr().out
    with pytest.raises(SystemExit):
        cg.main(["inconnu"])


# ------------------------------------------------------------------------------------------- configuration réelle
@pytest.mark.parametrize("name", list(cg.GENERATORS))
def test_real_generators_always_write_to_temporary_dir(name, tmp_path):
    """Chaque générateur réel reçoit --sortie vers le dossier temporaire, jamais sa sortie par défaut (le dépôt)."""
    gen = cg.GENERATORS[name]
    assert gen.script.is_file()
    args = gen.output_args(tmp_path)
    assert "--sortie" in args
    target = Path(args[args.index("--sortie") + 1])
    assert target == tmp_path or target.parent == tmp_path
    assert all(out.is_file() and out.is_relative_to(cg.REPO) for out in gen.outputs)


def test_real_generators_leave_repo_untouched():
    """Exécution réelle (formules seulement, sans LibreOffice) : aucun fichier du dépôt créé ni modifié."""
    folders = sorted({out.parent for gen in cg.GENERATORS.values() for out in gen.outputs})

    outputs = [out for gen in cg.GENERATORS.values() for out in gen.outputs]

    def snapshot() -> tuple[dict[Path, str], set[Path]]:
        listing = {p for folder in folders for p in folder.iterdir() if p.name != "__pycache__"}
        return {out: _digest(out) for out in outputs}, listing

    before = snapshot()
    for name, gen in cg.GENERATORS.items():
        result = cg.check_generator(name, gen, recalc=False)
        assert not result.errors, result.errors
    assert snapshot() == before
