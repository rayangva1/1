"""Import Excel (.xlsx) piloté par le dictionnaire de champs (feuille, ligne d'en-tête, cellule d'export).

Les nombres Excel (flottants) sont convertis par leur représentation décimale la plus
courte (``excel_number``), jamais par le binaire. Les formules sont lues par leur dernière
valeur calculée ; une formule jamais calculée apparaît vide (=> prix absent => quarantaine).
"""

from __future__ import annotations

import io
import zipfile
from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path
from typing import Any

import openpyxl  # type: ignore[import-untyped]
from openpyxl.utils.exceptions import InvalidFileException  # type: ignore[import-untyped]

from pokeshop.importers.base import ImportResult, RawRecord, RawSnapshot, RecordBatch, SourceReadError
from pokeshop.importers.mapping import FileFormat, MappingError, SupplierMapping
from pokeshop.importers.values import CellValue, excel_number

__all__ = ["XlsxReader", "import_xlsx", "excel_cell"]


def excel_cell(value: Any) -> CellValue:
    """Valeur de cellule Excel normalisée (``float`` -> ``int``/``Decimal``, texte rogné)."""
    if value is None or isinstance(value, (bool, int, Decimal, datetime, date)):
        return value
    if isinstance(value, float):
        return excel_number(value)
    if isinstance(value, time):
        return value.isoformat()
    text = str(value).strip()
    return text or None


class XlsxReader:
    """Lit une feuille Excel en enregistrements bruts (numéro de ligne Excel conservé)."""

    def read(self, snapshot: RawSnapshot, mapping: SupplierMapping) -> RecordBatch:
        """Ouvre le classeur en lecture seule des valeurs ; erreurs -> :class:`SourceReadError`."""
        if mapping.file_format is not FileFormat.XLSX:
            raise MappingError(f"dictionnaire {mapping.supplier_id} : format {mapping.file_format.value}, pas xlsx")
        opts = mapping.xlsx
        try:
            wb = openpyxl.load_workbook(io.BytesIO(snapshot.content), data_only=True)
        except (zipfile.BadZipFile, InvalidFileException, KeyError, OSError, ValueError) as exc:
            raise SourceReadError(f"classeur Excel illisible : {type(exc).__name__}: {exc}") from exc
        try:
            if opts.sheet is not None:
                if opts.sheet not in wb.sheetnames:
                    raise SourceReadError(f"feuille {opts.sheet!r} absente (feuilles : {wb.sheetnames})")
                ws = wb[opts.sheet]
            else:
                ws = wb.worksheets[0]
            rows = ws.iter_rows(min_row=opts.header_row, values_only=True)
            first = next(rows, None)
            if first is None:
                raise SourceReadError(f"ligne d'en-tête {opts.header_row} absente")
            header = tuple("" if v is None else str(v).strip() for v in first)
            named = [h for h in header if h]
            if not named:
                raise SourceReadError(f"ligne d'en-tête {opts.header_row} vide")
            if len(set(named)) != len(named):
                raise SourceReadError("en-tête : noms de colonnes en double")
            records: list[RawRecord] = []
            blank = 0
            for offset, row in enumerate(rows, start=1):
                values = [excel_cell(v) for v in row]
                cells: dict[str, CellValue] = {}
                for i, name in enumerate(header):
                    if name:
                        cells[name] = values[i] if i < len(values) else None
                if all(v is None for v in cells.values()):
                    blank += 1
                    continue
                records.append(RawRecord(line_no=opts.header_row + offset, cells=cells))
            source_ts_raw = excel_cell(ws[opts.source_ts_cell].value) if opts.source_ts_cell else None
        finally:
            wb.close()
        return RecordBatch(records=tuple(records), header=tuple(named), source_ts_raw=source_ts_raw, blank_rows=blank)


def import_xlsx(
    source: str | Path | bytes,
    mapping: SupplierMapping | str | Path,
    *,
    now: datetime,
    **kwargs: Any,
) -> ImportResult:
    """Raccourci : importe un classeur Excel (chemin ou contenu) en simulation."""
    from pokeshop.importers.runner import run_import

    return run_import(mapping, source, now=now, **kwargs)
